# Incentivi RS module regression tests (config, import foglio, conferma, calcolo, override, storno, chiusura, export, compensazioni, mio/online, entita)
import io
import json
import os
import time

import pyotp
import pytest
import requests
from dotenv import load_dotenv

load_dotenv("/app/frontend/.env")
BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
TOTP_SECRET = "ZNL4MQSH7OUEIOECQ6TI346P6NWNPNPV"
KEVIN = "9f826340-38a8-44c6-a426-b51c9f5020ff"
DEBORAH = "f32e7569-9bbd-43de-aa55-88d352ae0aa6"
GRAVEDONA = "0916114d-79f9-4f06-8d7e-7e030ebe6cdc"
MORBEGNO = "348642fd-364f-48ab-897b-83c6236b061f"
SONDRIO = "abbe8ad2-8102-4951-b4f7-8db24634a205"
MESE = "2026-06"
STATE = {}


@pytest.fixture(scope="module")
def admin():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": "rsriparazioni@gmail.com", "password": "EnricoRS@26"})
    assert r.status_code == 200, r.text
    mfa = r.json()["mfa_token"]
    for _ in range(3):
        r2 = s.post(f"{API}/auth/login/mfa", json={"mfa_token": mfa, "code": pyotp.TOTP(TOTP_SECRET).now()})
        if r2.status_code == 200:
            break
        time.sleep(31)
    assert r2.status_code == 200, r2.text
    s.headers["Authorization"] = f"Bearer {r2.json()['token']}"
    return s


@pytest.fixture(scope="module")
def kevin(admin):
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": "kevin@cambiaora.local", "password": "Kevin2026!"})
    assert r.status_code == 200, r.text
    d = r.json()
    if d.get("mfa_required"):
        pytest.skip("Kevin has 2FA already, secret unknown")
    s.headers["Authorization"] = f"Bearer {d['token']}"
    e = s.post(f"{API}/auth/2fa/enroll")
    assert e.status_code == 200, e.text
    c = s.post(f"{API}/auth/2fa/enroll/confirm", json={"code": pyotp.TOTP(e.json()["manual_key"]).now()})
    assert c.status_code == 200, c.text
    yield s
    admin.post(f"{API}/users/{KEVIN}/2fa/reset")


class TestConfig:
    def test_get_config(self, admin):
        r = admin.get(f"{API}/incentivi/config")
        assert r.status_code == 200
        c = r.json()
        assert len(c["partecipanti"]) == 4 and len(c["negozi_fondo"]) == 2
        assert "_id" not in c

    def test_put_and_log(self, admin):
        n0 = len(admin.get(f"{API}/incentivi/config/log").json())
        r = admin.put(f"{API}/incentivi/config", json={"gettone_lordo": 30})
        assert r.status_code == 200 and r.json()["gettone_lordo"] == 30
        log = admin.get(f"{API}/incentivi/config/log").json()
        assert len(log) == n0 + 1 and "gettone_lordo" in log[0]["modifiche"]
        r = admin.put(f"{API}/incentivi/config", json={"gettone_lordo": 25})
        assert r.json()["gettone_lordo"] == 25

    def test_premi_over_100(self, admin):
        r = admin.put(f"{API}/incentivi/config", json={"premi_pct": [60, 30, 20]})
        assert r.status_code == 400


class TestImportFoglio:
    body = {"sheet_id": "1bYRn-pAzYXPI1BCnWph8XIMGXhfHsEZz6o_22etuG1M", "gid": "1632338798", "mese": MESE}

    def test_anteprima(self, admin):
        r = admin.post(f"{API}/incentivi/import-foglio", json={**self.body, "anteprima": True})
        assert r.status_code == 200, r.text
        d = r.json()
        print("anteprima righe", len(d["righe"]), "totale", d["totale"], "nuove", d["nuove"])
        assert len(d["righe"]) > 0 and d["totale"] == 2333

    def test_import_and_dedupe(self, admin):
        r = admin.post(f"{API}/incentivi/import-foglio", json={**self.body, "anteprima": False})
        assert r.status_code == 200 and r.json()["importate"] > 0
        r = admin.post(f"{API}/incentivi/import-foglio", json={**self.body, "anteprima": False})
        assert r.json()["importate"] == 0


class TestConfermaCalcolo:
    def test_conferma(self, admin):
        righe = [{"fonte": "eolo", "store_id": GRAVEDONA, "user_id": KEVIN, "importo": 150, "descrizione": "TEST_eolo"},
                 {"fonte": "energia", "store_id": MORBEGNO, "user_id": DEBORAH, "importo": 50, "descrizione": "TEST_energia"}]
        r = admin.post(f"{API}/incentivi/conferma", json={"mese": MESE, "righe": righe})
        assert r.status_code == 200 and r.json()["inseriti"] == 2

    def test_calcolo(self, admin):
        d = admin.get(f"{API}/incentivi/calcolo", params={"mese": MESE}).json()
        f, k, cl = d["fondo"], d["costi"], d["classifica"]
        provv = admin.get(f"{API}/incentivi/provvigioni", params={"mese": MESE}).json()["righe"]
        g_eolo = sum(1 for p in provv if p["fonte"] == "eolo" and p.get("user_id") in (KEVIN, DEBORAH) and p["importo"] > 0)
        expected = round(0.5 * (f["eolo_gruppo_lordo"] - 36.25 * g_eolo), 2)
        print("fondo", f, "gettoni", k["gettoni_n"], "expected", expected)
        assert abs(f["fondo_disponibile"] - expected) < 0.02
        assert k["gettoni_n"] == 2
        assert [c["valore"] for c in cl] == sorted([c["valore"] for c in cl], reverse=True)
        pct_attesi = sum(p for c, p in zip(cl, [50, 30, 20]) if c["valore"] > 0)
        print("premi sum", sum(c["premio_lordo"] for c in cl), "residuo", f["residuo"])
        assert abs(sum(c["premio_lordo"] for c in cl) - f["fondo_disponibile"] * pct_attesi / 100) < 0.05
        assert cl[0]["user_id"] == KEVIN
        assert abs(cl[0]["premio_lordo"] - f["fondo_disponibile"] * 0.5) < 0.02
        STATE["fondo"] = f["fondo_disponibile"]
        STATE["kevin_eolo_id"] = next(p["id"] for p in provv if p.get("descrizione") == "TEST_eolo")

    def test_override(self, admin):
        d = admin.get(f"{API}/incentivi/calcolo", params={"mese": MESE, "override": json.dumps({"eolo_fondo_pct": 30})}).json()
        assert d["fondo"]["fondo_pct"] == 30 and abs(d["fondo"]["fondo_disponibile"] - STATE["fondo"] * 0.6) < 0.05
        assert admin.get(f"{API}/incentivi/config").json()["eolo_fondo_pct"] == 50

    def test_override_bad_type(self, admin):
        r = admin.get(f"{API}/incentivi/calcolo", params={"mese": MESE, "override": json.dumps({"eolo_fondo_pct": "abc"})})
        print("override bad type ->", r.status_code)
        assert r.status_code == 400


class TestStorno:
    def test_storno(self, admin):
        r = admin.post(f"{API}/incentivi/storno", json={"provvigione_id": STATE["kevin_eolo_id"], "mese": "2026-07"})
        assert r.status_code == 200, r.text
        assert r.json()["importo"] == -150 and r.json()["mese"] == "2026-07"
        STATE["storno_id"] = r.json()["id"]
        r = admin.post(f"{API}/incentivi/storno", json={"provvigione_id": STATE["kevin_eolo_id"], "mese": "2026-07"})
        assert r.status_code == 400
        d = admin.get(f"{API}/incentivi/calcolo", params={"mese": "2026-07"}).json()
        assert d["fondo"]["fondo_disponibile"] >= 0


class TestChiusura:
    def test_chiusura_flow(self, admin):
        assert admin.post(f"{API}/incentivi/chiusura/{MESE}").status_code == 200
        r = admin.post(f"{API}/incentivi/conferma", json={"mese": MESE, "righe": [{"fonte": "eolo", "importo": 1}]})
        assert r.status_code == 400 and "chiuso" in r.json()["detail"]
        st = {s["mese"]: s for s in admin.get(f"{API}/incentivi/storico").json()}
        assert st[MESE]["chiuso"] is True
        assert admin.delete(f"{API}/incentivi/chiusura/{MESE}").status_code == 200
        assert admin.get(f"{API}/incentivi/calcolo", params={"mese": MESE}).json()["chiuso"] is False


class TestExport:
    def test_export(self, admin):
        from openpyxl import load_workbook
        r = admin.get(f"{API}/incentivi/export", params={"mese": MESE})
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Fondo", "Classifica", "Estratti conto", "Provvigioni", "Anomalie"]


class TestCompensazioni:
    def test_compensazione(self, admin):
        r = admin.post(f"{API}/incentivi/compensazioni", json={"mese": MESE, "store_id": SONDRIO, "societa": "srls", "descrizione": "TEST_comp", "importo": 30, "categoria": "riparazioni"})
        assert r.status_code == 200
        cid = r.json()["id"]
        d = admin.get(f"{API}/incentivi/calcolo", params={"mese": MESE}).json()
        e = next((x for x in d["estratti_conto"] if x["store_id"] == SONDRIO), None)
        admin.delete(f"{API}/incentivi/compensazioni/{cid}")
        assert e is not None, "Sondrio not in estratti_conto (no provvigioni for Sondrio in month?)"
        assert e["avere_srls"] == 30 and e["saldo_srls"] == -30


class TestMioOnline:
    def test_mio_flow(self, admin, kevin):
        try:
            assert kevin.get(f"{API}/incentivi/config").status_code == 403
            r = kevin.get(f"{API}/incentivi/mio", params={"mese": MESE})
            assert r.status_code == 403, r.text
            assert admin.post(f"{API}/incentivi/online", json={"online": True}).json()["online"] is True
            r = kevin.get(f"{API}/incentivi/mio", params={"mese": MESE})
            assert r.status_code == 200, r.text
            d = r.json()
            for c in d["classifica"]:
                assert not any(k in c for k in ("premio_lordo", "totale_lordo", "valore", "gettoni_lordo"))
            assert d["mio"]["premio_lordo"] > 0 and d["mio"]["posizione"] == 1
        finally:
            admin.post(f"{API}/incentivi/online", json={"online": False})


class TestEntita:
    def test_entita(self, admin):
        try:
            r = admin.post(f"{API}/incentivi/entita", json={"store_id": GRAVEDONA, "entita": "srls", "da": "2026-08"})
            assert r.status_code == 200
            assert {"store_id": GRAVEDONA, "entita": "srls", "da": "2026-08"} in r.json()["entita"]
            k6 = next(c for c in admin.get(f"{API}/incentivi/calcolo", params={"mese": MESE}).json()["classifica"] if c["user_id"] == KEVIN)
            k9 = next(c for c in admin.get(f"{API}/incentivi/calcolo", params={"mese": "2026-09"}).json()["classifica"] if c["user_id"] == KEVIN)
            assert k6["entita"] == "ditta" and k9["entita"] == "srls"
        finally:
            admin.post(f"{API}/incentivi/entita", json={"store_id": GRAVEDONA, "entita": "ditta", "da": "2026-08"})


class TestCompensazioniMeseConProvvigioni:
    def test_compensazione_2026_09(self, admin):
        r = admin.post(f"{API}/incentivi/compensazioni", json={"mese": "2026-09", "store_id": SONDRIO, "societa": "srls", "descrizione": "TEST_comp", "importo": 30, "categoria": "riparazioni"})
        cid = r.json()["id"]
        try:
            d = admin.get(f"{API}/incentivi/calcolo", params={"mese": "2026-09"}).json()
            e = next(x for x in d["estratti_conto"] if x["store_id"] == SONDRIO)
            assert e["avere_srls"] == 30 and e["saldo_srls"] == -30
        finally:
            assert admin.delete(f"{API}/incentivi/compensazioni/{cid}").status_code == 200
        d = admin.get(f"{API}/incentivi/calcolo", params={"mese": "2026-09"}).json()
        assert next(x for x in d["estratti_conto"] if x["store_id"] == SONDRIO)["avere_srls"] == 0
