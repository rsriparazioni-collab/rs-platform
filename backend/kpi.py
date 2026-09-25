import io
import uuid
from calendar import monthrange
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

router = APIRouter()
_d: dict = {}

# (id, label, area, fonte) — fonte: auto | manuale | mista
KPI_VOCI = [
    ("clienti_serviti", "Clienti serviti", "generale", "mista"),
    ("utenza_analizzata", "Utenze analizzate", "energia", "mista"),
    ("swa", "Inserite ENEL SWA / VSA", "energia", "mista"),
    ("voltura", "Volture", "energia", "mista"),
    ("ricontratto", "Ricontratti", "energia", "manuale"),
    ("subentro", "Subentri", "energia", "mista"),
    ("allaccio", "Allacci", "energia", "mista"),
    ("prima_attivazione", "Prima attivazione", "energia", "mista"),
    ("switch_cambiaora", "Switch CambiaOra effettuati", "energia", "mista"),
    ("lavorazione_kevin", "Lavorazioni per Kevin", "energia", "manuale"),
    ("chiamata_enel_eco", "Chiamate ENEL eco", "energia", "manuale"),
    ("mr_servizi_enel", "M&R servizi ENEL", "energia", "manuale"),
    ("investitore_pro", "Clienti Investitore PRO", "energia", "manuale"),
    ("passato_a_silvio", "Clienti passati a Silvio (telefonia)", "energia", "manuale"),
    ("sim_nuova", "SIM numero nuovo — totale", "telefonia", "auto"),
    ("sim_mnp", "SIM portabilità (MNP) — totale", "telefonia", "auto"),
    ("sim_cambio_offerta", "SIM cambio offerta / ricarica", "telefonia", "auto"),
    ("fisso", "Contratti fisso/internet — totale", "telefonia", "auto"),
    ("telefono_proposto", "Telefoni proposti in acquisto", "telefonia", "manuale"),
    ("ricondizionato_venduto", "Ricondizionati venduti", "telefonia", "mista"),
    ("nuovo_venduto", "Telefoni nuovi venduti", "telefonia", "mista"),
    ("accessorio", "Accessori venduti", "telefonia", "mista"),
    ("pellicola", "Pellicole", "telefonia", "manuale"),
    ("eolo_proposto", "Eolo proposto", "internet", "manuale"),
    ("chiamata_eolo", "Chiamate Eolo", "internet", "manuale"),
    ("chiamata_positiva", "Chiamate chiuse positive", "chiamate", "mista"),
    ("chiamata_negativa", "Chiamate chiuse negative", "chiamate", "mista"),
    ("chiamata_non_risposto", "Chiamate non risposto", "chiamate", "mista"),
    ("passato_a_deborah", "Clienti passati a Deborah (utenze)", "chiamate", "manuale"),
    ("rip_ritirata", "Riparazioni ritirate (ingressi)", "riparazioni", "auto"),
    ("rip_normale", "Riparazioni normali", "riparazioni", "auto"),
    ("rip_garanzia", "Riparazioni in garanzia", "riparazioni", "auto"),
    ("rip_ricambio_usato", "Riparazioni con ricambio usato", "riparazioni", "auto"),
    ("rip_scheda_madre", "Riparazioni scheda madre", "riparazioni", "auto"),
    ("rip_non_riparabile", "Non riparabili", "riparazioni", "auto"),
    ("rip_preventivo_rifiutato", "Preventivi non accettati", "riparazioni", "auto"),
    ("rip_senza_tipo", "Riparazioni chiuse senza tipo indicato", "riparazioni", "auto"),
]
VOCI_IDS = {v[0] for v in KPI_VOCI}
RIP_ESITI = {"normale": "rip_normale", "garanzia": "rip_garanzia", "ricambio_usato": "rip_ricambio_usato", "scheda_madre": "rip_scheda_madre",
             "non_riparabile": "rip_non_riparabile", "preventivo_rifiutato": "rip_preventivo_rifiutato"}
RIP_CHIUSI = ("pronto", "pronto_ritiro", "consegnato", "non_riparabile")
MNP_OPERATORI = ["fastweb", "ho", "kena", "iliad", "very", "wind3", "tim", "vodafone", "altro"]
ENEL_MAP = {"swa": "swa", "voltura_swa": "swa", "voltura": "voltura", "subentro": "subentro", "allaccio": "allaccio", "prima_attivazione": "prima_attivazione"}


def setup(**deps):
    _d.update(deps)


async def _current_user(request: Request):
    auth = request.headers.get("authorization", "")
    creds = None
    if auth.lower().startswith("bearer "):
        from fastapi.security import HTTPAuthorizationCredentials
        creds = HTTPAuthorizationCredentials(scheme="Bearer", credentials=auth[7:])
    return await _d["get_current_user"](request, creds)


def _vede_tutti(user: dict) -> bool:
    return user["role"] == "admin" or bool(user.get("can_view_all"))


def _oggi() -> str:
    return datetime.now(timezone.utc).date().isoformat()


class AttivitaInput(BaseModel):
    voce: str
    quantita: int = 1
    client_id: str = ""
    dettaglio: str = ""
    note: str = ""
    data: str = ""


@router.get("/kpi/voci")
async def kpi_voci():
    return {"voci": [{"id": i, "label": l, "area": a, "fonte": f} for i, l, a, f in KPI_VOCI], "mnp_operatori": MNP_OPERATORI}


@router.post("/attivita")
async def registra_attivita(input: AttivitaInput, user: dict = Depends(_current_user)):
    if input.voce not in VOCI_IDS:
        raise HTTPException(status_code=400, detail="Voce non valida")
    if not 1 <= input.quantita <= 50:
        raise HTTPException(status_code=400, detail="Quantità non valida")
    db = _d["db"]
    cli = await db.clients.find_one({"id": input.client_id}, {"_id": 0, "nome": 1, "cognome": 1}) if input.client_id else None
    doc = {"id": str(uuid.uuid4()), "user_id": user["id"], "user_name": user["name"], "store_id": (user.get("store_ids") or [""])[0],
           "voce": input.voce, "quantita": input.quantita, "client_id": input.client_id,
           "client_nome": f"{cli.get('cognome', '')} {cli.get('nome', '')}".strip() if cli else "",
           "dettaglio": input.dettaglio.strip(), "note": input.note.strip(), "data": input.data or _oggi(),
           "at": datetime.now(timezone.utc).isoformat()}
    await db.attivita.insert_one(dict(doc))
    return doc


@router.get("/attivita")
async def lista_attivita(user: dict = Depends(_current_user), data: str = "", user_id: str = ""):
    db = _d["db"]
    q = {"data": data or _oggi()}
    q["user_id"] = user_id if (user_id and _vede_tutti(user)) else user["id"]
    return await db.attivita.find(q, {"_id": 0}).sort("at", -1).to_list(500)


@router.delete("/attivita/{attivita_id}")
async def elimina_attivita(attivita_id: str, user: dict = Depends(_current_user)):
    q = {"id": attivita_id}
    if not _vede_tutti(user):
        q["user_id"] = user["id"]
    res = await _d["db"].attivita.delete_one(q)
    if not res.deleted_count:
        raise HTTPException(status_code=404, detail="Attività non trovata")
    return {"status": "ok"}


def _giorno(iso: str) -> Optional[int]:
    try:
        return int(str(iso)[8:10])
    except (ValueError, TypeError):
        return None


KPI_START_DEFAULT = "2026-10-01"


async def _kpi_start() -> str:
    doc = await _d["db"].settings.find_one({"key": "kpi_start"}, {"_id": 0, "value": 1})
    return (doc or {}).get("value") or KPI_START_DEFAULT


def _griglia_vuota(anno: int, mese: int, ndays: int, kpi_start: str) -> dict:
    voci = [{"id": i, "label": l, "area": a, "fonte": f, "giorni": {}, "totale": 0, "fonti_usate": []} for i, l, a, f in KPI_VOCI]
    return {"anno": anno, "mese": mese, "giorni": ndays, "voci": voci, "totale_mese": 0, "kpi_start": kpi_start, "prima_inizio": True}


class KpiStartInput(BaseModel):
    data: str


@router.get("/kpi/inizio")
async def kpi_inizio(user: dict = Depends(_current_user)):
    return {"kpi_start": await _kpi_start()}


@router.put("/kpi/inizio")
async def set_kpi_inizio(input: KpiStartInput, user: dict = Depends(_current_user)):
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Solo admin")
    try:
        datetime.strptime(input.data, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail="Data non valida (AAAA-MM-GG)")
    await _d["db"].settings.update_one({"key": "kpi_start"}, {"$set": {"key": "kpi_start", "value": input.data}}, upsert=True)
    return {"kpi_start": input.data}


async def calcola_kpi(user_id: str, anno: int, mese: int) -> dict:
    """Griglia voce -> {giorno: n} per un operatore, unendo dati automatici e registro rapido."""
    db = _d["db"]
    ndays = monthrange(anno, mese)[1]
    start = f"{anno:04d}-{mese:02d}-01"
    end = (datetime(anno, mese, ndays) + timedelta(days=1)).strftime("%Y-%m-%d")
    rng = {"$gte": start, "$lt": end}
    kpi_start = await _kpi_start()
    if kpi_start >= end:
        return _griglia_vuota(anno, mese, ndays, kpi_start)
    if kpi_start > start:
        rng["$gte"] = kpi_start
    grid = defaultdict(lambda: defaultdict(int))
    serviti = defaultdict(set)
    fonti = defaultdict(set)
    dinamiche: dict = {}
    target = await db.users.find_one({"id": user_id}, {"_id": 0, "role": 1, "store_ids": 1})

    def add(voce, iso, n=1, client_id=None, fonte="auto"):
        g = _giorno(iso)
        if g and 1 <= g <= ndays and str(iso)[:10] >= kpi_start:
            grid[voce][g] += n
            fonti[voce].add(fonte)
            if client_id:
                serviti[g].add(client_id)

    async for c in db.clients.find({"created_by": user_id, "created_at": rng}, {"_id": 0, "id": 1, "created_at": 1, "gestione": 1, "enel_operazione": 1}):
        add("utenza_analizzata", c["created_at"], client_id=c["id"])
        if c.get("gestione") == "enel" and ENEL_MAP.get(c.get("enel_operazione") or ""):
            add(ENEL_MAP[c["enel_operazione"]], c["created_at"])
    async for l in db.lavorazioni_log.find({"operatore_id": user_id, "created_at": rng, "status": "cambio_effettuato"}, {"_id": 0, "client_id": 1, "created_at": 1}):
        add("switch_cambiaora", l["created_at"], client_id=l.get("client_id"))
    async for l in db.lavorazioni_log.find({"operatore_id": user_id, "created_at": rng, "status": {"$ne": "cambio_effettuato"}}, {"_id": 0, "client_id": 1, "created_at": 1}):
        g = _giorno(l["created_at"])
        if g and l.get("client_id"):
            serviti[g].add(l["client_id"])
    async for s in db.servizi.find({"created_by": user_id, "created_at": rng, "tipo": {"$in": ["sim", "fisso", "internet"]}}, {"_id": 0, "client_id": 1, "created_at": 1, "tipo": 1, "prodotto": 1, "operatore_tel": 1, "note": 1, "tipo_attivazione": 1}):
        op = (s.get("operatore_tel") or "altro").strip().upper() or "ALTRO"
        if s["tipo"] == "sim":
            testo = f"{s.get('prodotto', '')} {s.get('note', '')}".lower()
            ta = s.get("tipo_attivazione") or ("mnp" if any(k in testo for k in ("portabilit", "mnp")) else "nuovo")
            base = {"nuovo": "sim_nuova", "mnp": "sim_mnp", "cambio_offerta": "sim_cambio_offerta"}.get(ta, "sim_nuova")
            add(base, s["created_at"], client_id=s.get("client_id"))
            if base != "sim_cambio_offerta":
                add(f"{base}:{op}", s["created_at"])
                dinamiche[f"{base}:{op}"] = (f"{'SIM numero nuovo' if base == 'sim_nuova' else 'SIM MNP'} — {op}", "telefonia")
        else:
            add("fisso", s["created_at"], client_id=s.get("client_id"))
            add(f"fisso:{op}", s["created_at"])
            dinamiche[f"fisso:{op}"] = (f"Fisso/Internet — {op}", "internet")
    rip_q = {"tipo": "riparazione"}
    if target and target.get("role") == "tecnico":
        pass
    elif target and target.get("store_ids"):
        rip_q["venditore_id"] = {"$in": target["store_ids"]}
    else:
        rip_q["created_by"] = user_id
    async for r in db.servizi.find({**rip_q, "$or": [{"created_at": rng}, {"updated_at": rng}, {"data_uscita": rng}]}, {"_id": 0, "client_id": 1, "created_at": 1, "updated_at": 1, "stato": 1, "data_uscita": 1, "esito_riparazione": 1, "ricambi_usati": 1}):
        if start <= (r.get("created_at") or "")[:10] < end:
            add("rip_ritirata", r["created_at"], client_id=r.get("client_id"))
        fine = r.get("data_uscita") or r.get("updated_at") or ""
        if r.get("stato") in RIP_CHIUSI and start <= fine[:10] < end:
            esito = r.get("esito_riparazione") or ("non_riparabile" if r.get("stato") == "non_riparabile" else "")
            if not esito and any((u.get("nome") or "").lower().startswith(("usato", "rigenerat")) for u in r.get("ricambi_usati") or []):
                esito = "ricambio_usato"
            add(RIP_ESITI.get(esito, "rip_senza_tipo"), fine, client_id=r.get("client_id"))
    vend_map = {"accessori": "accessorio", "telefono_rigenerato": "ricondizionato_venduto", "telefono_usato": "ricondizionato_venduto", "telefono_nuovo": "nuovo_venduto"}
    async for v in db.vendite.find({"created_by": user_id, "data_vendita": rng}, {"_id": 0, "client_id": 1, "data_vendita": 1, "categoria": 1, "quantita": 1}):
        voce = vend_map.get(v.get("categoria") or "")
        if voce:
            add(voce, v["data_vendita"], int(v.get("quantita") or 1), client_id=v.get("client_id"))
        elif v.get("client_id"):
            g = _giorno(v["data_vendita"])
            if g:
                serviti[g].add(v["client_id"])
    esiti = {"contattato": "chiamata_positiva", "positivo": "chiamata_positiva", "chiuso_positivo": "chiamata_positiva", "non_risponde": "chiamata_non_risposto",
             "non_risposto": "chiamata_non_risposto", "non_interessato": "chiamata_negativa", "negativo": "chiamata_negativa", "rifiutato": "chiamata_negativa"}
    async for f in db.followup.find({"user_id": user_id, "created_at": rng}, {"_id": 0, "client_id": 1, "created_at": 1, "esito": 1}):
        voce = esiti.get(f.get("esito") or "")
        if voce:
            add(voce, f["created_at"], client_id=f.get("client_id"))
    async for a in db.attivita.find({"user_id": user_id, "data": rng}, {"_id": 0, "voce": 1, "quantita": 1, "data": 1, "client_id": 1}):
        if a["voce"] != "clienti_serviti":
            add(a["voce"], a["data"], int(a.get("quantita") or 1), client_id=a.get("client_id") or None, fonte="manuale")
        else:
            g = _giorno(a["data"])
            if g:
                grid["clienti_serviti_extra"][g] += int(a.get("quantita") or 1)
                fonti["clienti_serviti"].add("manuale")
    async for a in db.audit_log.find({"user_id": user_id, "at": rng, "entity": {"$in": ["cliente", "servizio"]}, "entity_id": {"$nin": [None, ""]}}, {"_id": 0, "at": 1, "entity": 1, "entity_id": 1}):
        g = _giorno(a["at"])
        if g:
            serviti[g].add(a["entity_id"])
    for g, ids in serviti.items():
        if f"{anno:04d}-{mese:02d}-{g:02d}" < kpi_start:
            continue
        grid["clienti_serviti"][g] += len(ids)
        fonti["clienti_serviti"].add("auto")
    for g, n in grid.pop("clienti_serviti_extra", {}).items():
        grid["clienti_serviti"][g] += n
    out = []
    for vid, label, area, fonte in KPI_VOCI:
        giorni = {str(g): n for g, n in grid.get(vid, {}).items()}
        out.append({"id": vid, "label": label, "area": area, "fonte": fonte, "giorni": giorni, "totale": sum(giorni.values()),
                    "fonti_usate": sorted(fonti.get(vid, []))})
        for did in sorted(k for k in dinamiche if k.startswith(vid + ":")):
            dg = {str(g): n for g, n in grid.get(did, {}).items()}
            out.append({"id": did, "label": "   ↳ " + dinamiche[did][0], "area": dinamiche[did][1] if vid == "fisso" else area, "fonte": "auto",
                        "giorni": dg, "totale": sum(dg.values()), "fonti_usate": ["auto"], "dettaglio": True})
    tot = sum(v["totale"] for v in out if v["id"] != "clienti_serviti" and not v.get("dettaglio"))
    return {"anno": anno, "mese": mese, "giorni": ndays, "voci": out, "totale_mese": tot, "kpi_start": kpi_start}


async def _utenti_kpi(user: dict) -> list:
    db = _d["db"]
    if _vede_tutti(user):
        return await db.users.find({"active": {"$ne": False}, "role": {"$ne": "admin"}}, {"_id": 0, "id": 1, "name": 1, "role": 1, "store_ids": 1}).sort("name", 1).to_list(100)
    return [{"id": user["id"], "name": user["name"], "role": user["role"], "store_ids": user.get("store_ids", [])}]


@router.get("/kpi/utenti")
async def kpi_utenti(user: dict = Depends(_current_user)):
    return await _utenti_kpi(user)


@router.get("/kpi")
async def kpi(anno: int, mese: int, user_id: str = "", user: dict = Depends(_current_user)):
    if not _vede_tutti(user) or not user_id:
        user_id = user["id"] if not _vede_tutti(user) else user_id
    if not user_id:
        raise HTTPException(status_code=400, detail="Seleziona un operatore")
    res = await calcola_kpi(user_id, anno, mese)
    u = await _d["db"].users.find_one({"id": user_id}, {"_id": 0, "name": 1})
    res["user_name"] = (u or {}).get("name", "")
    return res


@router.get("/kpi/riepilogo")
async def kpi_riepilogo(anno: int, mese: int, user: dict = Depends(_current_user)):
    prev_anno, prev_mese = (anno - 1, 12) if mese == 1 else (anno, mese - 1)
    out = []
    for u in await _utenti_kpi(user):
        r = await calcola_kpi(u["id"], anno, mese)
        p = await calcola_kpi(u["id"], prev_anno, prev_mese)
        top = sorted((v for v in r["voci"] if v["totale"] and v["id"] != "clienti_serviti" and not v.get("dettaglio")), key=lambda v: -v["totale"])[:4]
        serviti = next((v["totale"] for v in r["voci"] if v["id"] == "clienti_serviti"), 0)
        serviti_prev = next((v["totale"] for v in p["voci"] if v["id"] == "clienti_serviti"), 0)
        prev_tot = {v["id"]: v["totale"] for v in p["voci"]}
        out.append({"user_id": u["id"], "name": u["name"], "role": u["role"], "totale_mese": r["totale_mese"], "clienti_serviti": serviti,
                    "prev_totale_mese": p["totale_mese"], "prev_clienti_serviti": serviti_prev, "prev_label": f"{prev_mese:02d}/{prev_anno}",
                    "top": [{"label": v["label"], "totale": v["totale"], "prev": prev_tot.get(v["id"], 0)} for v in top]})
    return out


async def promemoria_registro() -> dict:
    """18:30: WhatsApp a chi oggi non ha segnato nulla nel registro rapido (numero utente o telefono avvisi del negozio)."""
    db = _d["db"]
    oggi = _oggi()
    report = []
    async for u in db.users.find({"active": {"$ne": False}, "role": {"$in": ["negozio", "operatore", "tecnico"]}}, {"_id": 0, "id": 1, "name": 1, "telefono": 1, "store_ids": 1}):
        if await db.attivita.count_documents({"user_id": u["id"], "data": oggi}):
            continue
        store = await db.stores.find_one({"id": (u.get("store_ids") or [""])[0]}, {"_id": 0, "telefono_avvisi": 1, "nome": 1}) or {}
        phone = u.get("telefono") or store.get("telefono_avvisi")
        if not phone:
            report.append({"user": u["name"], "esito": "nessun numero"})
            continue
        msg = (f"Ciao {u['name']}! ⏰ Promemoria KPI: oggi non hai ancora segnato nulla nel Registro rapido del gestionale. "
               f"Prima di chiudere, segna al volo accessori, pellicole, proposte Eolo/telefoni e chiamate fatte: ci vogliono 2 minuti. Grazie!")
        try:
            await _d["wa_send"](phone, msg, session=(u.get("store_ids") or ["default"])[0], tipo="promemoria_kpi")
            report.append({"user": u["name"], "esito": "inviato", "phone": phone})
        except Exception as e:
            report.append({"user": u["name"], "esito": f"errore: {e}"[:120]})
    await db.cron_log.insert_one({"job": "promemoria-registro", "at": datetime.now(timezone.utc).isoformat(), "report": report})
    return {"report": report}


@router.post("/kpi/promemoria/invia-ora")
async def promemoria_ora(user: dict = Depends(_current_user)):
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Solo admin")
    return await promemoria_registro()


@router.get("/kpi/export")
async def kpi_export(anno: int, mese: int, user_id: str = "", user: dict = Depends(_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    if not _vede_tutti(user):
        user_id = user["id"]
    utenti = [u for u in await _utenti_kpi(user) if not user_id or u["id"] == user_id]
    wb = Workbook()
    wb.remove(wb.active)
    for u in utenti:
        r = await calcola_kpi(u["id"], anno, mese)
        ws = wb.create_sheet(u["name"][:30])
        ws.append(["giorno:"] + list(range(1, r["giorni"] + 1)) + ["totale mese"])
        ws["A1"].font = Font(bold=True)
        area = None
        for v in r["voci"]:
            if v["area"] != area:
                area = v["area"]
                ws.append([area])
                ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
                ws.cell(row=ws.max_row, column=1).fill = PatternFill("solid", fgColor="E2E8F0")
            ws.append([v["label"]] + [v["giorni"].get(str(g), "") or "" for g in range(1, r["giorni"] + 1)] + [v["totale"]])
        ws.column_dimensions["A"].width = 38
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="kpi_{anno}_{mese:02d}.xlsx"'})
