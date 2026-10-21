"""Report inserimenti mensili: lavorazioni create nel gestionale per mese × negozio × persona (servizi, vendite, ritiri, contratti energia)."""
import io
from collections import defaultdict
from datetime import datetime, timezone

from fastapi import APIRouter, Request
from fastapi.responses import StreamingResponse

router = APIRouter(prefix="/report", tags=["report"])
_d: dict = {}

LAVORAZIONI = [("riparazione", "Riparazioni"), ("vendita_tel", "Vendite telefoni"), ("accessori", "Accessori"), ("sim", "SIM / Mobile"), ("internet", "Fisso / Internet"),
               ("fisso", "Fisso (voce)"), ("vendita", "Vendite (ordini)"), ("ritiro", "Ritiri usato"), ("energia_luce", "Contratti luce"), ("energia_gas", "Contratti gas"), ("cliente", "Nuovi clienti")]
LABEL = dict(LAVORAZIONI)


def setup(**deps):
    _d.update(deps)


def _ym(s: str) -> str:
    return (s or "")[:7]


async def _dati(user: dict, anno: int, mesi: int):
    db = _d["db"]
    fine = datetime.now(timezone.utc)
    if anno < fine.year:
        fine = datetime(anno, 12, 31, tzinfo=timezone.utc)
    ym_list = []
    y, m = fine.year, fine.month
    for _ in range(mesi):
        ym_list.append(f"{y:04d}-{m:02d}")
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    ym_list.reverse()
    ym_set = set(ym_list)
    stores = {s["id"]: s["nome"] async for s in db.stores.find({}, {"_id": 0, "id": 1, "nome": 1})}
    users = {u["id"]: u["name"] async for u in db.users.find({}, {"_id": 0, "id": 1, "name": 1})}
    scope_store = None if user["role"] in ("admin", "tecnico") or user.get("can_view_all") else set(user.get("store_ids") or [])
    cnt = defaultdict(int)  # (ym, store_id, user_id, lav) -> n

    def add(ym, sid, uid, lav):
        if ym in ym_set and (scope_store is None or sid in scope_store):
            cnt[(ym, sid or "", uid if uid in users else "", lav)] += 1

    async for s in db.servizi.find({}, {"_id": 0, "tipo": 1, "venditore_id": 1, "created_by": 1, "created_at": 1}):
        add(_ym(s.get("created_at")), s.get("venditore_id"), s.get("created_by"), "vendita_tel" if s.get("tipo") == "vendita" else (s.get("tipo") or "altro"))
    async for v in db.vendite.find({}, {"_id": 0, "store_id": 1, "created_by": 1, "created_at": 1}):
        add(_ym(v.get("created_at")), v.get("store_id"), v.get("created_by"), "vendita")
    async for r in db.ritiri.find({}, {"_id": 0, "store_id": 1, "created_by": 1, "created_at": 1}):
        add(_ym(r.get("created_at")), r.get("store_id"), r.get("created_by"), "ritiro")
    async for c in db.clients.find({}, {"_id": 0, "venditore_id": 1, "created_by": 1, "created_at": 1, "utenze": 1}):
        add(_ym(c.get("created_at")), c.get("venditore_id"), c.get("created_by"), "cliente")
        for u in c.get("utenze") or []:
            add(_ym(u.get("created_at") or u.get("data_contratto") or c.get("created_at")), c.get("venditore_id"), u.get("created_by") or c.get("created_by"), "energia_gas" if u.get("tipo_bolletta") == "gas" else "energia_luce")
    return ym_list, stores, users, cnt


@router.get("/inserimenti")
async def report_inserimenti(request: Request, anno: int = 0, mesi: int = 12, store_id: str = "", user_id: str = ""):
    user = await _d["current_user"](request)
    return await _inserimenti(anno, mesi, store_id, user_id, user)


async def _inserimenti(anno, mesi, store_id, user_id, user):
    anno = anno or datetime.now(timezone.utc).year
    ym_list, stores, users, cnt = await _dati(user, anno, max(1, min(mesi, 24)))
    righe = defaultdict(lambda: defaultdict(int))   # lav -> ym -> n
    per_negozio = defaultdict(lambda: defaultdict(int))
    per_persona = defaultdict(lambda: defaultdict(int))
    dettaglio = defaultdict(lambda: defaultdict(int))  # (store,user) -> lav -> n (filtered months total)
    for (ym, sid, uid, lav), n in cnt.items():
        if (store_id and sid != store_id) or (user_id and uid != user_id):
            continue
        righe[lav][ym] += n
        per_negozio[sid][ym] += n
        per_persona[uid][ym] += n
        dettaglio[(sid, uid)][lav] += n
    ordine = [k for k, _ in LAVORAZIONI] + sorted(k for k in righe if k not in LABEL)
    return {"mesi": ym_list,
            "lavorazioni": [{"id": k, "label": LABEL.get(k, k), "valori": [righe[k].get(ym, 0) for ym in ym_list], "totale": sum(righe[k].values())} for k in ordine if k in righe or k in LABEL],
            "totale_mese": [sum(righe[k].get(ym, 0) for k in righe) for ym in ym_list],
            "negozi": sorted([{"id": sid, "nome": stores.get(sid, "— senza negozio —"), "valori": [v.get(ym, 0) for ym in ym_list], "totale": sum(v.values())} for sid, v in per_negozio.items()], key=lambda x: -x["totale"]),
            "persone": sorted([{"id": uid, "nome": users.get(uid, "— sconosciuto —"), "valori": [v.get(ym, 0) for ym in ym_list], "totale": sum(v.values())} for uid, v in per_persona.items()], key=lambda x: -x["totale"]),
            "dettaglio": sorted([{"store_id": sid, "negozio": stores.get(sid, "—"), "user_id": uid, "persona": users.get(uid, "—"), **{LABEL.get(k, k): n for k, n in lav.items()}, "totale": sum(lav.values())} for (sid, uid), lav in dettaglio.items()], key=lambda x: -x["totale"]),
            "filtri": {"stores": [{"id": k, "nome": v} for k, v in sorted(stores.items(), key=lambda x: x[1])], "users": [{"id": k, "nome": v} for k, v in sorted(users.items(), key=lambda x: x[1])]}}


@router.get("/inserimenti/export")
async def export_inserimenti(request: Request, anno: int = 0, mesi: int = 12, store_id: str = "", user_id: str = ""):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    user = await _d["current_user"](request)
    d = await _inserimenti(anno, mesi, store_id, user_id, user)
    wb = Workbook(); wb.remove(wb.active)
    for titolo, righe, key in (("Per lavorazione", d["lavorazioni"], "label"), ("Per negozio", d["negozi"], "nome"), ("Per persona", d["persone"], "nome")):
        ws = wb.create_sheet(titolo); ws.append([titolo] + d["mesi"] + ["Totale"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in righe:
            ws.append([r[key]] + r["valori"] + [r["totale"]])
        ws.column_dimensions["A"].width = 30
    ws = wb.create_sheet("Negozio × persona"); cols = ["Negozio", "Persona"] + [l for _, l in LAVORAZIONI] + ["Totale"]; ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in d["dettaglio"]:
        ws.append([r["negozio"], r["persona"]] + [r.get(l, 0) for _, l in LAVORAZIONI] + [r["totale"]])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="inserimenti_mensili.xlsx"'})
