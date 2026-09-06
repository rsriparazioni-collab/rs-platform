import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

VERDE = RGBColor(0x5A, 0x14, 0x6E)
ACCENT = RGBColor(0xBE, 0x28, 0x82)
SCURO = RGBColor(0x1E, 0x29, 0x3B)
GRIGIO = RGBColor(0x64, 0x74, 0x8B)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def add_slide():
    return prs.slides.add_slide(BLANK)

def textbox(slide, x, y, w, h, lines):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (txt, bold, sz, col) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.space_after = Pt(8)
        for r in p.runs:
            r.font.size = Pt(sz)
            r.font.bold = bold
            r.font.color.rgb = col
            r.font.name = "Calibri"
    return tb

def title_slide(s, num, title, subtitle=""):
    bar = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(3), Inches(0.5))
    p = bar.text_frame.paragraphs[0]
    p.text = f"{num:02d}"
    p.runs[0].font.size = Pt(20)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = ACCENT
    textbox(s, 0.6, 0.9, 12, 1, [(title, True, 34, VERDE)])
    if subtitle:
        textbox(s, 0.6, 1.55, 12, 0.6, [(subtitle, False, 16, GRIGIO)])

def bullets_slide(s, items, y=2.1, size=17):
    textbox(s, 0.9, y, 11.6, 5, [(f"▪  {i}", False, size, SCURO) for i in items])

# 1 COVER
s = add_slide()
if os.path.exists("/app/frontend/public/rs-logo.png"):
    s.shapes.add_picture("/app/frontend/public/rs-logo.png", Inches(3.9), Inches(0.8), height=Inches(2.1))
if os.path.exists("/app/frontend/public/cambiaora-logo.jpg"):
    s.shapes.add_picture("/app/frontend/public/cambiaora-logo.jpg", Inches(6.6), Inches(0.8), height=Inches(2.1))
bg = s.shapes.add_textbox(Inches(0), Inches(3.3), Inches(13.33), Inches(3.4))
tf = bg.text_frame
tf.word_wrap = True
for i, (t, sz, c, b) in enumerate([
    ("Gestionale RS & CambiaOra", 44, VERDE, True),
    ("Tutto sotto controllo", 24, ACCENT, True),
    ("Energia, riparazioni, telefonia, magazzino, ritiri usato e WhatsApp dei negozi — in una sola piattaforma", 17, SCURO, False),
    ("Presentazione per la squadra · giugno 2026", 13, GRIGIO, False)]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = t
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(14)
    for r in p.runs:
        r.font.size = Pt(sz)
        r.font.color.rgb = c
        r.font.bold = b
        r.font.name = "Calibri"

# 2 PROBLEMA
s = add_slide()
title_slide(s, 2, "Perché questo progetto", "La situazione prima e dopo")
textbox(s, 0.9, 2.1, 5.8, 4.5, [("PRIMA", True, 18, GRIGIO)] + [(f"▪  {t}", False, 15, SCURO) for t in [
    "Dati sparsi in 7+ fogli Google diversi",
    "Nessun controllo su chi vede cosa",
    "Rinnovi energia segnati a mano, facili da perdere",
    "Numerazione bolle e riparazioni manuale",
    "Magazzino a memoria, giacenze mai certe",
    "Privacy e recensioni inviate a mano una a una"]])
textbox(s, 7, 2.1, 5.8, 4.5, [("OGGI", True, 18, ACCENT)] + [(f"▪  {t}", False, 15, SCURO) for t in [
    "Un solo gestionale condiviso e sempre aggiornato",
    "Permessi per ruolo e per negozio",
    "Alert automatici su scadenze e rinnovi",
    "Numerazione automatica documenti per negozio",
    "Magazzino per negozio con scarico automatico",
    "WhatsApp automatici dal numero del negozio"]])

# 3 PANORAMICA
s = add_slide()
title_slide(s, 3, "I moduli della piattaforma", "Tutto quello che serve al punto vendita, in un posto solo")
bullets_slide(s, [
    "DASHBOARD — KPI del giorno, 'Scadenze della settimana' e alert magazzino sotto scorta",
    "CLIENTI — anagrafica unica multi-servizio con badge Premium (3 step)",
    "ENERGIA — contratti luce/gas, rinnovo automatico a 10 mesi, export Excel, PDF unificato bollette",
    "RIPARAZIONI — stati lavorazione, prezzo consigliato automatico, scheda PDF da firmare",
    "TELEFONIA — SIM / internet / fisso con tracciamento vincoli (es. 48 mesi)",
    "MAGAZZINO — giacenze per negozio, rigenerati/usati, disponibilità cross-negozio, export Excel",
    "RITIRI USATO — bolla PDF con numerazione per negozio (M, SO, G, T, SA, GR)",
    "WHATSAPP — un numero per negozio, privacy + recensione Google automatiche, log invii",
], size=16)

# 4 RIPARAZIONI
s = add_slide()
title_slide(s, 4, "Riparazioni: il flusso completo", "Dal banco alla consegna, tutto tracciato")
bullets_slide(s, [
    "Numero automatico per negozio (es. RM109) — separato dai ritiri",
    "Codice sblocco (simbolo/PIN/password) + account del dispositivo registrati in scheda",
    "Prezzo consigliato automatico: costo ricambio + ricarico + manodopera + IVA",
    "Ricambi prelevati dal magazzino: la giacenza scala da sola",
    "Se il pezzo manca, il sistema dice in quale altro negozio è disponibile",
    "Scheda riparazione PDF stampabile con firme tecnico/cliente",
    "Storico completo importato: 167 riparazioni già dentro il sistema",
], size=16)

# 5 MAGAZZINO E RITIRI
s = add_slide()
title_slide(s, 5, "Magazzino e ritiri usato", "Inventario vero, documenti in regola")
bullets_slide(s, [
    "Giacenza separata per negozio: l'inventario è sempre chiaro",
    "Categorie: display, ricambi, accessori, SIM, rigenerati/usati (prezzo manuale)",
    "Scarico automatico dalle riparazioni, carico/scarico rapido con +/-",
    "Alert 'sotto scorta' in dashboard quando restano 0-2 pezzi",
    "Bolla di ritiro usato: cliente, CF, IMEI, documento, prezzo — PDF pronto da firmare",
    "Numerazione che prosegue i fogli esistenti (Morbegno M8, Sondrio SO7, Gravedona G3...)",
    "Export Excel di magazzino e ritiri per la contabilità",
], size=16)

# 6 WHATSAPP
s = add_slide()
title_slide(s, 6, "WhatsApp: un numero per negozio", "Il cliente riceve i messaggi dal SUO negozio")
bullets_slide(s, [
    "Ogni negozio collega il proprio numero una sola volta (QR o codice numerico)",
    "Link privacy inviato dalla scheda cliente in 1 clic",
    "Richiesta recensione Google automatica 5 minuti dopo — con il link del negozio giusto",
    "Fallback sicuro: se un numero non è collegato, parte il numero principale",
    "Log invii visibile: cosa è partito, quando, da quale numero",
    "Sessioni persistenti sul server: nessun telefono deve restare acceso in negozio",
], size=16)

# 7 RUOLI
s = add_slide()
title_slide(s, 7, "Ruoli e permessi", "Ognuno vede solo quello che gli serve")
textbox(s, 0.9, 2.1, 5.8, 4.5, [(f"▪  {t}", False, 16, SCURO) for t in [
    "ADMIN — tutto, tutti i negozi, compensi",
    "OPERATORE — clienti e servizi su tutti i negozi",
    "NEGOZIO — solo il proprio negozio",
    "TECNICO — tutte le riparazioni, niente vendite"]])
textbox(s, 7, 2.1, 5.8, 4.5, [("Sezioni attivabili per utente:", True, 16, VERDE)] + [(f"▪  {t}", False, 16, SCURO) for t in [
    "Energia", "Riparazioni", "Telefonia",
    "Combinazioni libere (es. solo riparazioni per il tecnico)",
    "Badge 'WhatsApp connessi X/9' in dashboard per l'admin"]])

# 8 NUMERI
s = add_slide()
title_slide(s, 8, "I numeri ad oggi", "Sistema già popolato e operativo")
bullets_slide(s, [
    "646 clienti in anagrafica (importati dai fogli Google)",
    "167 riparazioni storiche importate e numerate",
    "9 punti vendita / collaboratori configurati",
    "6 negozi con numerazione documenti autonoma",
    "Badge Premium: clienti a 3 categorie identificati automaticamente",
    "53 test automatici verdi a ogni modifica",
], size=17)

# 9 ROADMAP
s = add_slide()
title_slide(s, 9, "Prossimi passi", "Cosa arriva dopo")
bullets_slide(s, [
    "Notifiche email automatiche ai negozi sui compensi",
    "Statistiche avanzate per fornitore e periodo",
    "Scrittura bolle direttamente su Google Sheet (archivio storico)",
    "Import storico SIM/internet se servono i fogli",
    "Richiesta recensione automatica alla consegna riparazione",
], size=17)

# 10 CLOSING
s = add_slide()
tb = s.shapes.add_textbox(Inches(0), Inches(2.6), Inches(13.33), Inches(2.5))
tf = tb.text_frame
tf.word_wrap = True
for i, (t, sz, b, c) in enumerate([
    ("Meno carta, meno errori, più tempo per i clienti.", 32, True, VERDE),
    ("Gestionale RS & CambiaOra — tutto sotto controllo. Si parte.", 20, False, SCURO)]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = t
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(16)
    for r in p.runs:
        r.font.size = Pt(sz)
        r.font.bold = b
        r.font.color.rgb = c
        r.font.name = "Calibri"

OUT = "/app/frontend/public/Presentazione_Gestionale_RS_CambiaOra.pptx"
prs.save(OUT)
print("PPTX OK:", OUT, os.path.getsize(OUT), "bytes,", len(prs.slides._sldIdLst), "slide")
