from dotenv import load_dotenv
load_dotenv()

import os
import re
import hmac
import io
import uuid
import hashlib
import ipaddress
import logging
from datetime import datetime, date, timezone, timedelta
from html import escape
from html.parser import HTMLParser
from urllib.parse import urlparse

import csv
import bcrypt
import jwt
import httpx
import pyotp
import qrcode
import base64
import secrets as pysecrets
from cryptography.fernet import Fernet
import pandas as pd
import pymupdf as fitz
from dateutil.relativedelta import relativedelta
from fastapi import FastAPI, APIRouter, HTTPException, Request, Depends, BackgroundTasks, UploadFile, File
from fastapi.responses import StreamingResponse, Response
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel
from typing import Optional, List

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

JWT_ALGORITHM = "HS256"
security = HTTPBearer(auto_error=False)
MFA_SETUP_ALLOWED = ("/api/auth/me", "/api/auth/logout", "/api/auth/2fa/")

SUPPLIERS = [
    "Enel Energia", "Eni Plenitude", "A2A Energia", "Edison Energia", "Hera Comm",
    "Iren Luce e Gas", "E.ON Energia", "Acea Energia", "Engie Italia", "Sorgenia",
    "Illumia", "Alperia", "Estra Energie", "Dolomiti Energia", "Axpo Italia",
    "Green Network Energy", "NeN", "Octopus Energy", "Tate", "Pulsee", "Wekiwi",
    "Repower", "AGSM AIM Energia", "Ascotrade", "Bluenergy Group", "Unogas",
    "Gas Sales", "Vivigas", "Energit", "Egea", "EstEnergy", "Optima Italia",
    "Audax Energia", "Eva Energia", "Nord Energia", "Altri (maggior tutela)"
]

LAVORAZIONI = [
    "cambiare", "non_cambiare", "cambio_effettuato", "in_quotazione",
    "richieste_bollette", "in_attesa_ok", "problema_tecnico", "da_quotare",
    "contattare_cliente", "non_vuole_cambiare", "passa_in_negozio",
    "attesa_documenti", "rinnovato"
]

# ---------------- Auth helpers ----------------

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))

def create_token(user_id: str) -> str:
    payload = {"sub": user_id, "type": "access",
               "exp": datetime.now(timezone.utc) + timedelta(hours=24)}
    return jwt.encode(payload, os.environ["JWT_SECRET"], algorithm=JWT_ALGORITHM)

def serialize_user(u: dict) -> dict:
    return {"id": u["id"], "name": u["name"], "email": u["email"], "role": u["role"],
            "store_ids": u.get("store_ids", []), "can_view_all": u.get("can_view_all", False),
            "active": u.get("active", True), "totp_enabled": bool(u.get("totp_enabled")),
            "sections": u.get("sections") or ["energia", "riparazioni", "telefonia"]}

# ---------------- 2FA (TOTP, Google/Microsoft Authenticator) ----------------
TOTP_ISSUER = "Gestionale RS & CambiaOra"
TOTP_CODE_RE = re.compile(r"^\d{6}$")

def _fernet() -> Fernet:
    return Fernet(os.environ["TOTP_ENCRYPTION_KEY"].encode())

def totp_encrypt(secret: str) -> str:
    return _fernet().encrypt(secret.encode()).decode()

def totp_decrypt(value: str) -> str:
    return _fernet().decrypt(value.encode()).decode()

def totp_qr_data_uri(uri: str) -> str:
    buf = io.BytesIO()
    qrcode.make(uri).save(buf, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()

def make_recovery_codes(count: int = 8) -> tuple:
    plain = [pysecrets.token_hex(4).upper() + "-" + pysecrets.token_hex(4).upper() for _ in range(count)]
    stored = [{"hash": hash_password(c), "used_at": None} for c in plain]
    return plain, stored

def create_mfa_token(user_id: str) -> str:
    payload = {"sub": user_id, "type": "mfa", "exp": datetime.now(timezone.utc) + timedelta(minutes=5)}
    return jwt.encode(payload, os.environ["JWT_SECRET"], algorithm=JWT_ALGORITHM)

async def verify_totp_or_recovery(user: dict, supplied: str) -> bool:
    supplied = supplied.strip().upper().replace(" ", "")
    if TOTP_CODE_RE.fullmatch(supplied):
        totp = pyotp.TOTP(totp_decrypt(user["totp_secret_enc"]), interval=30)
        now = datetime.now(timezone.utc)
        if not totp.verify(supplied, for_time=now, valid_window=1):
            return False
        current = totp.timecode(now)
        previous = user.get("totp_last_timecode", -1)
        if current <= previous:
            return False
        res = await db.users.update_one({"id": user["id"], "totp_last_timecode": user.get("totp_last_timecode")},
                                        {"$set": {"totp_last_timecode": current}})
        return res.modified_count == 1
    for item in user.get("recovery_codes", []):
        if item.get("used_at") is None and verify_password(supplied, item["hash"]):
            res = await db.users.update_one(
                {"id": user["id"], "recovery_codes": {"$elemMatch": {"hash": item["hash"], "used_at": None}}},
                {"$set": {"recovery_codes.$.used_at": datetime.now(timezone.utc).isoformat()}})
            return res.modified_count == 1
    return False

async def get_current_user(request: Request, creds: HTTPAuthorizationCredentials = Depends(security)) -> dict:
    token = request.cookies.get("gu_token") or (creds.credentials if creds else None)
    if not token:
        raise HTTPException(status_code=401, detail="Non autenticato")
    try:
        payload = jwt.decode(token, os.environ["JWT_SECRET"], algorithms=[JWT_ALGORITHM])
        user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0})
        if not user or not user.get("active", True):
            raise HTTPException(status_code=401, detail="Utente non trovato")
        if not user.get("totp_enabled") and not request.url.path.startswith(MFA_SETUP_ALLOWED):
            raise HTTPException(status_code=403, detail="Attiva la verifica in due passaggi per continuare",
                                headers={"X-MFA-Setup-Required": "1"})
        request.state.user = user
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sessione scaduta")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token non valido")

def require_admin(user: dict = Depends(get_current_user)) -> dict:
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Solo amministratore")
    return user

def client_scope_filter(user: dict) -> dict:
    if user["role"] == "admin" or user.get("can_view_all"):
        return {}
    return {"venditore_id": {"$in": user.get("store_ids", [])}}

# ---------------- Business logic ----------------

def effective_pagato(pagato: bool, last_payment_date: Optional[str]) -> bool:
    if not pagato:
        return False
    if not last_payment_date:
        return True
    paid_on = None
    try:
        paid_on = date.fromisoformat(last_payment_date[:10])
    except ValueError:
        return True
    return date.today() < paid_on + relativedelta(months=6)

def compute_dates(c: dict) -> dict:
    dc = c.get("data_contratto")
    if dc:
        try:
            d0 = date.fromisoformat(dc[:10])
            att = d0 + relativedelta(months=2)
            rin = d0 + relativedelta(months=10)
            scad = att + relativedelta(months=12)
            c["data_attivazione"] = att.isoformat()
            c["data_rinnovo"] = rin.isoformat()
            c["data_scadenza"] = scad.isoformat()
            c["giorni_al_rinnovo"] = (rin - date.today()).days
        except ValueError:
            pass
    c["pagato_effettivo"] = effective_pagato(c.get("pagato", False), c.get("last_payment_date"))
    return c

async def build_alerts(user: dict) -> dict:
    scope = client_scope_filter(user)
    projection = {"_id": 0, "id": 1, "nome": 1, "cognome": 1, "data_contratto": 1,
                  "lavorazione": 1, "tipo_bolletta": 1, "pagato": 1, "last_payment_date": 1}
    clients = await db.clients.find(scope, projection).to_list(5000)
    stores = await db.stores.find({}, {"_id": 0}).to_list(500)
    rinnovi, pagamenti_clienti, pagamenti_negozi = [], [], []
    for c in clients:
        compute_dates(c)
        gr = c.get("giorni_al_rinnovo")
        if gr is not None and -30 <= gr <= 45 and c.get("lavorazione") != "rinnovato":
            rinnovi.append({"client_id": c["id"], "nome": c.get("nome", ""), "cognome": c.get("cognome", ""),
                            "data_rinnovo": c["data_rinnovo"], "giorni": gr,
                            "tipo_bolletta": c.get("tipo_bolletta", ""), "lavorazione": c.get("lavorazione", "")})
        if c.get("pagato") and not c.get("pagato_effettivo"):
            pagamenti_clienti.append({"client_id": c["id"], "nome": c.get("nome", ""), "cognome": c.get("cognome", ""),
                                      "last_payment_date": c.get("last_payment_date")})
    for s in stores:
        if s.get("pagato") and not effective_pagato(s.get("pagato", False), s.get("last_payment_date")):
            pagamenti_negozi.append({"store_id": s["id"], "nome": s.get("nome", ""),
                                     "referente": s.get("referente", ""), "last_payment_date": s.get("last_payment_date")})
    rinnovi.sort(key=lambda x: x["giorni"])
    mscope = magazzino_scope(user)
    mscope["quantita"] = {"$lte": 2}
    low = await db.magazzino.find(mscope, {"_id": 0, "id": 1, "nome": 1, "categoria": 1, "quantita": 1, "store_id": 1}).to_list(500)
    smap = {s["id"]: s.get("nome", "") for s in stores}
    sotto_scorta = [{"id": i["id"], "nome": i["nome"], "categoria": i.get("categoria", ""),
                     "quantita": i.get("quantita", 0), "store_name": smap.get(i.get("store_id", ""), "-")} for i in low]
    await attach_followups(rinnovi, "client", "client_id", "rinnovo")
    await attach_followups(pagamenti_clienti, "client", "client_id", "pagamento")
    return {"rinnovi": rinnovi, "pagamenti_clienti": pagamenti_clienti, "pagamenti_negozi": pagamenti_negozi,
            "sotto_scorta": sotto_scorta}

# ---------------- Email (managed Resend) ----------------

EMAIL_BASE_URL = "https://integrations.emergentagent.com"
EMAIL_KEY = os.environ["EMERGENT_EMAIL_KEY"]
EMAIL_FROM_NAME = os.environ["EMAIL_FROM_NAME"]
EMAIL_REPLY_TO = os.environ.get("EMAIL_REPLY_TO")

_SHORTENERS = ("bit.ly", "tinyurl.com", "t.co", "is.gd", "cutt.ly", "goo.gl", "rebrand.ly")
_CRED_ASK = ("reply with your password", "reply with the code", "send your password", "cvv",
             "send us your password", "enter your password below", "confirm your card number",
             "your full card number", "seed phrase", "recovery phrase", "verify your card",
             "social security number", "confirm your bank details")
_HOSTISH = re.compile(r"\b(?:https?://)?((?:[a-z0-9-]+\.)+[a-z]{2,})", re.I)

def _host_ok(host: str) -> bool:
    if not host or "xn--" in host:
        return False
    try:
        ipaddress.ip_address(host)
        return False
    except ValueError:
        pass
    return not any(host == s or host.endswith("." + s) for s in _SHORTENERS)

def _same_site(shown: str, real: str) -> bool:
    return shown == real or real.endswith("." + shown) or shown.endswith("." + real)

class _EmailScan(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags, self.urls, self.anchors = set(), [], []
        self._href, self._text = None, []
    def handle_starttag(self, tag, attrs):
        self.tags.add(tag.lower())
        self.urls += [v for k, v in attrs if k.lower() in ("href", "src") and v]
        if tag.lower() == "a":
            self._href = dict((k.lower(), v) for k, v in attrs).get("href")
            self._text = []
    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == "a" and self._href is not None:
            self.anchors.append((self._href, "".join(self._text)))
            self._href, self._text = None, []

def _assert_safe_email(subject: str, html: str) -> None:
    scan = _EmailScan(); scan.feed(html)
    if scan.tags & {"form", "input", "textarea", "select"}:
        raise ValueError("No forms or input fields in email (G2)")
    body = f"{subject}\n{html}".lower()
    for p in _CRED_ASK:
        if p in body:
            raise ValueError(f"Email asks the recipient for credentials: {p!r} (G2)")
    for url in scan.urls:
        low = url.strip().lower()
        if low.startswith(("mailto:", "tel:", "cid:", "#")):
            continue
        if not low.startswith("https://"):
            raise ValueError(f"Email links/assets must be absolute https: {url!r} (G3)")
        host = urlparse(low).hostname or ""
        if not _host_ok(host) or urlparse(low).username is not None:
            raise ValueError(f"Shortened, numeric-host or credential-bearing URL: {url!r} (G3)")
    for href, text in scan.anchors:
        real = urlparse(href.strip().lower()).hostname or ""
        if not real:
            continue
        for m in _HOSTISH.finditer(text):
            if not _same_site(m.group(1).lower(), real):
                raise ValueError(f"Anchor text {m.group(1)!r} != real link host {real!r} (G3)")

async def send_email(*, to: str, subject: str, html: str, reply_to: Optional[str] = None) -> Optional[str]:
    _assert_safe_email(subject, html)
    payload = {"to": [to], "subject": subject, "html": html, "from_name": EMAIL_FROM_NAME}
    if reply_to or EMAIL_REPLY_TO:
        payload["contact_email"] = reply_to or EMAIL_REPLY_TO
    try:
        async with httpx.AsyncClient(timeout=30) as http_client:
            resp = await http_client.post(f"{EMAIL_BASE_URL}/api/v1/email/send",
                                          headers={"X-Email-Key": EMAIL_KEY}, json=payload)
        resp.raise_for_status()
        return resp.json().get("id")
    except Exception as e:
        logger.error(f"Email send error: {str(e)}")
        raise HTTPException(status_code=502, detail="Invio email fallito")

async def send_digest_email():
    admin_user = {"role": "admin", "can_view_all": True, "store_ids": []}
    alerts = await build_alerts(admin_user)
    total = len(alerts["rinnovi"]) + len(alerts["pagamenti_clienti"]) + len(alerts["pagamenti_negozi"])
    if total == 0:
        logger.info("Digest: nessun alert, email non inviata")
        return None
    app_url = os.environ.get("FRONTEND_URL", "")
    rows_r = "".join(
        f'<tr><td style="padding:6px 10px;border-bottom:1px solid #e2e8f0">{escape(r["cognome"])} {escape(r["nome"])}</td>'
        f'<td style="padding:6px 10px;border-bottom:1px solid #e2e8f0">{escape(r["tipo_bolletta"])}</td>'
        f'<td style="padding:6px 10px;border-bottom:1px solid #e2e8f0">{escape(r["data_rinnovo"])}</td>'
        f'<td style="padding:6px 10px;border-bottom:1px solid #e2e8f0">{r["giorni"]} gg</td></tr>'
        for r in alerts["rinnovi"])
    rows_s = "".join(
        f'<tr><td style="padding:6px 10px;border-bottom:1px solid #e2e8f0">{escape(s["nome"])} ({escape(s["referente"])})</td>'
        f'<td style="padding:6px 10px;border-bottom:1px solid #e2e8f0">Ultimo pagamento: {escape(str(s["last_payment_date"] or "-"))}</td></tr>'
        for s in alerts["pagamenti_negozi"])
    sec_r = (f'<h3 style="color:#0f172a">Rinnovi in scadenza ({len(alerts["rinnovi"])})</h3>'
             f'<table width="100%" style="border-collapse:collapse;font-size:14px">{rows_r}</table>') if alerts["rinnovi"] else ""
    sec_s = (f'<h3 style="color:#0f172a">Negozi da ripagare ({len(alerts["pagamenti_negozi"])})</h3>'
             f'<table width="100%" style="border-collapse:collapse;font-size:14px">{rows_s}</table>') if alerts["pagamenti_negozi"] else ""
    html = (f'<table role="presentation" width="100%"><tr><td style="padding:24px;font-family:Arial,sans-serif">'
            f'<h2 style="color:#0284c7;margin:0 0 12px">Riepilogo alert giornaliero</h2>'
            f'<p>Ci sono <strong>{total}</strong> elementi che richiedono attenzione.</p>'
            f'{sec_r}{sec_s}'
            f'<p style="margin-top:16px"><a href="{escape(app_url)}">Apri il gestionale</a></p>'
            f'<p style="font-size:12px;color:#888">Inviato da {escape(EMAIL_FROM_NAME)}. Non chiediamo mai password o dati di pagamento via email.</p>'
            f'</td></tr></table>')
    try:
        return await send_email(to=os.environ["ALERT_EMAIL"], subject=f"Alert utenze: {total} elementi da gestire", html=html)
    except Exception as e:
        logger.error(f"Digest email non inviata: {getattr(e, 'detail', str(e))}")
        return None

# ---------------- Auth routes ----------------

class LoginInput(BaseModel):
    email: str
    password: str

@api_router.post("/auth/login")
async def login(input: LoginInput):
    email = input.email.strip().lower()
    identifier = f"login:{email}"
    attempts = await db.login_attempts.find_one({"identifier": identifier}, {"_id": 0})
    if attempts and attempts.get("count", 0) >= 5:
        locked_until = attempts.get("locked_until")
        if locked_until and datetime.fromisoformat(locked_until) > datetime.now(timezone.utc):
            raise HTTPException(status_code=429, detail="Troppi tentativi. Riprova tra 15 minuti.")
        await db.login_attempts.delete_one({"identifier": identifier})
    user = await db.users.find_one({"email": email}, {"_id": 0})
    if not user or not verify_password(input.password, user["password_hash"]):
        await db.login_attempts.update_one(
            {"identifier": identifier},
            {"$inc": {"count": 1},
             "$set": {"locked_until": (datetime.now(timezone.utc) + timedelta(minutes=15)).isoformat()}},
            upsert=True)
        raise HTTPException(status_code=401, detail="Email o password non corretti")
    if not user.get("active", True):
        raise HTTPException(status_code=403, detail="Account disattivato")
    await db.login_attempts.delete_one({"identifier": identifier})
    from fastapi.responses import JSONResponse
    if user.get("totp_enabled"):
        return JSONResponse({"mfa_required": True, "mfa_token": create_mfa_token(user["id"])})
    token = create_token(user["id"])
    resp = JSONResponse({"token": token, "user": serialize_user(user)})
    resp.set_cookie("gu_token", token, httponly=True, secure=True, samesite="lax", max_age=86400)
    return resp

class MfaLoginInput(BaseModel):
    mfa_token: str
    code: str

class TotpCodeInput(BaseModel):
    code: str

class TotpDisableInput(BaseModel):
    password: str
    code: str

@api_router.post("/auth/login/mfa")
async def login_mfa(input: MfaLoginInput, request: Request):
    try:
        payload = jwt.decode(input.mfa_token, os.environ["JWT_SECRET"], algorithms=[JWT_ALGORITHM])
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Verifica scaduta, ripeti l'accesso")
    if payload.get("type") != "mfa":
        raise HTTPException(status_code=401, detail="Token non valido")
    identifier = f"mfa:{payload['sub']}"
    attempts = await db.login_attempts.find_one({"identifier": identifier}, {"_id": 0})
    if attempts and attempts.get("count", 0) >= 5 and attempts.get("locked_until") \
            and datetime.fromisoformat(attempts["locked_until"]) > datetime.now(timezone.utc):
        raise HTTPException(status_code=429, detail="Troppi tentativi. Riprova tra 15 minuti.")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0})
    if not user or not user.get("totp_enabled") or not await verify_totp_or_recovery(user, input.code):
        await db.login_attempts.update_one({"identifier": identifier},
                                           {"$inc": {"count": 1},
                                            "$set": {"locked_until": (datetime.now(timezone.utc) + timedelta(minutes=15)).isoformat()}},
                                           upsert=True)
        raise HTTPException(status_code=401, detail="Codice non valido")
    await db.login_attempts.delete_one({"identifier": identifier})
    request.state.mfa_user = user
    from fastapi.responses import JSONResponse
    token = create_token(user["id"])
    resp = JSONResponse({"token": token, "user": serialize_user(user)})
    resp.set_cookie("gu_token", token, httponly=True, secure=True, samesite="lax", max_age=86400)
    return resp

@api_router.post("/auth/2fa/enroll")
async def totp_enroll(user: dict = Depends(get_current_user)):
    if user.get("totp_enabled"):
        raise HTTPException(status_code=409, detail="Autenticazione a due fattori già attiva")
    secret = pyotp.random_base32()
    uri = pyotp.TOTP(secret, interval=30).provisioning_uri(name=user["email"], issuer_name=TOTP_ISSUER)
    await db.users.update_one({"id": user["id"]}, {"$set": {"totp_pending_secret_enc": totp_encrypt(secret)}})
    return {"qr_data_uri": totp_qr_data_uri(uri), "manual_key": secret}

@api_router.post("/auth/2fa/enroll/confirm")
async def totp_enroll_confirm(input: TotpCodeInput, user: dict = Depends(get_current_user)):
    code = input.code.strip()
    if not TOTP_CODE_RE.fullmatch(code):
        raise HTTPException(status_code=400, detail="Il codice deve avere 6 cifre")
    enc = user.get("totp_pending_secret_enc")
    if not enc:
        raise HTTPException(status_code=400, detail="Avvia prima la configurazione")
    totp = pyotp.TOTP(totp_decrypt(enc), interval=30)
    if not totp.verify(code, valid_window=1):
        raise HTTPException(status_code=400, detail="Codice non valido: controlla l'ora del telefono e riprova")
    plain, stored = make_recovery_codes()
    await db.users.update_one({"id": user["id"]}, {
        "$set": {"totp_enabled": True, "totp_secret_enc": enc, "recovery_codes": stored,
                 "totp_last_timecode": totp.timecode(datetime.now(timezone.utc)),
                 "totp_enabled_at": datetime.now(timezone.utc).isoformat()},
        "$unset": {"totp_pending_secret_enc": ""}})
    return {"recovery_codes": plain}

@api_router.post("/auth/2fa/disable")
async def totp_disable(input: TotpDisableInput, user: dict = Depends(get_current_user)):
    if not user.get("totp_enabled"):
        raise HTTPException(status_code=400, detail="2FA non attiva")
    if not verify_password(input.password, user["password_hash"]) or not await verify_totp_or_recovery(user, input.code):
        raise HTTPException(status_code=401, detail="Password o codice non validi")
    await db.users.update_one({"id": user["id"]}, {"$set": {"totp_enabled": False},
                                                   "$unset": {"totp_secret_enc": "", "recovery_codes": "", "totp_last_timecode": ""}})
    return {"status": "ok"}

@api_router.get("/auth/2fa/status")
async def totp_status(user: dict = Depends(get_current_user)):
    codes = user.get("recovery_codes", [])
    return {"enabled": bool(user.get("totp_enabled")), "enabled_at": user.get("totp_enabled_at"),
            "recovery_codes_left": sum(1 for c in codes if c.get("used_at") is None)}

@api_router.post("/users/{user_id}/2fa/reset")
async def admin_reset_2fa(user_id: str, admin: dict = Depends(require_admin)):
    res = await db.users.update_one({"id": user_id}, {"$set": {"totp_enabled": False},
                                                      "$unset": {"totp_secret_enc": "", "recovery_codes": "", "totp_last_timecode": "", "totp_pending_secret_enc": ""}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Utente non trovato")
    return {"status": "ok"}

@api_router.get("/auth/me")
async def auth_me(user: dict = Depends(get_current_user)):
    return serialize_user(user)

@api_router.post("/auth/logout")
async def logout(user: dict = Depends(get_current_user)):
    from fastapi.responses import JSONResponse
    resp = JSONResponse({"status": "ok"})
    resp.delete_cookie("gu_token")
    return resp

# ---------------- Users (admin) ----------------

class UserCreate(BaseModel):
    name: str
    email: str
    password: str
    role: str = "negozio"
    store_ids: List[str] = []
    can_view_all: bool = False
    sections: List[str] = []

class UserUpdate(BaseModel):
    name: Optional[str] = None
    role: Optional[str] = None
    store_ids: Optional[List[str]] = None
    can_view_all: Optional[bool] = None
    active: Optional[bool] = None
    password: Optional[str] = None
    sections: Optional[List[str]] = None

@api_router.get("/users")
async def list_users(admin: dict = Depends(require_admin)):
    users = await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(500)
    return users

@api_router.post("/users")
async def create_user(input: UserCreate, user: dict = Depends(get_current_user)):
    if user["role"] != "admin":
        # Non-admin: solo account negozio, senza vista globale, limitati ai propri negozi
        if input.role != "negozio":
            raise HTTPException(status_code=403, detail="Puoi creare solo utenti di tipo Negozio")
        input.can_view_all = False
        own = user.get("store_ids", [])
        input.store_ids = [s for s in input.store_ids if s in own] or own[:1]
        caller_sections = user.get("sections") or ["energia", "riparazioni", "telefonia"]
        input.sections = [s for s in input.sections if s in caller_sections] or caller_sections
    email = input.email.strip().lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email già registrata")
    if input.role not in ("admin", "operatore", "negozio", "tecnico"):
        raise HTTPException(status_code=400, detail="Ruolo non valido")
    user = {"id": str(uuid.uuid4()), "name": input.name, "email": email,
            "password_hash": hash_password(input.password), "role": input.role,
            "store_ids": input.store_ids, "can_view_all": input.can_view_all,
            "sections": input.sections or ["energia", "riparazioni", "telefonia"],
            "active": True, "created_at": datetime.now(timezone.utc).isoformat()}
    await db.users.insert_one(user)
    return serialize_user(user)

@api_router.patch("/users/{user_id}")
async def update_user(user_id: str, input: UserUpdate, admin: dict = Depends(require_admin)):
    updates = {k: v for k, v in input.model_dump().items() if v is not None}
    if "password" in updates:
        updates["password_hash"] = hash_password(updates.pop("password"))
    if not updates:
        raise HTTPException(status_code=400, detail="Nessuna modifica")
    res = await db.users.update_one({"id": user_id}, {"$set": updates})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Utente non trovato")
    user = await db.users.find_one({"id": user_id}, {"_id": 0})
    return serialize_user(user)

@api_router.delete("/users/{user_id}")
async def delete_user(user_id: str, admin: dict = Depends(require_admin)):
    if user_id == admin["id"]:
        raise HTTPException(status_code=400, detail="Non puoi eliminare il tuo account")
    res = await db.users.delete_one({"id": user_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Utente non trovato")
    return {"status": "ok"}

# ---------------- Stores (venditori/negozi) ----------------

class StoreCreate(BaseModel):
    nome: str
    referente: str = ""
    tipo: str = "negozio"
    note: str = ""
    review_link: str = ""
    telefono_avvisi: str = ""
    msg_privacy: str = ""
    msg_pronto: str = ""
    msg_recensione: str = ""
    msg_promemoria: str = ""
    msg_vincolo: str = ""
    msg_offerta_annuale: str = ""
    msg_rinnovo_energia: str = ""
    msg_truffe: str = ""

class StoreUpdate(BaseModel):
    nome: Optional[str] = None
    referente: Optional[str] = None
    tipo: Optional[str] = None
    note: Optional[str] = None
    review_link: Optional[str] = None
    telefono_avvisi: Optional[str] = None
    msg_privacy: Optional[str] = None
    msg_pronto: Optional[str] = None
    msg_recensione: Optional[str] = None
    msg_promemoria: Optional[str] = None
    msg_vincolo: Optional[str] = None
    msg_offerta_annuale: Optional[str] = None
    msg_rinnovo_energia: Optional[str] = None
    msg_truffe: Optional[str] = None

MSG_DEFAULTS = {
    "msg_privacy": """RS Group – Grazie per averci scelto!
Ciao {nome}
per procedere con la tua richiesta e completare l'attivazione del servizio, è necessario firmare l'autorizzazione privacy.
Puoi farlo in modo semplice e veloce al link qui sotto:
https://rsriparazioni.it/privacy/
La firma è richiesta per attivare correttamente il servizio (riparazioni, contratti energia o telefonia).
Grazie per la fiducia
RS Group""",
    "msg_pronto": ("Ciao {nome}!\nIl tuo {dispositivo} (riparazione N. {numero}) è pronto per il ritiro presso RS Riparazioni {negozio}.\n"
                   "Ti aspettiamo in negozio negli orari di apertura. Grazie!"),
    "msg_recensione": ("Ciao!\nGrazie per aver scelto RS Riparazioni {negozio}\n"
                       "Se ti sei trovato bene, ci farebbe davvero piacere una tua recensione\n"
                       "Basta un clic qui\n{link}\nGrazie per il supporto"),
    "msg_promemoria": ("Ciao {nome}, ti ricordiamo che il tuo {dispositivo} (riparazione N. {numero}) è pronto da alcuni giorni "
                       "presso RS Riparazioni {negozio}.\nPassa a ritirarlo quando vuoi negli orari di apertura. Grazie!"),
    "msg_vincolo": ("Ciao {nome} {cognome}, il vincolo sulla tua offerta sta per scadere: passa in negozio per valutare il tuo prossimo risparmio."),
    "msg_offerta_annuale": ("Ciao {nome} {cognome}, la tua offerta sta per scadere: contattaci o passa in negozio per il tuo prossimo risparmio, "
                            "abbiamo offerte dedicate per te."),
    "msg_rinnovo_energia": """Ciao {nome}
Ti informiamo che il tuo contratto utenze è in prossima scadenza. Nei prossimi giorni verrai contattato dal numero 353 377 6535: parlerai con Deborah, che ti seguirà nell'aggiornamento della tua fornitura e nella verifica delle migliori opportunità di risparmio in base ai tuoi consumi.
Per effettuare un'analisi ancora più precisa, puoi inviarci direttamente qui le tue ultime bollette.
Come sempre, valuteremo le tue abitudini di consumo per garantirti la soluzione più conveniente e il miglior prezzo disponibile sul mercato.
A presto""",
    "msg_truffe": """Ciao {nome},
le truffe telefoniche sono sempre più frequenti. Se ricevi una chiamata da chi si presenta come operatore di luce, gas, telefonia, banca o altri servizi...
*Non prendere decisioni di fretta.*
*Non comunicare codici, dati personali o bancari.*
*Non dire "SÌ" senza aver capito con chi stai parlando.*
*Hai già i tuoi consulenti di fiducia.*
Prima di firmare, confermare o accettare qualsiasi proposta, *contatta noi*. Ti diremo gratuitamente se la chiamata è affidabile oppure se potrebbe trattarsi di un tentativo di truffa o di una proposta poco conveniente.
Un messaggio o una telefonata possono evitarti problemi e costi inutili.
*CambiaOra*
_Siamo al tuo fianco per aiutarti a scegliere in sicurezza._""",
}
MSG_PLACEHOLDERS = ["{nome}", "{cognome}", "{dispositivo}", "{numero}", "{negozio}", "{link}"]

def store_msg(store: Optional[dict], key: str, **vals) -> str:
    tpl = (store or {}).get(key) or MSG_DEFAULTS[key]
    vals = {"nome": "", "cognome": "", "dispositivo": "", "numero": "", "negozio": (store or {}).get("nome", ""),
            "link": (store or {}).get("review_link", ""), **vals}
    try:
        return tpl.format(**vals)
    except (KeyError, IndexError, ValueError):
        return tpl

@api_router.get("/messaggi-default")
async def messaggi_default(user: dict = Depends(get_current_user)):
    return {"defaults": MSG_DEFAULTS, "placeholders": MSG_PLACEHOLDERS}

@api_router.get("/stores")
async def list_stores(user: dict = Depends(get_current_user)):
    stores = await db.stores.find({}, {"_id": 0}).to_list(500)
    if not (user["role"] == "admin" or user.get("can_view_all")):
        ids = set(user.get("store_ids", []))
        stores = [s for s in stores if s["id"] in ids]
    month_start = date.today().replace(day=1).isoformat()
    for s in stores:
        s["pagato_effettivo"] = effective_pagato(s.get("pagato", False), s.get("last_payment_date"))
        s["totale_clienti"] = await db.clients.count_documents({"venditore_id": s["id"]})
        s["contratti_mese"] = await db.clients.count_documents(
            {"venditore_id": s["id"], "data_contratto": {"$gte": month_start}})
        # Clienti incassati dalla struttura (pagato effettivo, reset 6 mesi): base per la divisione compensi
        s["clienti_incassati"] = 0
        async for c in db.clients.find({"venditore_id": s["id"], "pagato": True},
                                       {"_id": 0, "pagato": 1, "last_payment_date": 1}):
            if effective_pagato(c.get("pagato", False), c.get("last_payment_date")):
                s["clienti_incassati"] += 1
    return stores

@api_router.post("/stores")
async def create_store(input: StoreCreate, admin: dict = Depends(require_admin)):
    store = {"id": str(uuid.uuid4()), "nome": input.nome, "referente": input.referente,
             "tipo": input.tipo, "note": input.note, "review_link": input.review_link,
             "telefono_avvisi": input.telefono_avvisi, "msg_privacy": input.msg_privacy, "msg_pronto": input.msg_pronto,
             "msg_recensione": input.msg_recensione, "msg_promemoria": input.msg_promemoria, "pagato": False, "last_payment_date": None,
             "created_at": datetime.now(timezone.utc).isoformat()}
    await db.stores.insert_one(store)
    store.pop("_id", None)
    return store

@api_router.patch("/stores/{store_id}")
async def update_store(store_id: str, input: StoreUpdate, admin: dict = Depends(require_admin)):
    updates = {k: v for k, v in input.model_dump().items() if v is not None}
    res = await db.stores.update_one({"id": store_id}, {"$set": updates})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Negozio non trovato")
    return await db.stores.find_one({"id": store_id}, {"_id": 0})

@api_router.post("/stores/{store_id}/mark-paid")
async def mark_store_paid(store_id: str, admin: dict = Depends(require_admin)):
    res = await db.stores.update_one({"id": store_id},
        {"$set": {"pagato": True, "last_payment_date": date.today().isoformat()}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Negozio non trovato")
    return {"status": "ok", "last_payment_date": date.today().isoformat()}

@api_router.delete("/stores/{store_id}")
async def delete_store(store_id: str, admin: dict = Depends(require_admin)):
    n_clients = await db.clients.count_documents({"venditore_id": store_id})
    if n_clients:
        raise HTTPException(status_code=400, detail=f"Il negozio ha {n_clients} clienti collegati: spostali prima di eliminarlo")
    res = await db.stores.delete_one({"id": store_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Negozio non trovato")
    await db.users.update_many({}, {"$pull": {"store_ids": store_id}})
    await db.venditori.update_many({"store_id": store_id}, {"$set": {"store_id": ""}})
    return {"status": "ok"}

# ---------------- Clients (utenze) ----------------

TIPI_BUSINESS = ["ditta_individuale", "societa", "business"]

def valida_tipo_cliente(data: dict) -> None:
    tipo = data.get("tipo_cliente") or "privato"
    piva = (data.get("p_iva") or "").strip()
    cf = (data.get("codice_fiscale") or "").strip()
    if tipo in TIPI_BUSINESS and not piva:
        raise HTTPException(status_code=400, detail="Partita IVA obbligatoria per i clienti business")
    if tipo == "ditta_individuale" and not cf:
        raise HTTPException(status_code=400, detail="Codice fiscale obbligatorio per la ditta individuale")
    if tipo == "privato":
        data["p_iva"] = ""

class ClientInput(BaseModel):
    nome: str
    cognome: str
    tipo_cliente: str = "privato"
    codice_fiscale: str = ""
    p_iva: str = ""
    indirizzo: str = ""
    pod: str = ""
    pdr: str = ""
    iban: str = ""
    email: str = ""
    telefono: str = ""
    kw_potenza: Optional[float] = None
    tipo_bolletta: str = "luce"
    fornitore_provenienza: str = ""
    costo_kwh_attuale: Optional[float] = None
    spese_fisse_attuale: Optional[float] = None
    costo_smc_attuale: Optional[float] = None
    data_contratto: Optional[str] = None
    data_verifica: Optional[str] = None
    data_cambio: Optional[str] = None
    tipo_contratto: str = "fisso"
    nuovo_fornitore: str = ""
    costo_kwh_nuovo: Optional[float] = None
    spese_fisse_nuovo: Optional[float] = None
    costo_smc_nuovo: Optional[float] = None
    privacy_firmata: bool = False
    note: str = ""
    lavorazione: str = "da_quotare"
    origine: str = ""
    venditore_id: str = ""
    operatore_id: str = ""
    provincia: str = ""
    venditore_pagato: bool = False
    gestione: str = "cambiaora"

@api_router.get("/clients")
async def list_clients(user: dict = Depends(get_current_user),
                       lavorazione: str = "", tipo_bolletta: str = "", tipo_servizio: str = "",
                       venditore_id: str = "", q: str = "", no_recensioni: str = "", tipo_cliente: str = ""):
    scope = client_scope_filter(user)
    if no_recensioni == "1":
        scope["no_recensioni"] = True
    if lavorazione:
        scope["lavorazione"] = lavorazione
    if tipo_cliente == "privato":
        scope["tipo_cliente"] = {"$in": ["privato", None, ""]}
    elif tipo_cliente == "business":
        scope["tipo_cliente"] = {"$in": TIPI_BUSINESS}
    elif tipo_cliente:
        scope["tipo_cliente"] = tipo_cliente
    if tipo_bolletta:
        scope["tipo_bolletta"] = tipo_bolletta
    if venditore_id and (user["role"] == "admin" or user.get("can_view_all")):
        scope["venditore_id"] = venditore_id
    if q:
        rx = {"$regex": re.escape(q), "$options": "i"}
        scope["$or"] = [{"nome": rx}, {"cognome": rx}, {"codice_fiscale": rx},
                        {"pod": rx}, {"pdr": rx}, {"telefono": rx}, {"email": rx}]
    proj = {"_id": 0, "id": 1, "nome": 1, "cognome": 1, "tipo_cliente": 1, "p_iva": 1,
            "telefono": 1, "email": 1, "codice_fiscale": 1, "pod": 1, "pdr": 1,
            "tipo_bolletta": 1, "fornitore_provenienza": 1, "lavorazione": 1,
            "data_contratto": 1, "venditore_id": 1, "pagato": 1, "last_payment_date": 1,
            "privacy_firmata": 1, "created_at": 1, "no_recensioni": 1}
    clients = await db.clients.find(scope, proj).sort("created_at", -1).to_list(5000)
    tipi_per_client = {}
    async for row in db.servizi.aggregate([{"$group": {"_id": "$client_id", "tipi": {"$addToSet": "$tipo"}}}]):
        tipi_per_client[row["_id"]] = row["tipi"]
    out = []
    for c in clients:
        compute_dates(c)
        tipi = set(tipi_per_client.get(c["id"], []))
        cats = set()
        if c.get("data_contratto"):
            cats.add("energia")
        if tipi & {"sim", "internet", "fisso"}:
            cats.add("telefonia")
        if tipi & {"riparazione", "accessori", "vendita"}:
            cats.add("riparazioni")
        c["premium_step"] = len(cats)
        tags = []
        if "energia" in cats or c.get("pod") or c.get("pdr") or c.get("fornitore_provenienza"):
            tags.append("gas" if c.get("tipo_bolletta") == "gas" else "luce")
        if tipi & {"riparazione", "accessori", "vendita"}:
            tags.append("rip")
        if "sim" in tipi:
            tags.append("mob")
        if tipi & {"internet", "fisso"}:
            tags.append("fis")
        c["tipi_servizi"] = tags
        if tipo_servizio and tipo_servizio not in tags:
            continue
        out.append(c)
    return out

@api_router.post("/clients")
async def create_client(input: ClientInput, background_tasks: BackgroundTasks, user: dict = Depends(get_current_user)):
    data = input.model_dump()
    valida_tipo_cliente(data)
    if user["role"] == "negozio" and data["venditore_id"] not in user.get("store_ids", []):
        if user.get("store_ids"):
            data["venditore_id"] = user["store_ids"][0]
    data.update({"id": str(uuid.uuid4()), "pagato": False, "last_payment_date": None,
                 "created_by": user["id"], "created_at": datetime.now(timezone.utc).isoformat(),
                 "updated_at": datetime.now(timezone.utc).isoformat()})
    await db.clients.insert_one(data)
    await log_lavorazione(data, user, data["lavorazione"], "Cliente inserito")
    data.pop("_id", None)
    if data.get("telefono"):
        background_tasks.add_task(privacy_automatica_nuovo_cliente, data["id"], data.get("origine", "") != "riparazione")
    return compute_dates(data)

async def privacy_automatica_nuovo_cliente(client_id: str, queue_review: bool) -> None:
    c = await db.clients.find_one({"id": client_id}, {"_id": 0})
    if not c or c.get("privacy_msg_sent_at"):
        return
    try:
        await invia_privacy_cliente(c, queue_review=queue_review)
    except HTTPException as e:
        logger.warning(f"Privacy automatica non inviata per {client_id}: {e.detail}")

@api_router.get("/clients/{client_id}")
async def get_client(client_id: str, user: dict = Depends(get_current_user)):
    scope = client_scope_filter(user)
    scope["id"] = client_id
    c = await db.clients.find_one(scope, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    history = await db.lavorazioni_log.find({"client_id": client_id}, {"_id": 0}).sort("created_at", -1).to_list(100)
    c["history"] = history
    tipi = {s["tipo"] async for s in db.servizi.find({"client_id": client_id}, {"_id": 0, "tipo": 1})}
    cats = set()
    if c.get("data_contratto"):
        cats.add("energia")
    if tipi & {"sim", "internet", "fisso"}:
        cats.add("telefonia")
    if tipi & {"riparazione", "accessori", "vendita"}:
        cats.add("riparazioni")
    c["premium_step"] = len(cats)
    return compute_dates(c)

@api_router.patch("/clients/{client_id}")
async def update_client(client_id: str, input: ClientInput, user: dict = Depends(get_current_user)):
    scope = client_scope_filter(user)
    scope["id"] = client_id
    old = await db.clients.find_one(scope, {"_id": 0})
    if not old:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    data = input.model_dump(exclude_unset=True)
    if "tipo_cliente" in data:
        valida_tipo_cliente({**old, **data})
        if (data.get("tipo_cliente") or "privato") == "privato":
            data["p_iva"] = ""
    if user["role"] == "negozio":
        data.pop("venditore_id", None)
    data["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.clients.update_one({"id": client_id}, {"$set": data})
    if data.get("lavorazione") and data["lavorazione"] != old.get("lavorazione"):
        await log_lavorazione({**old, **data}, user, data["lavorazione"], "Cambio stato lavorazione")
    updated = await db.clients.find_one({"id": client_id}, {"_id": 0})
    return compute_dates(updated)

@api_router.delete("/clients/{client_id}")
async def delete_client(client_id: str, user: dict = Depends(get_current_user)):
    scope = client_scope_filter(user)
    scope["id"] = client_id
    res = await db.clients.delete_one(scope)
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    await db.lavorazioni_log.delete_many({"client_id": client_id})
    return {"status": "ok"}

@api_router.post("/clients/{client_id}/mark-paid")
async def mark_client_paid(client_id: str, user: dict = Depends(get_current_user)):
    scope = client_scope_filter(user)
    scope["id"] = client_id
    res = await db.clients.update_one(scope,
        {"$set": {"pagato": True, "last_payment_date": date.today().isoformat()}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    return {"status": "ok", "last_payment_date": date.today().isoformat()}

async def log_lavorazione(client_doc: dict, user: dict, status: str, note: str):
    await db.lavorazioni_log.insert_one({
        "id": str(uuid.uuid4()), "client_id": client_doc["id"],
        "client_name": f"{client_doc.get('cognome', '')} {client_doc.get('nome', '')}".strip(),
        "operatore_id": client_doc.get("operatore_id") or user["id"],
        "operatore_name": user["name"], "status": status, "note": note,
        "created_at": datetime.now(timezone.utc).isoformat()})

# ---------------- Venditori ----------------

class VenditoreInput(BaseModel):
    nome: str
    store_id: str = ""
    tipo: str = "interno"
    attivo: bool = True

@api_router.get("/venditori")
async def list_venditori(user: dict = Depends(get_current_user)):
    venditori = await db.venditori.find({}, {"_id": 0}).sort("nome", 1).to_list(200)
    store_ids = [v["store_id"] for v in venditori if v.get("store_id")]
    stores = await db.stores.find({"id": {"$in": store_ids}}, {"_id": 0, "id": 1, "nome": 1}).to_list(100)
    smap = {s["id"]: s["nome"] for s in stores}
    for v in venditori:
        v["store_name"] = smap.get(v.get("store_id", ""), "Esterno")
        v["vendite"] = await db.clients.count_documents({"operatore_id": v["id"]})
        v["da_pagare"] = await db.clients.count_documents({"operatore_id": v["id"], "venditore_pagato": {"$ne": True}})
    return venditori

@api_router.post("/venditori")
async def create_venditore(input: VenditoreInput, admin: dict = Depends(require_admin)):
    data = input.model_dump()
    data.update({"id": str(uuid.uuid4()), "created_at": datetime.now(timezone.utc).isoformat()})
    await db.venditori.insert_one(data)
    data.pop("_id", None)
    return data

@api_router.patch("/venditori/{venditore_id}")
async def update_venditore(venditore_id: str, input: VenditoreInput, admin: dict = Depends(require_admin)):
    res = await db.venditori.update_one({"id": venditore_id}, {"$set": input.model_dump(exclude_unset=True)})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Venditore non trovato")
    return await db.venditori.find_one({"id": venditore_id}, {"_id": 0})

@api_router.delete("/venditori/{venditore_id}")
async def delete_venditore(venditore_id: str, admin: dict = Depends(require_admin)):
    res = await db.venditori.delete_one({"id": venditore_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Venditore non trovato")
    return {"status": "ok"}

@api_router.get("/venditori/{venditore_id}/vendite")
async def vendite_venditore(venditore_id: str, user: dict = Depends(get_current_user)):
    if user["role"] not in ("admin", "operatore") and not user.get("can_view_all"):
        raise HTTPException(status_code=403, detail="Non autorizzato")
    clients = await db.clients.find({"operatore_id": venditore_id},
        {"_id": 0, "id": 1, "nome": 1, "cognome": 1, "data_contratto": 1, "lavorazione": 1,
         "venditore_pagato": 1, "venditore_id": 1, "created_at": 1,
         "pagato": 1, "last_payment_date": 1}).sort("created_at", -1).to_list(2000)
    store_ids = list({c.get("venditore_id") for c in clients if c.get("venditore_id")})
    stores = await db.stores.find({"id": {"$in": store_ids}}, {"_id": 0, "id": 1, "nome": 1}).to_list(100)
    smap = {s["id"]: s["nome"] for s in stores}
    for c in clients:
        c["store_name"] = smap.get(c.get("venditore_id", ""), "-")
        c["incassato_struttura"] = effective_pagato(c.get("pagato", False), c.get("last_payment_date"))
        c.pop("venditore_id", None)
        c.pop("last_payment_date", None)
    return clients

@api_router.post("/clients/{client_id}/venditore-pagato")
async def toggle_venditore_pagato(client_id: str, user: dict = Depends(get_current_user)):
    scope = client_scope_filter(user)
    scope["id"] = client_id
    c = await db.clients.find_one(scope, {"_id": 0, "venditore_pagato": 1})
    if not c:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    new_val = not c.get("venditore_pagato", False)
    await db.clients.update_one({"id": client_id}, {"$set": {"venditore_pagato": new_val}})
    return {"venditore_pagato": new_val}


# ---------------- Meta, Dashboard, Alerts, Operatori ----------------

@api_router.get("/meta")
async def get_meta(user: dict = Depends(get_current_user)):
    stores = await db.stores.find({}, {"_id": 0}).to_list(500)
    if not (user["role"] == "admin" or user.get("can_view_all")):
        ids = set(user.get("store_ids", []))
        stores = [s for s in stores if s["id"] in ids]
    operators = await db.users.find({"role": {"$in": ["operatore", "admin"]}, "active": True},
                                    {"_id": 0, "password_hash": 0}).to_list(200)
    return {"suppliers": SUPPLIERS, "lavorazioni": LAVORAZIONI,
            "tel_operators": {"sim": SIM_OPERATORS, "internet": INTERNET_OPERATORS, "fisso": FISSO_OPERATORS},
            "rip_stati": RIP_STATI, "servizio_tipi": SERVIZIO_TIPI,
            "stores": [{"id": s["id"], "nome": s["nome"], "referente": s.get("referente", "")} for s in stores],
            "operators": [{"id": o["id"], "name": o["name"]} for o in operators]}

@api_router.get("/dashboard/stats")
async def dashboard_stats(user: dict = Depends(get_current_user)):
    scope = client_scope_filter(user)
    clients = await db.clients.find(scope, {"_id": 0}).to_list(5000)
    month_start = date.today().replace(day=1).isoformat()
    stats = {"totale_clienti": len(clients), "luce": 0, "gas": 0,
             "contratti_mese": 0, "rinnovi_30gg": 0, "non_pagati": 0,
             "per_lavorazione": {}, "per_negozio_mese": {}}
    for c in clients:
        compute_dates(c)
        if c.get("tipo_bolletta") == "gas":
            stats["gas"] += 1
        else:
            stats["luce"] += 1
        if (c.get("data_contratto") or "") >= month_start:
            stats["contratti_mese"] += 1
            vid = c.get("venditore_id", "")
            stats["per_negozio_mese"][vid] = stats["per_negozio_mese"].get(vid, 0) + 1
        gr = c.get("giorni_al_rinnovo")
        if gr is not None and 0 <= gr <= 30 and c.get("lavorazione") != "rinnovato":
            stats["rinnovi_30gg"] += 1
        if not c.get("pagato_effettivo"):
            stats["non_pagati"] += 1
        lav = c.get("lavorazione", "da_quotare")
        stats["per_lavorazione"][lav] = stats["per_lavorazione"].get(lav, 0) + 1
    svc_scope = servizio_scope(user)
    stats["servizi_attivi"] = await db.servizi.count_documents(
        {**svc_scope, "stato": {"$nin": ["consegnato", "non_riparabile"]}})
    vinc = await db.servizi.find({**svc_scope, "vincolo_mesi": {"$nin": [None, 0]},
                                  "data_attivazione": {"$ne": None}},
                                 {"_id": 0, "data_attivazione": 1, "vincolo_mesi": 1}).to_list(5000)
    stats["vincoli_60gg"] = 0
    for v in vinc:
        try:
            scad = date.fromisoformat(str(v["data_attivazione"])[:10]) + relativedelta(months=int(v["vincolo_mesi"]))
            if 0 <= (scad - date.today()).days <= 60:
                stats["vincoli_60gg"] += 1
        except (ValueError, TypeError):
            pass
    stats["tempi_riparazione"] = await tempi_riparazione_per_negozio(svc_scope)
    return stats

async def tempi_riparazione_per_negozio(svc_scope: dict) -> list:
    allowed = svc_scope.get("tipo")
    if isinstance(allowed, dict) and "riparazione" not in allowed.get("$in", []):
        return []
    if isinstance(allowed, str) and allowed != "riparazione":
        return []
    rips = await db.servizi.find({**svc_scope, "tipo": "riparazione"},
                                 {"_id": 0, "venditore_id": 1, "stato": 1, "data_ingresso": 1,
                                  "data_uscita": 1, "created_at": 1}).to_list(10000)
    stores = {s["id"]: s["nome"] for s in await db.stores.find({}, {"_id": 0, "id": 1, "nome": 1}).to_list(200)}
    acc: dict = {}
    today = date.today()
    for r in rips:
        vid = r.get("venditore_id", "")
        a = acc.setdefault(vid, {"store_id": vid, "store_name": stores.get(vid, "-"), "chiuse": 0,
                                 "giorni_totali": 0, "aperte": 0, "aperte_oltre_7gg": 0})
        try:
            ingresso = date.fromisoformat(str(r.get("data_ingresso") or r.get("created_at"))[:10])
        except (ValueError, TypeError):
            continue
        if r.get("stato") in ("consegnato", "non_riparabile"):
            if r.get("data_uscita"):
                a["chiuse"] += 1
                a["giorni_totali"] += max((date.fromisoformat(str(r["data_uscita"])[:10]) - ingresso).days, 0)
        else:
            a["aperte"] += 1
            if (today - ingresso).days > 7:
                a["aperte_oltre_7gg"] += 1
    out = []
    for a in acc.values():
        a["media_giorni"] = round(a["giorni_totali"] / a["chiuse"], 1) if a["chiuse"] else None
        out.append(a)
    return sorted(out, key=lambda x: x["store_name"])

@api_router.get("/alerts")
async def get_alerts(user: dict = Depends(get_current_user)):
    return await build_alerts(user)

# ---------------- Follow-up contatti (esito chiamate su scadenze/rinnovi) ----------------
FOLLOWUP_ESITI = {"non_risponde", "richiamare", "contattato", "non_interessato", "chiuso"}

class FollowupInput(BaseModel):
    target_type: str  # client | servizio | store
    target_id: str
    motivo: str = "altro"  # rinnovo | vincolo | riparazione_pronta | pagamento | altro
    esito: str
    richiamare_il: Optional[str] = None
    note: str = ""

async def _followup_target_ok(user: dict, target_type: str, target_id: str) -> Optional[dict]:
    if target_type == "client":
        c = await db.clients.find_one({**client_scope_filter(user), "id": target_id}, {"_id": 0, "id": 1, "nome": 1, "cognome": 1, "venditore_id": 1, "telefono": 1})
        return {"label": f"{c.get('cognome', '')} {c.get('nome', '')}".strip(), "store_id": c.get("venditore_id", ""), "client_id": c["id"], "telefono": c.get("telefono", "")} if c else None
    if target_type == "servizio":
        s = await db.servizi.find_one({**servizio_scope(user), "id": target_id}, {"_id": 0, "id": 1, "client_id": 1, "venditore_id": 1, "dispositivo": 1, "tipo": 1})
        if not s:
            return None
        c = await db.clients.find_one({"id": s["client_id"]}, {"_id": 0, "nome": 1, "cognome": 1, "telefono": 1}) or {}
        return {"label": f"{c.get('cognome', '')} {c.get('nome', '')}".strip(), "store_id": s.get("venditore_id", ""), "client_id": s["client_id"],
                "telefono": c.get("telefono", ""), "dettaglio": s.get("dispositivo") or s.get("tipo", "")}
    if target_type == "store" and user["role"] == "admin":
        s = await db.stores.find_one({"id": target_id}, {"_id": 0, "nome": 1})
        return {"label": s["nome"], "store_id": target_id, "client_id": "", "telefono": ""} if s else None
    return None

@api_router.post("/followup")
async def create_followup(input: FollowupInput, user: dict = Depends(get_current_user)):
    if input.esito not in FOLLOWUP_ESITI:
        raise HTTPException(status_code=400, detail="Esito non valido")
    if input.esito in ("non_risponde", "richiamare") and not input.richiamare_il:
        raise HTTPException(status_code=400, detail="Indica quando richiamare")
    target = await _followup_target_ok(user, input.target_type, input.target_id)
    if not target:
        raise HTTPException(status_code=404, detail="Elemento non trovato")
    now = datetime.now(timezone.utc).isoformat()
    doc = {"id": str(uuid.uuid4()), **input.model_dump(), **target, "user_id": user["id"], "user_name": user["name"], "created_at": now,
           "aperto": input.esito in ("non_risponde", "richiamare")}
    await db.followup.update_many({"target_type": input.target_type, "target_id": input.target_id, "motivo": input.motivo, "aperto": True},
                                  {"$set": {"aperto": False, "chiuso_at": now}})
    await db.followup.insert_one(doc)
    doc.pop("_id", None)
    return doc

def _followup_scope(user: dict) -> dict:
    if user["role"] == "admin" or user.get("can_view_all"):
        return {}
    return {"store_id": {"$in": user.get("store_ids", [])}}

@api_router.get("/followup")
async def list_followup(target_type: str, target_id: str, user: dict = Depends(get_current_user)):
    if not await _followup_target_ok(user, target_type, target_id):
        raise HTTPException(status_code=404, detail="Elemento non trovato")
    return await db.followup.find({"target_type": target_type, "target_id": target_id}, {"_id": 0}).sort("created_at", -1).to_list(200)

@api_router.get("/followup/da-richiamare")
async def followup_da_richiamare(user: dict = Depends(get_current_user)):
    rows = await db.followup.find({**_followup_scope(user), "aperto": True}, {"_id": 0}).sort("richiamare_il", 1).to_list(1000)
    today = date.today().isoformat()
    for r in rows:
        r["scaduto"] = bool(r.get("richiamare_il")) and r["richiamare_il"] < today
        r["oggi"] = r.get("richiamare_il") == today
    return rows

async def attach_followups(items: list, target_type: str, id_key: str, motivo: str):
    """Aggiunge a ogni voce l'ultimo esito contatto (ultimo_contatto) per lo stesso motivo."""
    ids = [i[id_key] for i in items]
    if not ids:
        return items
    last: dict = {}
    async for f in db.followup.find({"target_type": target_type, "target_id": {"$in": ids}, "motivo": motivo}, {"_id": 0}).sort("created_at", -1):
        last.setdefault(f["target_id"], {k: f.get(k) for k in ("esito", "richiamare_il", "note", "user_name", "created_at", "aperto")})
    for i in items:
        i["ultimo_contatto"] = last.get(i[id_key])
    return items

@api_router.get("/scadenze-settimana")
async def scadenze_settimana(user: dict = Depends(get_current_user)):
    sections = user_sections(user)
    out = {"rinnovi": [], "vincoli": [], "riparazioni_pronte": []}
    if "energia" in sections:
        scope = client_scope_filter(user)
        projection = {"_id": 0, "id": 1, "nome": 1, "cognome": 1, "data_contratto": 1,
                      "lavorazione": 1, "tipo_bolletta": 1, "pagato": 1, "last_payment_date": 1}
        clients = await db.clients.find(scope, projection).to_list(5000)
        for c in clients:
            compute_dates(c)
            gr = c.get("giorni_al_rinnovo")
            if gr is not None and 0 <= gr <= 7 and c.get("lavorazione") != "rinnovato":
                out["rinnovi"].append({"client_id": c["id"], "nome": c.get("nome", ""), "cognome": c.get("cognome", ""),
                                       "data_rinnovo": c["data_rinnovo"], "giorni": gr,
                                       "tipo_bolletta": c.get("tipo_bolletta", "")})
        out["rinnovi"].sort(key=lambda x: x["giorni"])
    if "telefonia" in sections:
        svc_scope = servizio_scope(user)
        svc_scope.update({"vincolo_mesi": {"$nin": [None, 0]}, "data_attivazione": {"$ne": None}})
        vinc = await db.servizi.find(svc_scope, {"_id": 0}).to_list(5000)
        names = await clients_name_map(list({s["client_id"] for s in vinc}))
        for s in vinc:
            ser = serialize_servizio(s, names.get(s["client_id"], ""))
            g = ser.get("giorni_alla_scadenza")
            if g is not None and 0 <= g <= 7:
                out["vincoli"].append({"id": s["id"], "client_name": ser["client_name"], "tipo": s["tipo"],
                                       "operatore_tel": s.get("operatore_tel", ""), "numero": s.get("numero", ""),
                                       "scadenza_vincolo": ser["scadenza_vincolo"], "giorni": g})
        out["vincoli"].sort(key=lambda x: x["giorni"])
    if "riparazioni" in sections:
        svc_scope = servizio_scope(user)
        svc_scope["tipo"] = "riparazione"
        svc_scope["stato"] = "pronto"
        rips = await db.servizi.find(svc_scope, {"_id": 0, "id": 1, "client_id": 1, "dispositivo": 1,
                                                 "problema": 1, "updated_at": 1}).to_list(1000)
        names = await clients_name_map(list({s["client_id"] for s in rips}))
        out["riparazioni_pronte"] = [
            {"id": s["id"], "client_name": names.get(s["client_id"], ""),
             "dispositivo": s.get("dispositivo", ""), "problema": s.get("problema", "")}
            for s in rips]
    out["totale"] = len(out["rinnovi"]) + len(out["vincoli"]) + len(out["riparazioni_pronte"])
    await attach_followups(out["rinnovi"], "client", "client_id", "rinnovo")
    await attach_followups(out["vincoli"], "servizio", "id", "vincolo")
    await attach_followups(out["riparazioni_pronte"], "servizio", "id", "riparazione_pronta")
    return out


@api_router.post("/alerts/send-digest")
async def send_digest_now(admin: dict = Depends(require_admin)):
    email_id = await send_digest_email()
    if email_id is None:
        return {"status": "ok", "message": "Nessun alert attivo, email non inviata"}
    return {"status": "ok", "email_id": email_id}

@api_router.get("/operators/stats")
async def operators_stats(user: dict = Depends(get_current_user)):
    if user["role"] == "negozio":
        raise HTTPException(status_code=403, detail="Accesso non consentito")
    operators = await db.users.find({"role": {"$in": ["operatore", "admin"]}},
                                    {"_id": 0, "password_hash": 0}).to_list(200)
    month_start = datetime.now(timezone.utc).replace(day=1, hour=0, minute=0, second=0).isoformat()
    result = []
    for op in operators:
        logs = await db.lavorazioni_log.find(
            {"operatore_id": op["id"], "created_at": {"$gte": month_start}}, {"_id": 0}).to_list(5000)
        per_status = {}
        for l in logs:
            per_status[l["status"]] = per_status.get(l["status"], 0) + 1
        tot_clients = await db.clients.count_documents({"operatore_id": op["id"]})
        result.append({"id": op["id"], "name": op["name"], "role": op["role"],
                       "lavorazioni_mese": len(logs), "per_status": per_status,
                       "clienti_gestiti": tot_clients,
                       "chiusi_mese": per_status.get("cambio_effettuato", 0) + per_status.get("rinnovato", 0)})
    return result

# ---------------- Import Google Sheet ----------------

HEADER_MAP = {
    "nome": "nome", "cognome": "cognome", "ragione sociale": "cognome",
    "codice fiscale": "codice_fiscale", "cf": "codice_fiscale",
    "p iva": "p_iva", "piva": "p_iva", "partita iva": "p_iva",
    "indirizzo": "indirizzo", "pod": "pod", "pdr": "pdr", "iban": "iban",
    "mail": "email", "email": "email", "e mail": "email",
    "telefono": "telefono", "tel": "telefono", "cellulare": "telefono",
    "kw": "kw_potenza", "potenza": "kw_potenza", "potenza kw": "kw_potenza", "potenza luce": "kw_potenza",
    "tipo bolletta": "tipo_bolletta", "tipo": "tipo_bolletta", "bolletta": "tipo_bolletta",
    "fornitore": "fornitore_provenienza", "fornitore di provenienza": "fornitore_provenienza",
    "fornitore attuale": "fornitore_provenienza",
    "costo al kw": "costo_kwh_attuale", "costo kwh": "costo_kwh_attuale",
    "costo al kwh": "costo_kwh_attuale", "prezzo kwh": "costo_kwh_attuale",
    "spese fisse": "spese_fisse_attuale", "spese fisse attuali": "spese_fisse_attuale",
    "costo smc": "costo_smc_attuale", "costo al smc": "costo_smc_attuale", "prezzo smc": "costo_smc_attuale",
    "data contratto": "data_contratto", "data di contratto": "data_contratto",
    "data verifica": "data_verifica", "data di verifica": "data_verifica",
    "data cambio": "data_cambio", "data di cambio": "data_cambio",
    "tipo contratto": "tipo_contratto", "fisso o variabile": "tipo_contratto",
    "nuovo fornitore": "nuovo_fornitore", "fornitore nuovo": "nuovo_fornitore",
    "costo kwh nuovo": "costo_kwh_nuovo", "nuovo costo kwh": "costo_kwh_nuovo", "nuovo costo al kw": "costo_kwh_nuovo",
    "spese fisse nuovo": "spese_fisse_nuovo", "nuove spese fisse": "spese_fisse_nuovo",
    "costo smc nuovo": "costo_smc_nuovo", "nuovo costo smc": "costo_smc_nuovo",
    "privacy": "privacy_firmata", "privacy firmata": "privacy_firmata",
    "pagato": "pagato", "note": "note",
    "lavorazione": "lavorazione", "stato": "lavorazione", "tipologia lavorazione": "lavorazione",
    "venditore": "venditore", "venduto da": "venditore",
}

VENDITORE_ALIASES = {"davis": "devis", "debby": "deborah", "debora": "deborah"}

async def _venditori_name_map():
    m = {}
    async for v in db.venditori.find({"attivo": {"$ne": False}}, {"_id": 0, "id": 1, "nome": 1}):
        nome = str(v.get("nome", "")).strip().lower()
        if nome:
            m[nome] = v["id"]
    return m

def _venditore_match(raw, vend_map: dict, fallback_id: str):
    v = re.sub(r"\s+", " ", str(raw or "").strip().lower())
    if not v:
        return "", False
    pagato = "pagat" in v
    v = re.sub(r"\bpagat[oa]?\b", "", v).strip()
    vid = vend_map.get(v) or vend_map.get(VENDITORE_ALIASES.get(v, "")) or fallback_id
    return vid, pagato

LAV_BY_LABEL = {
    "cambiare": "cambiare", "da cambiare": "cambiare",
    "non cambiare": "non_cambiare", "cambio effettuato": "cambio_effettuato",
    "in quotazione": "in_quotazione", "richieste bollette": "richieste_bollette",
    "richieste le bollette": "richieste_bollette", "in attesa ok": "in_attesa_ok",
    "in attesa ok cliente": "in_attesa_ok", "problema tecnico": "problema_tecnico",
    "da quotare": "da_quotare", "contattare cliente": "contattare_cliente",
    "non vuole cambiare": "non_vuole_cambiare", "passa in negozio": "passa_in_negozio",
    "attesa documenti": "attesa_documenti", "rinnovato": "rinnovato",
    "chieste le bollette": "richieste_bollette", "in attesa di ok del cliente": "in_attesa_ok",
    "attesa ok cliente": "in_attesa_ok", "in attesa di ok cliente": "in_attesa_ok",
}
LAV_BY_LABEL.update({l.replace("_", " "): l for l in LAVORAZIONI})

def _norm_header(h: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[._/\\-]", " ", str(h)).strip().lower())

def _parse_float(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == "":
        return None
    try:
        return float(re.sub(r"[^\d,.\-]", "", str(v)).replace(",", "."))
    except ValueError:
        return None

def _parse_date(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == "":
        return None
    try:
        return pd.to_datetime(str(v).strip(), dayfirst=True).date().isoformat()
    except Exception:
        return None

def _parse_bool(v) -> bool:
    return str(v).strip().lower() in ("si", "sì", "yes", "true", "1", "x", "ok", "pagato", "firmata")

def _norm_lavorazione(v) -> str:
    return LAV_BY_LABEL.get(str(v).strip().lower(), "da_quotare")

def _parse_date_flex(v):
    v = str(v).strip()
    if not v:
        return None
    if re.fullmatch(r"\d{5}", v):
        return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
    m = re.match(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", v)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        try:
            return date(y, mo, d).isoformat()
        except ValueError:
            return None
    return _parse_date(v)

def _find_gestionale_blocks(r: list) -> list:
    starts = []
    for j, c in enumerate(r):
        c = str(c).strip()
        if re.fullmatch(r"\d{1,4}", c) and j + 1 < len(r) and str(r[j + 1]).strip():
            window = [str(x).strip().upper() for x in r[j:j + 12]]
            if "LUCE" in window or "GAS" in window:
                starts.append(j)
    return starts

def _gestionale_block_to_doc(b: list, store_id: str, admin: dict, now: str,
                             vend_map: Optional[dict] = None, vend_fallback: str = "") -> Optional[dict]:
    servizio = b[10].upper()
    if servizio not in ("LUCE", "GAS"):
        return None
    toks = b[1].split()
    if not toks:
        return None
    tar = [float(n.replace(",", ".")) for n in re.findall(r"\d+[.,]\d+|\d+", b[16])]
    prezzo_n, fisse_n = None, None
    if len(tar) == 2:
        prezzo_n, fisse_n = tar
    elif len(tar) == 1 and tar[0] < 5:
        prezzo_n = tar[0]
    note = b[20]
    if b[16] and len(tar) > 2:
        note = (note + " | Tariffa nuova: " + b[16]).strip(" |")
    lav_raw = b[21].strip().lower()
    lav = _norm_lavorazione(lav_raw)
    if lav_raw and lav == "da_quotare" and lav_raw != "da quotare":
        note = (note + " | Stato foglio: " + b[21]).strip(" |")
    if b[19].strip().lower() == "annullato":
        note = (note + " | Pagamento: annullato").strip(" |")
    pagato = _parse_bool(b[19])
    op_id, vend_pagato = _venditore_match(b[22], vend_map or {}, vend_fallback)
    is_gas = servizio == "GAS"
    return {
        "id": str(uuid.uuid4()),
        "nome": toks[0], "cognome": " ".join(toks[1:]),
        "tipo_cliente": "business" if b[3] else "privato",
        "codice_fiscale": b[2].upper(), "p_iva": b[3], "indirizzo": b[4],
        "pod": "" if is_gas else b[5].upper(), "pdr": b[5] if is_gas else "",
        "iban": b[6].upper(), "email": b[7], "telefono": b[8],
        "kw_potenza": _parse_float(b[9]),
        "tipo_bolletta": "gas" if is_gas else "luce",
        "fornitore_provenienza": b[11],
        "costo_kwh_attuale": None, "spese_fisse_attuale": None, "costo_smc_attuale": None,
        "data_contratto": _parse_date_flex(b[14]),
        "data_verifica": _parse_date_flex(b[12]),
        "data_cambio": _parse_date_flex(b[13]),
        "tipo_contratto": "variabile" if "var" in b[15].lower() else "fisso",
        "nuovo_fornitore": b[17],
        "costo_kwh_nuovo": None if is_gas else prezzo_n,
        "spese_fisse_nuovo": fisse_n,
        "costo_smc_nuovo": prezzo_n if is_gas else None,
        "privacy_firmata": False,
        "note": note,
        "lavorazione": lav,
        "venditore_id": store_id, "operatore_id": op_id,
        "venditore_pagato": vend_pagato,
        "pagato": pagato,
        "last_payment_date": date.today().isoformat() if pagato else None,
        "created_by": admin["id"], "created_at": now, "updated_at": now,
    }

def _parse_gestionale_csv(text: str, store_id: str, admin: dict,
                          vend_map: Optional[dict] = None, vend_fallback: str = ""):
    rows = list(csv.reader(io.StringIO(text)))
    hdr_i = None
    for i, r in enumerate(rows[:10]):
        if any(re.sub(r"[^a-z0-9 ]", "", str(c).lower()).strip() == "nome cognome" for c in r):
            hdr_i = i
            break
    if hdr_i is None:
        return None, "Intestazioni 'nome cognome' non trovate"
    now = datetime.now(timezone.utc).isoformat()
    docs, skipped = [], 0
    for r in rows[hdr_i + 1:]:
        starts = _find_gestionale_blocks(r)
        if not starts:
            skipped += 1
        for s in starts:
            b = [str(x).strip() for x in r[s:s + 23]] + [""] * max(0, 23 - len(r[s:s + 23]))
            doc = _gestionale_block_to_doc(b, store_id, admin, now, vend_map, vend_fallback)
            if doc:
                docs.append(doc)
            else:
                skipped += 1
    if not docs:
        return None, "Nessun cliente riconosciuto nel formato gestionale"
    return (docs, skipped), None

def _parse_any_csv(text: str, store_id: str, admin: dict,
                   vend_map: Optional[dict] = None, vend_fallback: str = ""):
    head = text[:2000].lower()
    if "nome cognome" in head and "servizio" in head:
        return _parse_gestionale_csv(text, store_id, admin, vend_map, vend_fallback)
    return _parse_sheet_csv(text, store_id, admin, vend_map, vend_fallback)

def _build_col_map(df) -> dict:
    col_map = {}
    for col in df.columns:
        field = HEADER_MAP.get(_norm_header(col))
        if field and field not in col_map.values():
            col_map[col] = field
    return col_map

def _row_to_doc(data: dict, store_id: str, admin: dict, now: str,
                vend_map: Optional[dict] = None, vend_fallback: str = "") -> dict:
    tb = str(data.get("tipo_bolletta", "")).strip().lower()
    tc = str(data.get("tipo_contratto", "")).strip().lower()
    pagato = _parse_bool(data.get("pagato", ""))
    op_id, vend_pagato = _venditore_match(data.get("venditore", ""), vend_map or {}, vend_fallback)
    return {
        "id": str(uuid.uuid4()),
        "nome": str(data.get("nome", "")).strip(), "cognome": str(data.get("cognome", "")).strip(),
        "tipo_cliente": "business" if str(data.get("p_iva", "")).strip() else "privato",
        "codice_fiscale": str(data.get("codice_fiscale", "")).strip().upper(),
        "p_iva": str(data.get("p_iva", "")).strip(),
        "indirizzo": str(data.get("indirizzo", "")).strip(),
        "pod": str(data.get("pod", "")).strip().upper(),
        "pdr": str(data.get("pdr", "")).strip(),
        "iban": str(data.get("iban", "")).strip().upper(),
        "email": str(data.get("email", "")).strip(),
        "telefono": str(data.get("telefono", "")).strip(),
        "kw_potenza": _parse_float(data.get("kw_potenza")),
        "tipo_bolletta": "gas" if "gas" in tb else "luce",
        "fornitore_provenienza": str(data.get("fornitore_provenienza", "")).strip(),
        "costo_kwh_attuale": _parse_float(data.get("costo_kwh_attuale")),
        "spese_fisse_attuale": _parse_float(data.get("spese_fisse_attuale")),
        "costo_smc_attuale": _parse_float(data.get("costo_smc_attuale")),
        "data_contratto": _parse_date(data.get("data_contratto")),
        "data_verifica": _parse_date(data.get("data_verifica")),
        "data_cambio": _parse_date(data.get("data_cambio")),
        "tipo_contratto": "variabile" if "var" in tc else "fisso",
        "nuovo_fornitore": str(data.get("nuovo_fornitore", "")).strip(),
        "costo_kwh_nuovo": _parse_float(data.get("costo_kwh_nuovo")),
        "spese_fisse_nuovo": _parse_float(data.get("spese_fisse_nuovo")),
        "costo_smc_nuovo": _parse_float(data.get("costo_smc_nuovo")),
        "privacy_firmata": _parse_bool(data.get("privacy_firmata", "")),
        "note": str(data.get("note", "")).strip(),
        "lavorazione": _norm_lavorazione(data.get("lavorazione", "")),
        "venditore_id": store_id, "operatore_id": op_id,
        "venditore_pagato": vend_pagato,
        "pagato": pagato,
        "last_payment_date": date.today().isoformat() if pagato else None,
        "created_by": admin["id"], "created_at": now, "updated_at": now,
    }

def _parse_sheet_csv(text: str, store_id: str, admin: dict,
                     vend_map: Optional[dict] = None, vend_fallback: str = ""):
    try:
        df = pd.read_csv(io.StringIO(text), dtype=str)
    except Exception:
        return None, "Formato del foglio non leggibile"
    col_map = _build_col_map(df)
    if "nome" not in col_map.values() and "cognome" not in col_map.values():
        return None, f"Colonne 'nome'/'cognome' non trovate. Intestazioni lette: {', '.join(str(c) for c in df.columns[:15])}"
    now = datetime.now(timezone.utc).isoformat()
    docs, skipped = [], 0
    for _, row in df.iterrows():
        data = {field: (row[col] if pd.notna(row[col]) else "") for col, field in col_map.items()}
        if not str(data.get("nome", "")).strip() and not str(data.get("cognome", "")).strip():
            skipped += 1
            continue
        docs.append(_row_to_doc(data, store_id, admin, now, vend_map, vend_fallback))
    return (docs, skipped), None

async def _insert_imported(docs: list, admin: dict):
    now = datetime.now(timezone.utc).isoformat()
    await db.clients.insert_many(docs)
    await db.lavorazioni_log.insert_many([
        {"id": str(uuid.uuid4()), "client_id": d["id"], "client_name": f"{d['cognome']} {d['nome']}".strip(),
         "operatore_id": admin["id"], "operatore_name": admin["name"],
         "status": d["lavorazione"], "note": "Importato da Google Sheet", "created_at": now}
        for d in docs])

class SheetImportInput(BaseModel):
    sheet_url: str
    store_id: str = ""
    all_tabs: bool = False
    sheet_name: str = ""

async def _fetch_tab_csv(http_client, sheet_id: str, gid: Optional[str] = None, sheet_name: str = ""):
    if sheet_name:
        url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/gviz/tq?tqx=out:csv&sheet={sheet_name}"
    else:
        url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid or '0'}"
    return await http_client.get(url)

async def _import_store_tab(http_client, sheet_id: str, store: dict, fallback_hash: str, admin: dict,
                            vend_map: Optional[dict] = None, vend_fallback: str = "") -> dict:
    r = await _fetch_tab_csv(http_client, sheet_id, sheet_name=store["nome"])
    if r.status_code != 200 or hashlib.sha256(r.text.encode()).hexdigest() == fallback_hash:
        return {"store": store["nome"], "status": "pagina_non_trovata", "imported": 0}
    parsed, err = _parse_any_csv(r.text, store["id"], admin, vend_map, vend_fallback)
    if err:
        return {"store": store["nome"], "status": "errore", "detail": err, "imported": 0}
    docs, skipped = parsed
    if docs:
        await _insert_imported(docs, admin)
    return {"store": store["nome"], "status": "ok", "imported": len(docs), "skipped": skipped}

@api_router.post("/import/google-sheet")
async def import_google_sheet(input: SheetImportInput, admin: dict = Depends(require_admin)):
    m = re.search(r"/d/([a-zA-Z0-9\-_]+)", input.sheet_url)
    if not m:
        raise HTTPException(status_code=400, detail="Link Google Sheet non valido")
    sheet_id = m.group(1)
    gid_m = re.search(r"gid=(\d+)", input.sheet_url)
    vend_map = await _venditori_name_map()
    vend_fallback = vend_map.get("enrico", "")

    async with httpx.AsyncClient(timeout=30, follow_redirects=True) as http_client:
        if input.all_tabs:
            stores = await db.stores.find({}, {"_id": 0}).to_list(500)
            fallback_resp = await _fetch_tab_csv(http_client, sheet_id, sheet_name="__non_esiste__")
            if fallback_resp.status_code != 200:
                raise HTTPException(status_code=400, detail="Impossibile leggere il foglio. Verifica che sia condiviso: 'Chiunque abbia il link può visualizzare'.")
            fallback_hash = hashlib.sha256(fallback_resp.text.encode()).hexdigest()
            report, total_imported = [], 0
            for s in stores:
                rep = await _import_store_tab(http_client, sheet_id, s, fallback_hash, admin,
                                              vend_map, vend_fallback)
                report.append(rep)
                total_imported += rep["imported"]
            return {"mode": "all_tabs", "total_imported": total_imported, "report": report,
                    "hint": "Le pagine segnate come non trovate devono avere lo stesso nome del negozio, oppure importale singolarmente dal link della pagina (contiene gid)."}

        resp = await _fetch_tab_csv(http_client, sheet_id,
                                    gid=gid_m.group(1) if gid_m else None,
                                    sheet_name=input.sheet_name)
    if resp.status_code != 200 or ("text/csv" not in resp.headers.get("content-type", "") and "text/plain" not in resp.headers.get("content-type", "")):
        raise HTTPException(status_code=400, detail="Impossibile leggere il foglio. Verifica che sia condiviso: 'Chiunque abbia il link può visualizzare'.")
    parsed, err = _parse_any_csv(resp.text, input.store_id, admin, vend_map, vend_fallback)
    if err:
        raise HTTPException(status_code=400, detail=err)
    docs, skipped = parsed
    if docs:
        await _insert_imported(docs, admin)
    return {"imported": len(docs), "skipped": skipped}

# ---------------- Migrazione dati storici (idempotente, admin) ----------------

VENDITORI_SEED = [
    ("Deborah", "Morbegno", "interno"), ("Silvio", "Morbegno", "interno"),
    ("Bruno", "Morbegno", "interno"), ("Michael", "Tirano", "interno"),
    ("Lorenzo", "Sondalo", "interno"), ("Seba", "Grosio", "interno"),
    ("Enrico", "Sondrio", "interno"), ("Devis", "", "esterno"),
]
ENERGY_SHEET_ID = "19pEn41GLbJi6iI6BQ7v83idmro4GUoazUwdLGBUuksY"


def _nome_key(nome: str, cognome: str) -> str:
    return re.sub(r"\s+", " ", f"{nome} {cognome}".strip().lower())


RINNOVI_GIA_AVVISATI = [("3497732677", "2026-11-01"), ("3496179229", "2026-09-30"), ("3388029097", "2026-10-06"),
                        ("3665273901", "2026-10-08"), ("3332119225", "2026-10-08"), ("3896396032", "2026-10-20"),
                        ("3895182566", "2026-10-25"), ("335406976", "2026-09-28"), ("3280167807", "2026-10-05"),
                        ("3493200225", "2026-10-20")]

@api_router.post("/admin/migra-dati-storici")
async def migra_dati_storici(admin: dict = Depends(require_admin)):
    """Idempotente: rinomina/pulisce negozi, seed venditori, import colonna venditore dal foglio energia."""
    report = {"rinominati": [], "eliminati": [], "venditori_seedati": 0,
              "clienti_aggiornati": 0, "clienti_pagati": 0, "dettagli": [], "rinnovi_gia_avvisati": 0}
    now = datetime.now(timezone.utc).isoformat()
    # Avvisi rinnovo energia già inviati il 09/09/2026 (dall'ambiente di anteprima): evita doppio invio in produzione
    for tel, scad in RINNOVI_GIA_AVVISATI:
        res = await db.clients.update_many({"telefono": tel, "rinnovo_msg_sent_for": {"$ne": scad}},
                                           {"$set": {"rinnovo_msg_sent_for": scad, "rinnovo_msg_sent_at": "2026-09-09T10:32:00+00:00"}})
        report["rinnovi_gia_avvisati"] += res.modified_count

    r = await db.stores.update_one({"nome": "Sondrio Grosio"}, {"$set": {"nome": "Grosio"}})
    if r.modified_count:
        report["rinominati"].append("Sondrio Grosio -> Grosio")

    deriu = await db.stores.find_one({"nome": "Deriu"}, {"_id": 0})
    if deriu:
        cl_ids = [c["id"] for c in await db.clients.find({"venditore_id": deriu["id"]},
                                                         {"_id": 0, "id": 1}).to_list(500)]
        if cl_ids:
            await db.lavorazioni_log.delete_many({"client_id": {"$in": cl_ids}})
            await db.servizi.delete_many({"client_id": {"$in": cl_ids}})
            await db.clients.delete_many({"venditore_id": deriu["id"]})
        await db.stores.delete_one({"id": deriu["id"]})
        report["eliminati"].append("Deriu")
    for nome in ("Devis (Freelance)", "TEST_Negozio"):
        res = await db.stores.delete_many({"nome": nome})
        if res.deleted_count:
            report["eliminati"].append(f"{nome} x{res.deleted_count}")

    stores_by_name = {s["nome"]: s["id"] for s in await db.stores.find({}, {"_id": 0}).to_list(100)}
    for nome, store_nome, tipo in VENDITORI_SEED:
        if not await db.venditori.find_one({"nome": nome}):
            await db.venditori.insert_one({"id": str(uuid.uuid4()), "nome": nome,
                                           "store_id": stores_by_name.get(store_nome, ""),
                                           "tipo": tipo, "attivo": True, "created_at": now})
            report["venditori_seedati"] += 1

    vend_map = await _venditori_name_map()
    vend_fallback = vend_map.get("enrico", "")
    async with httpx.AsyncClient(timeout=30, follow_redirects=True) as http:
        fb = await _fetch_tab_csv(http, ENERGY_SHEET_ID, sheet_name="__non_esiste__")
        if fb.status_code != 200:
            raise HTTPException(status_code=400, detail="Foglio Google non raggiungibile")
        fb_hash = hashlib.sha256(fb.text.encode()).hexdigest()
        for store_name, store_id in stores_by_name.items():
            r = await _fetch_tab_csv(http, ENERGY_SHEET_ID, sheet_name=store_name)
            if r.status_code != 200:
                continue
            if hashlib.sha256(r.text.encode()).hexdigest() == fb_hash and store_name != "Sondrio":
                continue  # il fallback di Google restituisce il primo foglio (Sondrio)
            if "venditore" not in r.text[:3000].lower():
                continue
            parsed, err = _parse_any_csv(r.text, store_id, admin, vend_map, vend_fallback)
            if err:
                continue
            docs, _ = parsed
            sheet_map = {}
            for d in docs:
                if d.get("operatore_id"):
                    sheet_map[_nome_key(d["nome"], d["cognome"])] = (d["operatore_id"], d.get("venditore_pagato", False))
            updated = 0
            async for c in db.clients.find({"venditore_id": store_id},
                                           {"_id": 0, "id": 1, "nome": 1, "cognome": 1}):
                k = _nome_key(c.get("nome", ""), c.get("cognome", ""))
                if k not in sheet_map:
                    continue
                op_id, vpag = sheet_map[k]
                upd = {"operatore_id": op_id}
                if vpag:
                    upd["venditore_pagato"] = True
                    report["clienti_pagati"] += 1
                await db.clients.update_one({"id": c["id"]}, {"$set": upd})
                updated += 1
            if updated:
                report["dettagli"].append(f"{store_name}: {updated} clienti")
                report["clienti_aggiornati"] += updated
    return report

# ---------------- Import storico riparazioni dai fogli Google (idempotente, admin) ----------------

RIPARAZIONI_SHEETS = {
    "Morbegno": "1IogWl5ISTndU0_IzQEGE_NUZzGNtiSC1DbSt9XilT0E",
    "Gravedona": "1VmPtZXW72MWcyENfjkQ1wDuMj2PfxXCRi2w60nQn65E",
    "Tirano": "1lP7mJo7Y0-lOPCiAKBMjMj1WT91P8pbZoYbuEyesH7w",
    "Grosio": "15qcxX8Sf14WbDOYuTirggmzfrZ89_xctVKJKJgWgTCk",
    "Sondrio": "197nm49W3d6jT0jaE5qD5tE2mF-HHZLs48se_g1lWHhA",
    "Sondalo": "1LAOpPtEpz1FrMDXwR8wV1UghMXNSJa23wYmo6PvEmgc",
}
RIP_PREFIX = {"morbegno": "RM", "sondrio": "RSO", "gravedona": "RG",
              "tirano": "RT", "sondalo": "RSA", "grosio": "RGR"}
RIP_STATO_MAP = {
    "consegnato": "consegnato", "pronto": "pronto", "in lavorazione": "in_lavorazione",
    "preventivato": "preventivo", "non riparabile": "non_riparabile",
    "in attesa cliente": "in_attesa_cliente", "": "ingresso",
}

def _rip_stato(raw: str) -> str:
    s = re.sub(r"\s+", " ", str(raw or "").strip().lower())
    if s in RIP_STATO_MAP:
        return RIP_STATO_MAP[s]
    if "ricambio" in s:
        return "attesa_ricambio_carico"
    return "ingresso"

def _rip_date(v: str) -> str:
    m = re.match(r"(\d{1,2})-(\d{1,2})-(\d{4})", str(v or "").strip())
    if not m:
        return datetime.now(timezone.utc).isoformat()
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        return datetime(y, mo, d, 10, 0, tzinfo=timezone.utc).isoformat()
    except ValueError:
        return datetime.now(timezone.utc).isoformat()

async def _find_or_create_rip_client(nome: str, cognome: str, telefono: str, store_id: str, now: str) -> str:
    scope = {"venditore_id": store_id}
    if telefono:
        found = await db.clients.find_one({**scope, "telefono": telefono}, {"_id": 0, "id": 1})
        if found:
            return found["id"]
    key = _nome_key(nome, cognome)
    async for c in db.clients.find(scope, {"_id": 0, "id": 1, "nome": 1, "cognome": 1}):
        if _nome_key(c.get("nome", ""), c.get("cognome", "")) == key:
            return c["id"]
    cid = str(uuid.uuid4())
    await db.clients.insert_one({
        "id": cid, "nome": nome, "cognome": cognome, "tipo_cliente": "privato",
        "codice_fiscale": "", "p_iva": "", "indirizzo": "", "pod": "", "pdr": "",
        "iban": "", "email": "", "telefono": telefono, "kw_potenza": None,
        "tipo_bolletta": "luce", "fornitore_provenienza": "",
        "costo_kwh_attuale": None, "spese_fisse_attuale": None, "costo_smc_attuale": None,
        "data_contratto": None, "data_verifica": None, "data_cambio": None,
        "tipo_contratto": "fisso", "nuovo_fornitore": "",
        "costo_kwh_nuovo": None, "spese_fisse_nuovo": None, "costo_smc_nuovo": None,
        "privacy_firmata": False, "note": "", "lavorazione": "",
        "venditore_id": store_id, "operatore_id": "", "pagato": False,
        "last_payment_date": None, "import_source": "fogli_riparazioni",
        "created_by": "import", "created_at": now, "updated_at": now})
    return cid

@api_router.post("/admin/import-riparazioni-storiche")
async def import_riparazioni_storiche(admin: dict = Depends(require_admin)):
    """Idempotente: importa le riparazioni storiche dai 6 fogli Google (una per negozio)."""
    marker = await db.import_state.find_one({"_id": "riparazioni_fogli_v1"})
    if marker:
        return {"status": "gia_importato", **{k: v for k, v in marker.items() if k != "_id"}}
    stores_by_name = {s["nome"]: s["id"] for s in await db.stores.find({}, {"_id": 0}).to_list(100)}
    now = datetime.now(timezone.utc).isoformat()
    report, tot_servizi, tot_clienti_new = [], 0, 0
    async with httpx.AsyncClient(timeout=30, follow_redirects=True) as http:
        for store_name, sheet_id in RIPARAZIONI_SHEETS.items():
            store_id = stores_by_name.get(store_name)
            if not store_id:
                report.append({"store": store_name, "status": "negozio_mancante"})
                continue
            r = await http.get(f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid=0")
            if r.status_code != 200:
                report.append({"store": store_name, "status": "foglio_non_leggibile"})
                continue
            rows = list(csv.reader(io.StringIO(r.text)))
            hdr_i = next((i for i, row in enumerate(rows[:5])
                          if any(c.strip().lower() == "cognome" for c in row)), None)
            if hdr_i is None:
                report.append({"store": store_name, "status": "intestazioni_mancanti"})
                continue
            hdr = [c.strip().lower() for c in rows[hdr_i]]
            prefix = next((p for key, p in RIP_PREFIX.items() if key in store_name.lower()), "RIP")
            count, seq = 0, 1
            clients_before = await db.clients.count_documents({})
            for row in rows[hdr_i + 1:]:
                row = row + [""] * (len(hdr) - len(row))
                d = dict(zip(hdr, row))
                cognome, nome = d.get("cognome", "").strip(), d.get("nome", "").strip()
                if not cognome and not nome:
                    continue
                telefono = re.sub(r"[^\d]", "", d.get("numero", ""))
                cid = await _find_or_create_rip_client(nome.capitalize(), cognome.capitalize(),
                                                       telefono, store_id, now)
                await db.servizi.insert_one({
                    "id": str(uuid.uuid4()), "tipo": "riparazione", "client_id": cid,
                    "venditore_id": store_id, "operatore_id": "",
                    "dispositivo": d.get("modello", "").strip(),
                    "problema": d.get("intervento", "").strip(),
                    "note": d.get("note", "").strip(),
                    "stato": _rip_stato(d.get("stato", "")),
                    "pagato": _parse_bool(d.get("pagato", "")),
                    "numero_riparazione": f"{prefix}{seq}",
                    "created_at": _rip_date(d.get("data di consegna", "")),
                    "updated_at": now, "import_source": "fogli_riparazioni"})
                count += 1
                seq += 1
            tot_clienti_new += await db.clients.count_documents({}) - clients_before
            tot_servizi += count
            if count:
                await db.counters.update_one({"_id": f"riparazione:{store_id}"},
                                             {"$set": {"prefix": prefix, "seq": seq}}, upsert=True)
            report.append({"store": store_name, "status": "ok", "importate": count})
    await db.import_state.insert_one({"_id": "riparazioni_fogli_v1", "at": now,
                                      "servizi_creati": tot_servizi, "clienti_creati": tot_clienti_new})
    return {"status": "ok", "servizi_creati": tot_servizi,
            "clienti_creati": tot_clienti_new, "report": report}

# ---------------- Sincronizzazione incrementale fogli Google (admin, ripetibile) ----------------

SYNC_CLIENT_FIELDS = ["codice_fiscale", "p_iva", "indirizzo", "pod", "pdr", "iban", "email", "telefono", "kw_potenza",
                      "fornitore_provenienza", "costo_kwh_attuale", "spese_fisse_attuale", "costo_smc_attuale",
                      "data_contratto", "data_verifica", "data_cambio", "nuovo_fornitore", "costo_kwh_nuovo",
                      "spese_fisse_nuovo", "costo_smc_nuovo", "note", "lavorazione"]

def _client_sync_key(c: dict) -> tuple:
    ref = (c.get("pod") or c.get("pdr") or c.get("indirizzo") or "").strip().lower()
    return ((c.get("nome") or "").strip().lower(), (c.get("cognome") or "").strip().lower(), c.get("tipo_bolletta") or "luce", ref)

def _client_sync_key_loose(c: dict) -> tuple:
    return ((c.get("nome") or "").strip().lower(), (c.get("cognome") or "").strip().lower(), c.get("tipo_bolletta") or "luce")

def _client_sync_key_tokens(c: dict) -> tuple:
    toks = sorted(re.findall(r"\w+", f"{c.get('nome') or ''} {c.get('cognome') or ''}".lower()))
    return (" ".join(toks), c.get("tipo_bolletta") or "luce")

def _client_sync_key_pod(c: dict) -> Optional[str]:
    ref = (c.get("pod") or c.get("pdr") or "").strip().upper()
    return ref or None

class SyncFogliInput(BaseModel):
    dry_run: bool = True
    energia: bool = True
    riparazioni: bool = True

ENERGY_TAB_ALIAS = {"Sondrio": ["CAMBIAORA"], "Gravedona": ["KEVIN"]}

async def _sync_energia(admin: dict, dry_run: bool) -> dict:
    vend_map = await _venditori_name_map()
    vend_fallback = vend_map.get("enrico", "")
    rep = {"nuovi": 0, "aggiornati": 0, "invariati": 0, "campi_aggiornati": {}, "negozi": [], "esempi_nuovi": []}
    now = datetime.now(timezone.utc).isoformat()
    async with httpx.AsyncClient(timeout=30, follow_redirects=True) as http:
        page = await http.get(f"https://docs.google.com/spreadsheets/d/{ENERGY_SHEET_ID}/htmlview")
        tabs = {name.strip().lower(): gid for name, gid in re.findall(r'items\.push\(\{name: "([^"]*)".*?gid: "(-?\d+)"', page.text)}
        for s in await db.stores.find({}, {"_id": 0}).to_list(100):
            r, tab_used = None, None
            for tab in [s["nome"]] + ENERGY_TAB_ALIAS.get(s["nome"], []):
                gid = tabs.get(tab.lower())
                if gid is None:
                    continue
                r = await _fetch_tab_csv(http, ENERGY_SHEET_ID, gid=gid)
                if r.status_code == 200:
                    tab_used = tab
                    break
            if not tab_used:
                rep["negozi"].append({"store": s["nome"], "status": "pagina_non_trovata"})
                continue
            parsed, err = _parse_any_csv(r.text, s["id"], admin, vend_map, vend_fallback)
            if err:
                rep["negozi"].append({"store": s["nome"], "status": "errore", "detail": err})
                continue
            docs, _ = parsed
            seen, uniq = set(), []
            for d in docs:
                k = (_client_sync_key(d), d.get("data_contratto"))
                if k in seen:
                    continue
                seen.add(k)
                uniq.append(d)
            docs = uniq
            existing = await db.clients.find({"venditore_id": s["id"], "anonimizzato": {"$ne": True}}, {"_id": 0}).to_list(5000)
            indexes = [(_client_sync_key, {}), (_client_sync_key_pod, {}), (_client_sync_key_loose, {}), (_client_sync_key_tokens, {})]
            for c in existing:
                for fn, idx in indexes:
                    k = fn(c)
                    if k:
                        idx.setdefault(k, []).append(c)
            n_new = n_upd = 0
            for d in docs:
                cands = next((idx.get(fn(d)) for fn, idx in indexes if fn(d) and idx.get(fn(d))), None)
                if not cands:
                    n_new += 1
                    if len(rep["esempi_nuovi"]) < 15:
                        rep["esempi_nuovi"].append(f"{s['nome']}: {d['cognome']} {d['nome']} ({d['tipo_bolletta']})")
                    if not dry_run:
                        d["import_source"] = "sync_fogli"
                        await _insert_imported([d], admin)
                    continue
                c = cands.pop(0)
                for fn, idx in indexes:
                    lst = idx.get(fn(c)) if fn(c) else None
                    if lst and c in lst:
                        lst.remove(c)
                changes = {}
                for f in SYNC_CLIENT_FIELDS:
                    v = d.get(f)
                    if v in (None, "", 0.0) or v == c.get(f):
                        continue
                    if f == "lavorazione" and not v:
                        continue
                    changes[f] = v
                if d.get("pagato") and not c.get("pagato"):
                    changes["pagato"] = True
                    changes["last_payment_date"] = date.today().isoformat()
                if d.get("privacy_firmata") and not c.get("privacy_firmata"):
                    changes["privacy_firmata"] = True
                if d.get("operatore_id") and d["operatore_id"] != vend_fallback and not c.get("operatore_id"):
                    changes["operatore_id"] = d["operatore_id"]
                    changes["venditore_pagato"] = d.get("venditore_pagato", False)
                if not changes:
                    rep["invariati"] += 1
                    continue
                n_upd += 1
                for f in changes:
                    rep["campi_aggiornati"][f] = rep["campi_aggiornati"].get(f, 0) + 1
                if not dry_run:
                    changes["updated_at"] = now
                    await db.clients.update_one({"id": c["id"]}, {"$set": changes})
                    if "lavorazione" in changes:
                        await db.lavorazioni_log.insert_one({
                            "id": str(uuid.uuid4()), "client_id": c["id"], "client_name": f"{c['cognome']} {c['nome']}".strip(),
                            "operatore_id": admin["id"], "operatore_name": admin["name"], "status": changes["lavorazione"],
                            "note": "Aggiornato da foglio Google (sync)", "created_at": now})
            rep["nuovi"] += n_new
            rep["aggiornati"] += n_upd
            rep["negozi"].append({"store": s["nome"], "tab": tab_used, "status": "ok", "righe_foglio": len(docs), "nuovi": n_new, "aggiornati": n_upd})
    return rep

async def _sync_riparazioni(admin: dict, dry_run: bool) -> dict:
    stores_by_name = {s["nome"]: s["id"] for s in await db.stores.find({}, {"_id": 0}).to_list(100)}
    marker = await db.import_state.find_one({"_id": "riparazioni_fogli_v1"})
    import_at = (marker or {}).get("at")
    now = datetime.now(timezone.utc).isoformat()
    rep = {"nuove": 0, "aggiornate": 0, "invariate": 0, "saltate_modificate_in_app": 0, "negozi": [], "esempi_nuove": []}
    async with httpx.AsyncClient(timeout=30, follow_redirects=True) as http:
        for store_name, sheet_id in RIPARAZIONI_SHEETS.items():
            store_id = stores_by_name.get(store_name)
            if not store_id:
                rep["negozi"].append({"store": store_name, "status": "negozio_mancante"})
                continue
            r = await http.get(f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid=0")
            if r.status_code != 200:
                rep["negozi"].append({"store": store_name, "status": "foglio_non_leggibile"})
                continue
            rows = list(csv.reader(io.StringIO(r.text)))
            hdr_i = next((i for i, row in enumerate(rows[:5]) if any(c.strip().lower() == "cognome" for c in row)), None)
            if hdr_i is None:
                rep["negozi"].append({"store": store_name, "status": "intestazioni_mancanti"})
                continue
            hdr = [c.strip().lower() for c in rows[hdr_i]]
            prefix = next((p for key, p in RIP_PREFIX.items() if key in store_name.lower()), "RIP")
            existing = await db.servizi.find({"tipo": "riparazione", "venditore_id": store_id}, {"_id": 0}).to_list(10000)
            cl_names = await clients_name_map(list({s["client_id"] for s in existing}))
            by_key, by_dev = {}, {}
            for s in existing:
                if s.get("sheet_key"):
                    by_key.setdefault(tuple(s["sheet_key"]), []).append(s)
                    continue
                nm = (cl_names.get(s["client_id"], "") or "").strip().lower()
                by_key.setdefault((nm, (s.get("dispositivo") or "").strip().lower(), (s.get("created_at") or "")[:10]), []).append(s)
                by_dev.setdefault(((s.get("dispositivo") or "").strip().lower(), (s.get("created_at") or "")[:10]), []).append((nm, s))
            counter = await db.counters.find_one({"_id": f"riparazione:{store_id}"}) or {"seq": 1}
            seq = counter.get("seq", 1)
            n_new = n_upd = 0
            for row in rows[hdr_i + 1:]:
                row = row + [""] * (len(hdr) - len(row))
                d = dict(zip(hdr, row))
                cognome, nome = d.get("cognome", "").strip(), d.get("nome", "").strip()
                if not cognome and not nome:
                    continue
                modello = d.get("modello", "").strip()
                created = _rip_date(d.get("data di consegna", ""))
                key = (f"{cognome} {nome}".strip().lower(), modello.lower(), created[:10])
                cands = by_key.get(key)
                if not cands:
                    toks = set(re.findall(r"\w+", key[0]))
                    alt = [s for nm, s in by_dev.get((key[1], key[2]), []) if toks & set(re.findall(r"\w+", nm))]
                    if alt:
                        cands = alt
                        by_dev[(key[1], key[2])] = [(nm, s) for nm, s in by_dev[(key[1], key[2])] if s is not alt[0]]
                if not cands:
                    n_new += 1
                    if len(rep["esempi_nuove"]) < 15:
                        rep["esempi_nuove"].append(f"{store_name}: {cognome} {nome} - {modello} ({created[:10]})")
                    if not dry_run:
                        telefono = re.sub(r"[^\d]", "", d.get("numero", ""))
                        cid = await _find_or_create_rip_client(nome.capitalize(), cognome.capitalize(), telefono, store_id, now)
                        await db.servizi.insert_one({
                            "id": str(uuid.uuid4()), "tipo": "riparazione", "client_id": cid, "venditore_id": store_id, "operatore_id": "",
                            "dispositivo": modello, "problema": d.get("intervento", "").strip(), "note": d.get("note", "").strip(),
                            "stato": _rip_stato(d.get("stato", "")), "pagato": _parse_bool(d.get("pagato", "")),
                            "numero_riparazione": f"{prefix}{seq}", "created_at": created, "updated_at": now,
                            "import_source": "fogli_riparazioni", "sheet_synced_at": now, "sheet_key": list(key)})
                        seq += 1
                    continue
                s = cands.pop(0)
                untouched = s.get("updated_at") in (import_at, s.get("sheet_synced_at"))
                changes = {}
                for f, v in (("stato", _rip_stato(d.get("stato", ""))), ("problema", d.get("intervento", "").strip()),
                             ("note", d.get("note", "").strip())):
                    if v and v != s.get(f):
                        changes[f] = v
                if _parse_bool(d.get("pagato", "")) and not s.get("pagato"):
                    changes["pagato"] = True
                if not changes:
                    rep["invariate"] += 1
                    continue
                if not untouched:
                    rep["saltate_modificate_in_app"] += 1
                    continue
                n_upd += 1
                if not dry_run:
                    changes.update({"updated_at": now, "sheet_synced_at": now})
                    await db.servizi.update_one({"id": s["id"]}, {"$set": changes})
            if not dry_run and n_new:
                await db.counters.update_one({"_id": f"riparazione:{store_id}"}, {"$set": {"prefix": prefix, "seq": seq}}, upsert=True)
            rep["nuove"] += n_new
            rep["aggiornate"] += n_upd
            rep["negozi"].append({"store": store_name, "status": "ok", "nuove": n_new, "aggiornate": n_upd})
    return rep

@api_router.post("/admin/sync-fogli")
async def sync_fogli(input: SyncFogliInput, admin: dict = Depends(require_admin)):
    """Sincronizzazione incrementale dai fogli Google: aggiunge le righe nuove e aggiorna quelle cambiate (ripetibile)."""
    out = {"dry_run": input.dry_run}
    if input.energia:
        out["energia"] = await _sync_energia(admin, input.dry_run)
    if input.riparazioni:
        out["riparazioni"] = await _sync_riparazioni(admin, input.dry_run)
    if not input.dry_run:
        await db.import_state.update_one({"_id": "sync_fogli_last"}, {"$set": {"at": datetime.now(timezone.utc).isoformat(), "report": out}}, upsert=True)
    return out

# ---------------- Formazione (slide, manuale, servizi & operatori) ----------------
from formazione_seed import build_seed as _formazione_seed

class FormazioneInput(BaseModel):
    sezione: str  # slide | manuale | servizi
    titolo: str
    sottotitolo: str = ""
    contenuto: str = ""
    categoria: str = ""
    link: str = ""
    immagine: str = ""
    ordine: int = 0

FORMAZIONE_SEZIONI = {"slide", "manuale", "servizi"}

async def seed_formazione():
    if await db.formazione.count_documents({}) > 0:
        return
    portali = await db.portali.find({}, {"_id": 0}).to_list(200)
    now = datetime.now(timezone.utc).isoformat()
    docs = [{**d, "id": str(uuid.uuid4()), "link": d.get("link", ""), "immagine": d.get("immagine", ""), "created_at": now, "updated_at": now} for d in _formazione_seed(portali)]
    if docs:
        await db.formazione.insert_many(docs)

def _formazione_public(d: dict) -> dict:
    return {k: v for k, v in d.items() if k != "_id"}

@api_router.get("/public/formazione/presentazione")
async def public_presentazione():
    rows = await db.formazione.find({"sezione": "slide"}, {"_id": 0}).sort("ordine", 1).to_list(500)
    return rows

@api_router.get("/formazione")
async def list_formazione(user: dict = Depends(get_current_user), sezione: str = ""):
    q = {"sezione": sezione} if sezione else {}
    return await db.formazione.find(q, {"_id": 0}).sort([("sezione", 1), ("ordine", 1)]).to_list(1000)

@api_router.post("/formazione")
async def create_formazione(input: FormazioneInput, admin: dict = Depends(require_admin)):
    if input.sezione not in FORMAZIONE_SEZIONI:
        raise HTTPException(status_code=400, detail="Sezione non valida")
    now = datetime.now(timezone.utc).isoformat()
    doc = input.model_dump()
    if not doc["ordine"]:
        last = await db.formazione.find({"sezione": input.sezione}).sort("ordine", -1).limit(1).to_list(1)
        doc["ordine"] = (last[0]["ordine"] + 1) if last else 1
    doc.update({"id": str(uuid.uuid4()), "created_at": now, "updated_at": now, "updated_by": admin["name"]})
    await db.formazione.insert_one(doc)
    return _formazione_public(doc)

@api_router.patch("/formazione/{item_id}")
async def update_formazione(item_id: str, input: FormazioneInput, admin: dict = Depends(require_admin)):
    if input.sezione not in FORMAZIONE_SEZIONI:
        raise HTTPException(status_code=400, detail="Sezione non valida")
    doc = input.model_dump()
    doc.update({"updated_at": datetime.now(timezone.utc).isoformat(), "updated_by": admin["name"]})
    res = await db.formazione.update_one({"id": item_id}, {"$set": doc})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Voce non trovata")
    return _formazione_public(await db.formazione.find_one({"id": item_id}))

@api_router.delete("/formazione/{item_id}")
async def delete_formazione(item_id: str, admin: dict = Depends(require_admin)):
    res = await db.formazione.delete_one({"id": item_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Voce non trovata")
    return {"status": "ok"}

@api_router.post("/formazione/ripristina-default")
async def reset_formazione(admin: dict = Depends(require_admin)):
    await db.formazione.delete_many({})
    await seed_formazione()
    return {"status": "ok", "voci": await db.formazione.count_documents({})}

def _pdf_txt(s: str) -> str:
    s = (s or "").replace("→", "->").replace("−", "-").replace("’", "'").replace("“", '"').replace("”", '"').replace("·", "-").replace("€", "EUR")
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    return s.encode("latin-1", "replace").decode("latin-1")

def _pdf_markdown(pdf, text: str):
    for raw in (text or "").split("\n"):
        line = _pdf_txt(raw.rstrip())
        if not line.strip():
            pdf.ln(2)
            continue
        if line.startswith("## "):
            pdf.ln(2); pdf.set_font("helvetica", "B", 12); pdf.multi_cell(0, 6, line[3:], new_x="LMARGIN", new_y="NEXT"); pdf.set_font("helvetica", "", 10.5)
        elif line.startswith("# "):
            pdf.ln(2); pdf.set_font("helvetica", "B", 13); pdf.multi_cell(0, 7, line[2:], new_x="LMARGIN", new_y="NEXT"); pdf.set_font("helvetica", "", 10.5)
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(set(c) <= set("-: ") for c in cells):
                continue
            pdf.multi_cell(0, 5.5, "   " + "  -  ".join(cells), new_x="LMARGIN", new_y="NEXT")
        elif line.lstrip().startswith(("- ", "* ")):
            pdf.multi_cell(0, 5.5, "   - " + line.lstrip()[2:], new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.multi_cell(0, 5.5, line, new_x="LMARGIN", new_y="NEXT")

def _build_formazione_pdf(sezione: str, rows: list) -> bytes:
    titles = {"slide": "Presentazione del Gestionale", "manuale": "Manuale operativo e gestionale", "servizi": "Servizi e Operatori"}
    pdf = _new_pdf(titles.get(sezione, "Formazione"), f"RS Group - Cambia Ora - aggiornato al {date.today().strftime('%d/%m/%Y')}")
    pdf.set_font("helvetica", "", 10.5)
    for i, r in enumerate(rows):
        if sezione == "slide" and i > 0:
            pdf.add_page()
        elif pdf.get_y() > 240:
            pdf.add_page()
        pdf.set_font("helvetica", "B", 15 if sezione == "slide" else 13)
        pdf.multi_cell(0, 8, _pdf_txt(f"{i + 1}. {r['titolo']}"), new_x="LMARGIN", new_y="NEXT")
        if r.get("sottotitolo"):
            pdf.set_font("helvetica", "I", 10.5); pdf.multi_cell(0, 6, _pdf_txt(r["sottotitolo"]), new_x="LMARGIN", new_y="NEXT")
        if r.get("categoria"):
            pdf.set_font("helvetica", "", 9); pdf.set_text_color(120, 120, 120)
            pdf.multi_cell(0, 5, _pdf_txt(f"Sezione: {r['categoria']}"), new_x="LMARGIN", new_y="NEXT"); pdf.set_text_color(0, 0, 0)
        if r.get("link"):
            pdf.set_font("helvetica", "", 9); pdf.multi_cell(0, 5, _pdf_txt(f"Link: {r['link'][:110]}"), new_x="LMARGIN", new_y="NEXT")
        img = r.get("immagine") or ""
        if img.startswith("/formazione/") and os.path.exists(f"/app/frontend/public{img}"):
            pdf.ln(1)
            pdf.image(f"/app/frontend/public{img}", w=120)
            pdf.ln(2)
        pdf.set_font("helvetica", "", 10.5)
        pdf.ln(1)
        _pdf_markdown(pdf, r.get("contenuto", ""))
        pdf.ln(4)
    return bytes(pdf.output())

@api_router.get("/formazione/pdf")
async def formazione_pdf(sezione: str = "manuale", user: dict = Depends(get_current_user)):
    if sezione not in FORMAZIONE_SEZIONI:
        raise HTTPException(status_code=400, detail="Sezione non valida")
    rows = await db.formazione.find({"sezione": sezione}, {"_id": 0}).sort("ordine", 1).to_list(1000)
    return Response(content=_build_formazione_pdf(sezione, rows), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="formazione_{sezione}.pdf"'})

@api_router.get("/public/formazione/presentazione.pdf")
async def public_presentazione_pdf():
    rows = await db.formazione.find({"sezione": "slide"}, {"_id": 0}).sort("ordine", 1).to_list(500)
    return Response(content=_build_formazione_pdf("slide", rows), media_type="application/pdf",
                    headers={"Content-Disposition": 'attachment; filename="presentazione_gestionale.pdf"'})

# ---------------- Object Storage (allegati bollette/documenti) ----------------

STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
APP_NAME = "gestionale-utenze"
storage_key = None

def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = httpx.post(f"{STORAGE_URL}/init", json={"emergent_key": os.environ["EMERGENT_LLM_KEY"]}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key

def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = httpx.put(f"{STORAGE_URL}/objects/{path}",
                     headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    resp.raise_for_status()
    return resp.json()

def get_object(path: str):
    key = init_storage()
    resp = httpx.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")

def delete_object(path: str) -> None:
    try:
        key = init_storage()
        httpx.delete(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    except Exception as e:
        logger.warning(f"Eliminazione oggetto storage fallita {path}: {e}")

ANON_CLIENT_FIELDS = ["nome", "cognome", "codice_fiscale", "piva", "telefono", "email", "indirizzo", "iban", "pod_pdr",
                      "note", "privacy_msg_sent_at", "review_msg_sent_at", "rinnovo_msg_sent_at", "truffe_msg_sent_at", "data_nascita"]
ANON_SERVIZIO_FIELDS = ["numero", "iccid", "account_email", "problema", "operazioni", "note",
                        "codice_sblocco", "account_password", "codice_sblocco_enc", "account_password_enc"]

@api_router.post("/clients/{client_id}/anonimizza")
async def anonimizza_cliente(client_id: str, admin: dict = Depends(require_admin)):
    c = await db.clients.find_one({"id": client_id}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    if c.get("anonimizzato_at"):
        raise HTTPException(status_code=400, detail="Cliente già anonimizzato")
    return await anonimizza_cliente_core(c, admin["id"])

async def anonimizza_cliente_core(c: dict, by: str) -> dict:
    client_id = c["id"]
    now = datetime.now(timezone.utc).isoformat()
    tel = c.get("telefono", "")
    atts = await db.attachments.find({"client_id": client_id}, {"_id": 0, "storage_path": 1}).to_list(500)
    for a in atts:
        if a.get("storage_path"):
            delete_object(a["storage_path"])
    n_att = (await db.attachments.delete_many({"client_id": client_id})).deleted_count
    n_wa = (await db.wa_log.delete_many({"$or": [{"client_id": client_id}] + ([{"phone": tel}] if tel else [])})).deleted_count
    await db.whatsapp_queue.delete_many({"client_id": client_id})
    n_svc = (await db.servizi.update_many({"client_id": client_id},
                                          {"$unset": {f: "" for f in ANON_SERVIZIO_FIELDS}, "$set": {"anonimizzato_at": now}})).modified_count
    n_rit = (await db.ritiri.update_many({"client_id": client_id},
                                         {"$set": {"nome": "Anonimo", "cognome": "Anonimo", "codice_fiscale": "", "numero_documento": "", "imei": "", "anonimizzato_at": now}})).modified_count
    await db.lavorazioni_log.update_many({"client_id": client_id}, {"$set": {"note": ""}})
    await db.audit_log.update_many({"entity": "cliente", "entity_id": client_id}, {"$set": {"label": "Cliente anonimizzato"}})
    await db.clients.update_one({"id": client_id}, {
        "$unset": {f: "" for f in ANON_CLIENT_FIELDS if f not in ("nome", "cognome")},
        "$set": {"nome": "Anonimo", "cognome": f"GDPR-{client_id[:8].upper()}", "anonimizzato_at": now, "anonimizzato_da": by,
                 "no_recensioni": True}})
    return {"status": "ok", "allegati_eliminati": n_att, "messaggi_eliminati": n_wa, "servizi_anonimizzati": n_svc, "ritiri_anonimizzati": n_rit}

RETENTION_ANNI = 5
STATI_RIP_APERTI = ["ingresso", "attesa_ricambio_cliente", "attesa_ricambio_carico", "in_attesa_cliente", "preventivo", "in_lavorazione", "pronto"]

async def pulizia_retention_gdpr() -> dict:
    """Anonimizza i clienti senza alcuna attività (contratti, servizi, ritiri) negli ultimi RETENTION_ANNI anni."""
    limite = (datetime.now(timezone.utc) - relativedelta(years=RETENTION_ANNI)).isoformat()
    limite_d = limite[:10]
    candidati = await db.clients.find({"anonimizzato_at": {"$in": [None]}, "created_at": {"$lt": limite, "$gte": "2000-01-01"}}, {"_id": 0}).to_list(50000)
    report = []
    for c in candidati:
        compute_dates(c)
        if (c.get("data_scadenza") or "") >= limite_d or (c.get("data_contratto") or "") >= limite_d or (c.get("updated_at") or "") >= limite:
            continue
        if await db.servizi.count_documents({"client_id": c["id"], "$or": [
                {"stato": {"$in": STATI_RIP_APERTI}}, {"created_at": {"$gte": limite}}, {"updated_at": {"$gte": limite}},
                {"data_uscita": {"$gte": limite_d}}, {"data_attivazione": {"$gte": limite_d}}]}):
            continue
        if await db.ritiri.count_documents({"client_id": c["id"], "created_at": {"$gte": limite}}):
            continue
        await anonimizza_cliente_core(c, "retention-automatica")
        report.append(c["id"])
    await db.cron_log.insert_one({"job": "retention-gdpr", "at": datetime.now(timezone.utc).isoformat(), "anonimizzati": len(report), "ids": report})
    return {"anonimizzati": len(report)}

@api_router.get("/admin/retention-anteprima")
async def retention_anteprima(admin: dict = Depends(require_admin)):
    limite = (datetime.now(timezone.utc) - relativedelta(years=RETENTION_ANNI)).isoformat()
    n = await db.clients.count_documents({"anonimizzato_at": {"$in": [None]}, "created_at": {"$lt": limite, "$gte": "2000-01-01"}})
    date_anomale = await db.clients.count_documents({"created_at": {"$lt": "2000-01-01"}})
    ultimo = await db.cron_log.find_one({"job": "retention-gdpr"}, {"_id": 0}, sort=[("at", -1)])
    return {"anni": RETENTION_ANNI, "candidati_per_data_registrazione": n, "date_anomale_ignorate": date_anomale, "ultima_esecuzione": ultimo}

@api_router.get("/dashboard/margini-12-mesi")
async def margini_12_mesi(admin: dict = Depends(require_admin)):
    oggi = date.today().replace(day=1)
    mesi = [(oggi - relativedelta(months=i)).strftime("%Y-%m") for i in range(11, -1, -1)]
    rips = await db.servizi.find({"tipo": "riparazione", "stato": {"$in": ["consegnato", "pronto", "in_lavorazione"]}}, {"_id": 0}).to_list(50000)
    stores = {s["id"]: s["nome"] for s in await db.stores.find({}, {"_id": 0, "id": 1, "nome": 1}).to_list(200)}
    rows = {m: {"mese": m, "label": m[5:] + "/" + m[2:4]} for m in mesi}
    usati = set()
    for s in rips:
        rif = (s.get("data_uscita") or s.get("created_at") or "")[:7]
        if rif not in rows:
            continue
        costi = costi_riparazione(s) or {"totale": 0}
        prezzo = float(s.get("prezzo_finale") if s.get("prezzo_finale") is not None else (calcola_prezzo_riparazione(s) or 0))
        nome = stores.get(s.get("venditore_id"), "Altro")
        usati.add(nome)
        rows[rif][nome] = round(rows[rif].get(nome, 0.0) + prezzo / 1.22 - costi["totale"], 2)
    return {"negozi": sorted(usati), "dati": [rows[m] for m in mesi]}

async def get_scoped_client(client_id: str, user: dict) -> dict:
    scope = client_scope_filter(user)
    scope["id"] = client_id
    c = await db.clients.find_one(scope, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    return c

@api_router.post("/clients/{client_id}/attachments")
async def upload_attachment(client_id: str, file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    await get_scoped_client(client_id, user)
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else "bin"
    if ext not in ("pdf", "jpg", "jpeg", "png", "webp"):
        raise HTTPException(status_code=400, detail="Formato non supportato (PDF, JPG, PNG, WEBP)")
    data = await file.read()
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File troppo grande (max 15 MB)")
    path = f"{APP_NAME}/attachments/{client_id}/{uuid.uuid4()}.{ext}"
    result = put_object(path, data, file.content_type or "application/octet-stream")
    doc = {"id": str(uuid.uuid4()), "client_id": client_id, "storage_path": result["path"],
           "original_filename": file.filename, "content_type": file.content_type,
           "size": result["size"], "is_deleted": False, "uploaded_by": user["id"],
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.attachments.insert_one(doc)
    doc.pop("_id", None)
    return doc

@api_router.get("/clients/{client_id}/attachments")
async def list_attachments(client_id: str, user: dict = Depends(get_current_user)):
    await get_scoped_client(client_id, user)
    return await db.attachments.find({"client_id": client_id, "is_deleted": False}, {"_id": 0}).to_list(200)

async def _check_attachment_scope(rec: dict, user: dict):
    if rec.get("client_id"):
        await get_scoped_client(rec["client_id"], user)
    else:
        await get_scoped_servizio(rec["servizio_id"], user)

@api_router.get("/attachments/{att_id}/download")
async def download_attachment(att_id: str, user: dict = Depends(get_current_user)):
    rec = await db.attachments.find_one({"id": att_id, "is_deleted": False}, {"_id": 0})
    if not rec:
        raise HTTPException(status_code=404, detail="Allegato non trovato")
    await _check_attachment_scope(rec, user)
    data, content_type = get_object(rec["storage_path"])
    return Response(content=data, media_type=rec.get("content_type", content_type),
                    headers={"Content-Disposition": f'attachment; filename="{rec["original_filename"]}"'})

@api_router.delete("/attachments/{att_id}")
async def delete_attachment(att_id: str, user: dict = Depends(get_current_user)):
    rec = await db.attachments.find_one({"id": att_id, "is_deleted": False}, {"_id": 0})
    if not rec:
        raise HTTPException(status_code=404, detail="Allegato non trovato")
    await _check_attachment_scope(rec, user)
    await db.attachments.update_one({"id": att_id}, {"$set": {"is_deleted": True}})
    return {"status": "ok"}

@api_router.get("/clients/{client_id}/attachments/merged")
async def merged_pdf(client_id: str, user: dict = Depends(get_current_user)):
    from pypdf import PdfWriter, PdfReader
    client_doc = await get_scoped_client(client_id, user)
    atts = await db.attachments.find({"client_id": client_id, "is_deleted": False,
                                      "content_type": "application/pdf"}, {"_id": 0}).to_list(200)
    if not atts:
        raise HTTPException(status_code=400, detail="Nessun PDF tra gli allegati di questo cliente")
    writer = PdfWriter()
    for att in atts:
        data, _ = get_object(att["storage_path"])
        reader = PdfReader(io.BytesIO(data))
        for page in reader.pages:
            writer.add_page(page)
    buf = io.BytesIO()
    writer.write(buf)
    buf.seek(0)
    fname = f"documenti_{client_doc.get('cognome', 'cliente')}_{client_doc.get('nome', '')}.pdf".replace(" ", "_")
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})

# ---------------- Numerazione automatica, Magazzino, Ritiri, PDF ----------------

RITIRO_SEEDS = {"morbegno": ("M", 8), "grosio": ("GR", 1), "sondrio": ("SO", 7),
                "gravedona": ("G", 3), "tirano": ("T", 1), "sondalo": ("SA", 1)}

async def next_numero(kind: str, store_id: str, default_prefix: str = "RIP") -> str:
    store = await db.stores.find_one({"id": store_id}, {"_id": 0, "nome": 1})
    sname = (store or {}).get("nome", "").lower()
    prefix, start = default_prefix, 1
    if kind == "ritiro":
        for key, (p, s) in RITIRO_SEEDS.items():
            if key in sname:
                prefix, start = p, s
                break
    ckey = f"{kind}:{store_id}"
    existing = await db.counters.find_one({"_id": ckey})
    if not existing:
        await db.counters.insert_one({"_id": ckey, "prefix": prefix, "seq": start})
        existing = {"prefix": prefix, "seq": start}
    seq = existing["seq"]
    await db.counters.update_one({"_id": ckey}, {"$inc": {"seq": 1}})
    return f"{existing.get('prefix', prefix)}{seq}"

MAGAZZINO_CATEGORIE = ["display", "ricambi", "accessori", "sim", "rigenerati", "altro"]

class MagazzinoInput(BaseModel):
    nome: str
    categoria: str = "altro"
    store_id: str = ""
    quantita: int = 0
    prezzo_acquisto: Optional[float] = None
    prezzo_vendita: Optional[float] = None
    note: str = ""

def magazzino_scope(user: dict) -> dict:
    if user["role"] in ("admin", "tecnico") or user.get("can_view_all"):
        return {}
    return {"store_id": {"$in": user.get("store_ids", [])}}

@api_router.get("/magazzino")
async def list_magazzino(user: dict = Depends(get_current_user), categoria: str = "",
                         venditore_id: str = "", q: str = ""):
    scope = magazzino_scope(user)
    if categoria:
        scope["categoria"] = categoria
    if venditore_id and (user["role"] in ("admin", "tecnico") or user.get("can_view_all")):
        scope["store_id"] = venditore_id
    if q:
        scope["nome"] = {"$regex": re.escape(q), "$options": "i"}
    items = await db.magazzino.find(scope, {"_id": 0}).sort("nome", 1).to_list(5000)
    store_ids = list({i.get("store_id", "") for i in items if i.get("store_id")})
    stores = await db.stores.find({"id": {"$in": store_ids}}, {"_id": 0, "id": 1, "nome": 1}).to_list(200)
    smap = {s["id"]: s["nome"] for s in stores}
    for i in items:
        i["store_name"] = smap.get(i.get("store_id", ""), "-")
    return items

@api_router.get("/magazzino/disponibilita")
async def disponibilita_magazzino(user: dict = Depends(get_current_user), q: str = "", categoria: str = ""):
    scope = {}
    if q:
        scope["nome"] = {"$regex": re.escape(q), "$options": "i"}
    if categoria:
        scope["categoria"] = categoria
    items = await db.magazzino.find(scope, {"_id": 0}).to_list(5000)
    store_ids = list({i.get("store_id", "") for i in items if i.get("store_id")})
    stores = await db.stores.find({"id": {"$in": store_ids}}, {"_id": 0, "id": 1, "nome": 1}).to_list(200)
    smap = {s["id"]: s["nome"] for s in stores}
    grouped = {}
    for i in items:
        key = (i.get("nome", "").strip().lower(), i.get("categoria", ""))
        g = grouped.setdefault(key, {"nome": i.get("nome", ""), "categoria": i.get("categoria", ""), "stores": []})
        can_price = user["role"] in ("admin", "tecnico") or user.get("can_view_all") \
            or i.get("store_id") in user.get("store_ids", [])
        g["stores"].append({"item_id": i["id"], "store_id": i.get("store_id", ""),
                            "store_name": smap.get(i.get("store_id", ""), "-"),
                            "quantita": i.get("quantita", 0),
                            "prezzo_vendita": i.get("prezzo_vendita") if can_price else None})
    return sorted(grouped.values(), key=lambda x: x["nome"])

@api_router.post("/magazzino")
async def create_magazzino(input: MagazzinoInput, user: dict = Depends(get_current_user)):
    data = input.model_dump()
    if data["categoria"] not in MAGAZZINO_CATEGORIE:
        raise HTTPException(status_code=400, detail="Categoria non valida")
    if user["role"] == "negozio" and user.get("store_ids"):
        data["store_id"] = user["store_ids"][0]
    if not data["store_id"]:
        raise HTTPException(status_code=400, detail="Negozio obbligatorio")
    now = datetime.now(timezone.utc).isoformat()
    data.update({"id": str(uuid.uuid4()), "created_by": user["id"], "created_at": now, "updated_at": now})
    await db.magazzino.insert_one(data)
    data.pop("_id", None)
    return data

@api_router.patch("/magazzino/{item_id}")
async def update_magazzino(item_id: str, input: MagazzinoInput, user: dict = Depends(get_current_user)):
    scope = magazzino_scope(user)
    scope["id"] = item_id
    old = await db.magazzino.find_one(scope, {"_id": 0})
    if not old:
        raise HTTPException(status_code=404, detail="Articolo non trovato")
    data = input.model_dump(exclude_unset=True)
    if user["role"] == "negozio":
        data.pop("store_id", None)
    data["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.magazzino.update_one({"id": item_id}, {"$set": data})
    return await db.magazzino.find_one({"id": item_id}, {"_id": 0})

@api_router.delete("/magazzino/{item_id}")
async def delete_magazzino(item_id: str, admin: dict = Depends(require_admin)):
    res = await db.magazzino.delete_one({"id": item_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Articolo non trovato")
    return {"status": "ok"}

class MovimentoInput(BaseModel):
    delta: int
    motivo: str = ""

class VenditaRigeneratoInput(BaseModel):
    prezzo_vendita: float
    note: str = ""

@api_router.post("/magazzino/{item_id}/vendi")
async def vendi_rigenerato(item_id: str, input: VenditaRigeneratoInput, user: dict = Depends(get_current_user)):
    scope = magazzino_scope(user)
    scope["id"] = item_id
    item = await db.magazzino.find_one(scope, {"_id": 0})
    if not item:
        raise HTTPException(status_code=404, detail="Articolo non trovato")
    if item.get("categoria") != "rigenerati":
        raise HTTPException(status_code=400, detail="Vendita tracciata solo per dispositivi rigenerati")
    if (item.get("quantita") or 0) < 1:
        raise HTTPException(status_code=400, detail="Dispositivo già venduto / giacenza zero")
    if input.prezzo_vendita < 0:
        raise HTTPException(status_code=400, detail="Prezzo non valido")
    costo = float(item.get("prezzo_acquisto") or 0)
    now = datetime.now(timezone.utc).isoformat()
    vendita = {"id": str(uuid.uuid4()), "magazzino_id": item["id"], "store_id": item.get("store_id", ""),
               "nome": item.get("nome", ""), "imei": item.get("imei", ""), "ritiro_id": item.get("ritiro_id", ""),
               "ritiro_numero": item.get("ritiro_numero", ""), "servizio_id": item.get("servizio_id", ""),
               "riparazione_numero": item.get("riparazione_numero", ""), "costo": round(costo, 2),
               "prezzo_vendita": round(input.prezzo_vendita, 2), "margine": round(input.prezzo_vendita / 1.22 - costo, 2),
               "note": input.note, "venduto_da": user["id"], "venduto_da_nome": user["name"], "venduto_at": now}
    await db.vendite_rigenerati.insert_one(vendita)
    await db.magazzino.update_one({"id": item_id}, {"$inc": {"quantita": -1},
                                                   "$set": {"venduto_at": now, "prezzo_venduto": vendita["prezzo_vendita"], "updated_at": now}})
    vendita.pop("_id", None)
    return vendita

@api_router.get("/rigenerati/venduti")
async def rigenerati_venduti(user: dict = Depends(get_current_user), mese: str = ""):
    scope = magazzino_scope(user)
    if mese:
        scope["venduto_at"] = {"$regex": f"^{re.escape(mese)}"}
    rows = await db.vendite_rigenerati.find(scope, {"_id": 0}).sort("venduto_at", -1).to_list(2000)
    stores = {s["id"]: s["nome"] for s in await db.stores.find({}, {"_id": 0, "id": 1, "nome": 1}).to_list(200)}
    for v in rows:
        v["store_name"] = stores.get(v.get("store_id", ""), "-")
    return {"vendite": rows, "totale_margine": round(sum(v["margine"] for v in rows), 2),
            "totale_vendite": round(sum(v["prezzo_vendita"] for v in rows), 2)}

@api_router.post("/magazzino/{item_id}/movimento")
async def movimento_magazzino(item_id: str, input: MovimentoInput, user: dict = Depends(get_current_user)):
    scope = magazzino_scope(user)
    scope["id"] = item_id
    item = await db.magazzino.find_one(scope, {"_id": 0})
    if not item:
        raise HTTPException(status_code=404, detail="Articolo non trovato")
    new_q = (item.get("quantita") or 0) + input.delta
    if new_q < 0:
        raise HTTPException(status_code=400, detail="Giacenza insufficiente")
    await db.magazzino.update_one({"id": item_id},
                                  {"$set": {"quantita": new_q, "updated_at": datetime.now(timezone.utc).isoformat()}})
    return {"status": "ok", "quantita": new_q}

def _xlsx_response(ws_title: str, headers: list, rows: list, fname: str):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = ws_title
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 3
        ws.column_dimensions[col[0].column_letter].width = min(width, 45)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf,
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})

@api_router.get("/magazzino/export")
async def export_magazzino(user: dict = Depends(get_current_user)):
    items = await db.magazzino.find(magazzino_scope(user), {"_id": 0}).sort("nome", 1).to_list(5000)
    store_ids = list({i.get("store_id", "") for i in items if i.get("store_id")})
    stores = await db.stores.find({"id": {"$in": store_ids}}, {"_id": 0, "id": 1, "nome": 1}).to_list(200)
    smap = {s["id"]: s["nome"] for s in stores}
    rows = [[i.get("nome", ""), i.get("categoria", ""), smap.get(i.get("store_id", ""), "-"),
             i.get("quantita", 0), i.get("prezzo_acquisto"), i.get("prezzo_vendita"), i.get("note", "")] for i in items]
    return _xlsx_response("Magazzino", ["Articolo", "Categoria", "Negozio", "Giacenza", "Prezzo acquisto", "Prezzo vendita", "Note"],
                          rows, f"magazzino_{date.today().isoformat()}.xlsx")

@api_router.get("/ritiri/export")
async def export_ritiri(user: dict = Depends(get_current_user)):
    scope = ritiri_scope(user)
    ritiri = await db.ritiri.find(scope, {"_id": 0}).sort("created_at", -1).to_list(5000)
    store_ids = list({r.get("store_id", "") for r in ritiri if r.get("store_id")})
    stores = await db.stores.find({"id": {"$in": store_ids}}, {"_id": 0, "id": 1, "nome": 1}).to_list(200)
    smap = {s["id"]: s["nome"] for s in stores}
    rows = [[r.get("numero", ""), _fmt_it(r.get("data_ritiro")), r.get("cognome", ""), r.get("nome", ""),
             r.get("codice_fiscale", ""), r.get("articolo", ""), r.get("imei", ""), r.get("prezzo_ritiro"),
             r.get("numero_documento", ""), r.get("n_allegati", ""), smap.get(r.get("store_id", ""), "-")] for r in ritiri]
    return _xlsx_response("Ritiri", ["N. ritiro", "Data", "Cognome", "Nome", "Codice fiscale", "Articolo", "IMEI",
                                     "Prezzo ritiro", "N. documento", "Allegati", "Negozio"],
                          rows, f"ritiri_usato_{date.today().isoformat()}.xlsx")

class RicambioUsoInput(BaseModel):
    magazzino_id: str
    quantita: int = 1
    prezzo_manuale: Optional[float] = None

@api_router.post("/servizi/{servizio_id}/ricambi")
async def usa_ricambio(servizio_id: str, input: RicambioUsoInput, user: dict = Depends(get_current_user)):
    svc = await get_scoped_servizio(servizio_id, user)
    if svc.get("tipo") != "riparazione":
        raise HTTPException(status_code=400, detail="Ricambi solo per riparazioni")
    item = await db.magazzino.find_one({"id": input.magazzino_id}, {"_id": 0})
    if not item:
        raise HTTPException(status_code=404, detail="Articolo non trovato in magazzino")
    if item.get("store_id") != svc.get("venditore_id"):
        raise HTTPException(status_code=400, detail="Ricambio di un altro negozio: fai prima uno spostamento di giacenza")
    if input.quantita < 1:
        raise HTTPException(status_code=400, detail="Quantità non valida")
    if (item.get("quantita") or 0) < input.quantita:
        raise HTTPException(status_code=400, detail="Giacenza insufficiente")
    now = datetime.now(timezone.utc).isoformat()
    await db.magazzino.update_one({"id": item["id"]},
                                  {"$inc": {"quantita": -input.quantita}, "$set": {"updated_at": now}})
    prezzo = input.prezzo_manuale if input.prezzo_manuale is not None else item.get("prezzo_vendita")
    uso = {"item_id": item["id"], "nome": item["nome"], "quantita": input.quantita,
           "prezzo_vendita": prezzo, "at": now}
    await db.servizi.update_one({"id": servizio_id},
                                {"$push": {"ricambi_usati": uso}, "$set": {"updated_at": now}})
    return {"status": "ok", "giacenza": (item.get("quantita") or 0) - input.quantita}

@api_router.delete("/servizi/{servizio_id}/ricambi/{index}")
async def annulla_ricambio(servizio_id: str, index: int, user: dict = Depends(get_current_user)):
    svc = await get_scoped_servizio(servizio_id, user)
    usati = list(svc.get("ricambi_usati") or [])
    if index < 0 or index >= len(usati):
        raise HTTPException(status_code=404, detail="Ricambio non trovato")
    uso = usati[index]
    await db.magazzino.update_one({"id": uso["item_id"]},
                                  {"$inc": {"quantita": uso["quantita"]},
                                   "$set": {"updated_at": datetime.now(timezone.utc).isoformat()}})
    usati.pop(index)
    await db.servizi.update_one({"id": servizio_id},
                                {"$set": {"ricambi_usati": usati,
                                          "updated_at": datetime.now(timezone.utc).isoformat()}})
    return {"status": "ok"}

def _fmt_it(iso: str) -> str:
    try:
        return datetime.fromisoformat(str(iso).replace("Z", "+00:00")).strftime("%d/%m/%Y")
    except (ValueError, TypeError):
        return str(iso or "-")

def _new_pdf(title: str, subtitle: str = ""):
    from fpdf import FPDF
    pdf = FPDF(format="A4", unit="mm")
    pdf.add_page()
    logo = "/app/frontend/public/rs-logo.png"
    if os.path.exists(logo):
        pdf.image(logo, x=(210 - 34) / 2, w=34)
        pdf.ln(2)
    pdf.set_font("helvetica", "B", 17)
    pdf.cell(0, 10, title, align="C", new_x="LMARGIN", new_y="NEXT")
    if subtitle:
        pdf.set_font("helvetica", "", 11)
        pdf.cell(0, 7, subtitle, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    return pdf

def _pdf_field(pdf, label: str, value: str):
    pdf.set_font("helvetica", "B", 11)
    pdf.cell(55, 9, label)
    pdf.set_font("helvetica", "", 11)
    pdf.cell(0, 9, str(value or "-"), new_x="LMARGIN", new_y="NEXT")

def _pdf_multiline(pdf, label: str, value: str):
    pdf.set_font("helvetica", "B", 11)
    pdf.cell(0, 8, label, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 11)
    pdf.multi_cell(0, 7, str(value or "-"))
    pdf.ln(1)

def _build_scheda_pdf(svc: dict, client_doc: dict) -> bytes:
    pdf = _new_pdf("Scheda riparazione",
                   f"N. {svc.get('numero_riparazione') or '-'}  -  Data: {_fmt_it(svc.get('created_at'))}")
    tel = (client_doc or {}).get("telefono", "")
    _pdf_field(pdf, "Cliente", svc.get("client_name") or f"{(client_doc or {}).get('cognome', '')} {(client_doc or {}).get('nome', '')}".strip())
    _pdf_field(pdf, "Telefono", tel)
    _pdf_field(pdf, "Dispositivo", svc.get("dispositivo"))
    sblocco = ""
    segreti = leggi_segreti(svc)
    if svc.get("codice_sblocco_tipo") and svc.get("codice_sblocco_tipo") != "nessuno":
        sblocco = f"{svc['codice_sblocco_tipo']}: {segreti.get('codice_sblocco') or '-'}"
    _pdf_field(pdf, "Codice sblocco", sblocco or "Nessuno")
    pdf.ln(2)
    _pdf_multiline(pdf, "Problema riscontrato", svc.get("problema"))
    _pdf_multiline(pdf, "Operazioni svolte", svc.get("operazioni"))
    usati = svc.get("ricambi_usati") or []
    pdf.set_font("helvetica", "B", 11)
    pdf.cell(0, 8, "Ricambi utilizzati", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 11)
    if usati:
        for u in usati:
            pdf.cell(0, 7, f"- {u.get('nome', '')} x{u.get('quantita', 1)}", new_x="LMARGIN", new_y="NEXT")
    else:
        pdf.cell(0, 7, "Nessun ricambio", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    if svc.get("prezzo_finale") is not None:
        _pdf_field(pdf, "Prezzo", f"EUR {float(svc['prezzo_finale']):.2f} (IVA inclusa)")
    elif svc.get("prezzo_consigliato") is not None:
        _pdf_field(pdf, "Prezzo", f"EUR {svc['prezzo_consigliato']:.2f} (IVA inclusa)")
    pdf.ln(12)
    pdf.set_font("helvetica", "", 11)
    pdf.cell(95, 8, "Firma tecnico: ______________________")
    pdf.cell(0, 8, "Firma cliente: ______________________", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())

@api_router.get("/servizi/{servizio_id}/scheda")
async def scheda_riparazione(servizio_id: str, user: dict = Depends(get_current_user)):
    svc = await get_scoped_servizio(servizio_id, user)
    if svc.get("tipo") != "riparazione":
        raise HTTPException(status_code=400, detail="Scheda disponibile solo per riparazioni")
    client_doc = await db.clients.find_one({"id": svc["client_id"]}, {"_id": 0})
    names = await clients_name_map([svc["client_id"]])
    svc["client_name"] = names.get(svc["client_id"], "")
    svc["prezzo_consigliato"] = calcola_prezzo_riparazione(svc)
    buf = io.BytesIO(_build_scheda_pdf(svc, client_doc or {}))
    fname = f"scheda_{svc.get('numero_riparazione') or servizio_id}.pdf".replace(" ", "_")
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})

def _build_bolla_pdf(r: dict) -> bytes:
    pdf = _new_pdf("Bolla di ritiro articoli usati",
                   f"N. {r.get('numero', '-')}  -  Data ritiro: {_fmt_it(r.get('data_ritiro'))}")
    pdf.ln(4)
    pdf.set_font("helvetica", "B", 10); pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 6, "DATI CLIENTE", new_x="LMARGIN", new_y="NEXT"); pdf.set_text_color(0, 0, 0)
    _pdf_field(pdf, "Nome", r.get("nome"))
    _pdf_field(pdf, "Cognome", r.get("cognome"))
    _pdf_field(pdf, "Codice Fiscale", r.get("codice_fiscale"))
    _pdf_field(pdf, "Numero documento", r.get("numero_documento"))
    pdf.ln(3)
    pdf.set_font("helvetica", "B", 10); pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 6, "ARTICOLO RITIRATO", new_x="LMARGIN", new_y="NEXT"); pdf.set_text_color(0, 0, 0)
    if r.get("marca") or r.get("modello"):
        _pdf_field(pdf, "Marca", r.get("marca"))
        _pdf_field(pdf, "Modello", r.get("modello"))
    else:
        _pdf_field(pdf, "Articolo (marca e modello)", r.get("articolo"))
    _pdf_field(pdf, "IMEI / Serial", r.get("imei"))
    prezzo = f"EUR {r['prezzo_ritiro']:.2f}" if r.get("prezzo_ritiro") is not None else "-"
    _pdf_field(pdf, "Prezzo ritiro", prezzo)
    _pdf_field(pdf, "Data ritiro", _fmt_it(r.get("data_ritiro")))
    _pdf_field(pdf, "Si allegano documenti n.", str(r.get("n_allegati", 2)))
    if r.get("riparazione_numero"):
        pdf.ln(3)
        _pdf_field(pdf, "Riparazione collegata", f"N. {r['riparazione_numero']}  -  {r.get('riparazione_dispositivo', '')}")
    pdf.ln(4)
    pdf.set_font("helvetica", "", 9)
    pdf.multi_cell(0, 5, "Il cliente dichiara di essere il legittimo proprietario dell'articolo sopra descritto, di cederlo liberamente a RS Riparazioni "
                         "al prezzo indicato e di aver provveduto al salvataggio e alla rimozione dei propri dati personali (o di autorizzarne la cancellazione). "
                         "Copia del documento di identita' e' allegata alla presente.", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(12)
    pdf.set_font("helvetica", "", 11)
    pdf.cell(95, 8, f"Luogo e data: ____________________, {_fmt_it(r.get('data_ritiro'))}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(10)
    pdf.cell(95, 8, "Firma negozio")
    pdf.cell(0, 8, "Firma cliente", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(95, 8, "_______________________________")
    pdf.cell(0, 8, "_______________________________", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())

class RitiroInput(BaseModel):
    store_id: str = ""
    client_id: str = ""
    nome: str
    cognome: str
    codice_fiscale: str = ""
    articolo: str
    marca: str = ""
    modello: str = ""
    imei: str = ""
    prezzo_ritiro: Optional[float] = None
    numero_documento: str = ""
    n_allegati: int = 2
    data_ritiro: Optional[str] = None
    servizio_id: str = ""
    costo_ricambi: Optional[float] = None  # costo a nostro carico (es. display), NON in bolla
    crea_rigenerato: bool = True

def ritiri_scope(user: dict) -> dict:
    if "riparazioni" not in user_sections(user):
        raise HTTPException(status_code=403, detail="Sezione non abilitata per questo utente")
    if user["role"] == "admin" or user.get("can_view_all"):
        return {}
    return {"store_id": {"$in": user.get("store_ids", [])}}

@api_router.get("/ritiri")
async def list_ritiri(user: dict = Depends(get_current_user), venditore_id: str = "", q: str = ""):
    scope = ritiri_scope(user)
    if venditore_id and (user["role"] == "admin" or user.get("can_view_all")):
        scope["store_id"] = venditore_id
    ritiri = await db.ritiri.find(scope, {"_id": 0}).sort("created_at", -1).to_list(5000)
    if q:
        ql = q.lower()
        ritiri = [r for r in ritiri if ql in f"{r.get('cognome', '')} {r.get('nome', '')}".lower()
                  or ql in r.get("articolo", "").lower() or ql in r.get("numero", "").lower()]
    store_ids = list({r.get("store_id", "") for r in ritiri if r.get("store_id")})
    stores = await db.stores.find({"id": {"$in": store_ids}}, {"_id": 0, "id": 1, "nome": 1}).to_list(200)
    smap = {s["id"]: s["nome"] for s in stores}
    for r in ritiri:
        r["store_name"] = smap.get(r.get("store_id", ""), "-")
    return ritiri

@api_router.post("/ritiri")
async def create_ritiro(input: RitiroInput, user: dict = Depends(get_current_user)):
    ritiri_scope(user)
    data = input.model_dump()
    if user["role"] == "negozio" and user.get("store_ids"):
        data["store_id"] = user["store_ids"][0]
    if not data["store_id"]:
        raise HTTPException(status_code=400, detail="Negozio obbligatorio")
    if not data["nome"].strip() or not data["cognome"].strip() or not data["articolo"].strip():
        raise HTTPException(status_code=400, detail="Nome, cognome e articolo sono obbligatori")
    data["numero"] = await next_numero("ritiro", data["store_id"])
    if not data.get("data_ritiro"):
        data["data_ritiro"] = date.today().isoformat()
    now = datetime.now(timezone.utc).isoformat()
    data.update({"id": str(uuid.uuid4()), "created_by": user["id"], "created_at": now})
    if data.get("servizio_id"):
        svc = await db.servizi.find_one({"id": data["servizio_id"]}, {"_id": 0, "numero_riparazione": 1, "dispositivo": 1})
        if svc:
            data["riparazione_numero"] = svc.get("numero_riparazione", "")
            data["riparazione_dispositivo"] = svc.get("dispositivo", "")
    pdf_bytes = _build_bolla_pdf(data)
    result = put_object(f"{APP_NAME}/ritiri/{data['numero']}.pdf", pdf_bytes, "application/pdf")
    data["storage_path"] = result["path"]
    data["bolla_base_path"] = result["path"]
    data["documenti"] = []
    if data.get("crea_rigenerato"):
        costo = float(data.get("prezzo_ritiro") or 0) + float(data.get("costo_ricambi") or 0)
        item = {"id": str(uuid.uuid4()), "nome": data["articolo"], "categoria": "rigenerati", "store_id": data["store_id"],
                "quantita": 1, "prezzo_acquisto": round(costo, 2), "prezzo_vendita": None, "imei": data.get("imei", ""),
                "ritiro_id": data["id"], "ritiro_numero": data["numero"], "servizio_id": data.get("servizio_id", ""),
                "riparazione_numero": data.get("riparazione_numero", ""),
                "note": f"Ritiro {data['numero']} - costo ritiro {float(data.get('prezzo_ritiro') or 0):.2f} + ricambi {float(data.get('costo_ricambi') or 0):.2f}",
                "created_by": user["id"], "created_at": now, "updated_at": now}
        await db.magazzino.insert_one(item)
        data["magazzino_id"] = item["id"]
    await db.ritiri.insert_one(data)
    data.pop("_id", None)
    if data.get("servizio_id"):
        await db.servizi.update_one({"id": data["servizio_id"]},
                                    {"$set": {"ritiro_id": data["id"], "ritiro_numero": data["numero"],
                                              "updated_at": now}})
    return data

RITIRO_DOC_TYPES = {"application/pdf", "image/jpeg", "image/png", "image/webp"}

def _merge_bolla_documenti(bolla: bytes, docs: list) -> bytes:
    """Unisce la bolla con i documenti (PDF o immagini) in un unico PDF."""
    out = fitz.open("pdf", bolla)
    for content_type, data in docs:
        if content_type == "application/pdf":
            src = fitz.open("pdf", data)
        else:
            img = fitz.open(stream=data, filetype=content_type.split("/")[1])
            src = fitz.open("pdf", img.convert_to_pdf())
        out.insert_pdf(src)
    return out.tobytes()

@api_router.post("/ritiri/{ritiro_id}/documenti")
async def upload_ritiro_documenti(ritiro_id: str, files: List[UploadFile] = File(...), user: dict = Depends(get_current_user)):
    scope = ritiri_scope(user)
    scope["id"] = ritiro_id
    r = await db.ritiri.find_one(scope, {"_id": 0})
    if not r:
        raise HTTPException(status_code=404, detail="Ritiro non trovato")
    docs, names = [], []
    for f in files:
        ct = (f.content_type or "").lower()
        if ct not in RITIRO_DOC_TYPES:
            raise HTTPException(status_code=400, detail=f"Formato non supportato: {f.filename} (usa PDF, JPG, PNG o WEBP)")
        content = await f.read()
        if len(content) > 15 * 1024 * 1024:
            raise HTTPException(status_code=400, detail=f"File troppo grande: {f.filename} (max 15 MB)")
        docs.append((ct, content))
        names.append(f.filename)
    if not docs:
        raise HTTPException(status_code=400, detail="Nessun file")
    current, _ = get_object(r["storage_path"])
    merged = _merge_bolla_documenti(current, docs)
    version = len(r.get("documenti") or []) + 1
    result = put_object(f"{APP_NAME}/ritiri/{r['numero']}_v{version}.pdf", merged, "application/pdf")
    now = datetime.now(timezone.utc).isoformat()
    await db.ritiri.update_one({"id": ritiro_id}, {
        "$set": {"storage_path": result["path"], "updated_at": now},
        "$push": {"documenti": {"$each": [{"nome": n, "at": now, "by": user["id"]} for n in names]}}})
    return {"status": "ok", "documenti": (r.get("documenti") or []) + [{"nome": n, "at": now} for n in names]}

@api_router.get("/ritiri/{ritiro_id}/pdf")
async def download_bolla(ritiro_id: str, user: dict = Depends(get_current_user)):
    scope = ritiri_scope(user)
    scope["id"] = ritiro_id
    r = await db.ritiri.find_one(scope, {"_id": 0})
    if not r:
        raise HTTPException(status_code=404, detail="Ritiro non trovato")
    data, _ = get_object(r["storage_path"])
    return Response(content=data, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{r["numero"]}.pdf"'})

@api_router.delete("/ritiri/{ritiro_id}")
async def delete_ritiro(ritiro_id: str, admin: dict = Depends(require_admin)):
    res = await db.ritiri.delete_one({"id": ritiro_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Ritiro non trovato")
    await db.servizi.update_many({"ritiro_id": ritiro_id}, {"$unset": {"ritiro_id": "", "ritiro_numero": ""}})
    return {"status": "ok"}

# ---------------- WhatsApp (Baileys service) ----------------

WA_SERVICE = os.environ.get("WA_SERVICE_URL", "http://127.0.0.1:3001")
WA_SERVICE_KEY = os.environ.get("WA_SERVICE_KEY", "")

def wa_headers() -> dict:
    return {"X-API-Key": WA_SERVICE_KEY} if WA_SERVICE_KEY else {}

PRIVACY_MSG = """RS Group – Grazie per averci scelto!
Ciao {nome}
per procedere con la tua richiesta e completare l'attivazione del servizio, è necessario firmare l'autorizzazione privacy.
Puoi farlo in modo semplice e veloce al link qui sotto:
https://rsriparazioni.it/privacy/
La firma è richiesta per attivare correttamente il servizio (riparazioni, contratti energia o telefonia).
Grazie per la fiducia
RS Group"""

REVIEW_MSG = """Ciao!
Grazie per aver scelto CAMBIAORA
Se ti sei trovato bene, ci farebbe davvero piacere una tua recensione
Basta un clic qui
https://g.page/r/CaSv3O7luiBPEAE/review
Grazie per il supporto"""

# Canali WhatsApp energia: CambiaOra sempre dal numero Sondrio (3519460591); ENEL da Deborah o dal negozio Gravedona
CAMBIAORA_STORE_NAME = "Sondrio"
ENEL_GRAVEDONA_STORE_NAME = "Gravedona"
ENEL_SESSION_DEBORAH = "enel-deborah"
ENEL_REVIEW_MSG = {
    ENEL_SESSION_DEBORAH: """Ciao! 😊 Grazie per aver scelto ENEL Morbegno 🙌
Se ti sei trovato bene, ci farebbe davvero piacere una tua recensione ⭐ Basta un clic qui 👇
👉 https://g.page/r/CT15UUSXSQoCEAE/review
Grazie per il supporto""",
    "gravedona": """Ciao! 😊 Grazie per aver scelto ENEL Gravedona 🙌
Se ti sei trovato bene, ci farebbe davvero piacere una tua recensione ⭐ Basta un clic qui 👇
https://g.page/r/CWOpR0o29GE-EAE/review
Grazie per il supporto 💙""",
}
ENEL_TRUFFE_MSG = """Le truffe telefoniche sono sempre più frequenti. Se ricevi una chiamata da chi si presenta come operatore di luce, gas, telefonia, banca o altri servizi...
❌ *Non prendere decisioni di fretta.*
❌ *Non comunicare codici, dati personali o bancari.*
❌ *Non dire "SÌ" senza aver capito con chi stai parlando.*
📞 *Hai già i tuoi consulenti di fiducia.*
Prima di firmare, confermare o accettare qualsiasi proposta, *contatta noi*. Ti diremo gratuitamente se la chiamata è affidabile oppure se potrebbe trattarsi di un tentativo di truffa o di una proposta poco conveniente.
🛡️ Un messaggio o una telefonata possono evitarti problemi e costi inutili.
_Siamo al tuo fianco per aiutarti a scegliere in sicurezza._"""

_STORE_ID_CACHE: dict = {}

async def store_id_by_name(nome: str) -> str:
    if nome not in _STORE_ID_CACHE:
        s = await db.stores.find_one({"nome": {"$regex": f"^{re.escape(nome)}$", "$options": "i"}}, {"_id": 0, "id": 1})
        if not s:
            return ""
        _STORE_ID_CACHE[nome] = s["id"]
    return _STORE_ID_CACHE[nome]

def is_enel(c: dict) -> bool:
    return (c or {}).get("gestione") == "enel"

async def wa_session_cliente(c: dict) -> str:
    """Sessione WhatsApp per i messaggi ENERGIA di un cliente."""
    if is_enel(c):
        grav = await store_id_by_name(ENEL_GRAVEDONA_STORE_NAME)
        return grav if grav and c.get("venditore_id") == grav else ENEL_SESSION_DEBORAH
    return await store_id_by_name(CAMBIAORA_STORE_NAME) or c.get("venditore_id") or "default"

async def review_msg_cliente(c: dict) -> str:
    if not is_enel(c):
        return REVIEW_MSG
    sess = await wa_session_cliente(c)
    return ENEL_REVIEW_MSG["gravedona" if sess != ENEL_SESSION_DEBORAH else ENEL_SESSION_DEBORAH]

async def wa_send(phone: str, message: str, session: str = "default", tipo: str = "", client_id: str = "", servizio_id: str = ""):
    log = {"id": str(uuid.uuid4()), "phone": phone, "session": session, "message": message[:1000], "tipo": tipo,
           "client_id": client_id, "servizio_id": servizio_id,
           "at": datetime.now(timezone.utc).isoformat(), "ok": False, "error": None, "session_used": None}
    if os.environ.get("WA_DRY_RUN") == "1":
        log.update({"ok": True, "session_used": "dry-run", "dry_run": True})
        await db.wa_log.insert_one(log)
        return {"status": "dry-run"}
    try:
        async with httpx.AsyncClient(timeout=30) as http_client:
            resp = await http_client.post(f"{WA_SERVICE}/send", json={"phone": phone, "message": message, "session": session}, headers=wa_headers())
        data = resp.json() if resp.headers.get("content-type", "").startswith("application/json") else {}
        log["ok"] = resp.status_code == 200 and bool(data.get("success"))
        log["session_used"] = data.get("session")
        if not log["ok"]:
            log["error"] = data.get("error", "Servizio WhatsApp non disponibile")
            raise HTTPException(status_code=502, detail=log["error"])
        return data
    except HTTPException as e:
        log["error"] = log["error"] or str(e.detail)
        raise
    except Exception as e:
        log["error"] = str(e)
        raise HTTPException(status_code=502, detail=f"Servizio WhatsApp non raggiungibile: {e}")
    finally:
        await db.wa_log.insert_one(log)

PRIVACY_PROXY = "https://rsriparazioni.com/api/proxy.php"

STORE_TO_SITO = {"tirano": "tirano", "sondalo": "sondalo", "sondrio": "sondrio",
                 "sondrio grosio": "grosio", "gravedona": "gravedona"}

async def register_privacy_site(client: dict, store_name: str, categoria: str = "CambiaOra", dettaglio: str = "CambiaOra") -> dict:
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/126.0.0.0"}
    async with httpx.AsyncClient(timeout=30, headers=headers) as hc:
        r1 = await hc.post(f"{PRIVACY_PROXY}?action=generateOtp", json={"email": client["email"]})
        d1 = r1.json()
        if d1.get("result") != "ok":
            raise HTTPException(status_code=502, detail=f"Generazione OTP fallita: {d1.get('message', 'errore')}")
        reg = {
            "nome": client.get("nome", ""), "cognome": client.get("cognome", ""),
            "codiceFiscale": client.get("codice_fiscale", ""), "telefono": client.get("telefono", ""),
            "email": client["email"],
            "servizioCategoria": categoria, "servizioDettaglio": dettaglio,
            "negozio": STORE_TO_SITO.get(store_name.strip().lower(), ""),
            "privacyAccepted": True, "otp": d1.get("otp", ""),
        }
        r2 = await hc.post(f"{PRIVACY_PROXY}?action=register", json=reg)
        d2 = r2.json()
        if d2.get("result") != "ok":
            raise HTTPException(status_code=502, detail=f"Registrazione privacy fallita: {d2.get('message', 'errore')}")
    return d2

@api_router.get("/whatsapp/log")
async def whatsapp_log(admin: dict = Depends(require_admin)):
    return await db.wa_log.find({}, {"_id": 0}).sort("at", -1).to_list(50)

@api_router.get("/whatsapp/sessions-summary")
async def whatsapp_sessions_summary(admin: dict = Depends(require_admin)):
    totale = await db.stores.count_documents({}) + 2
    try:
        async with httpx.AsyncClient(timeout=5) as http_client:
            resp = await http_client.get(f"{WA_SERVICE}/status", headers=wa_headers())
        sessions = resp.json().get("sessions", [])
        connessi = len([s for s in sessions if s.get("connected")])
        return {"connessi": connessi, "totale": totale}
    except Exception:
        return {"connessi": None, "totale": totale}

def wa_sessions_allowed(user: dict) -> Optional[set]:
    """None = tutte le sessioni (admin / chi vede tutto); altrimenti solo i propri negozi."""
    if user["role"] == "admin" or user.get("can_view_all"):
        return None
    return set(user.get("store_ids", []))

def check_wa_session(user: dict, session: str) -> None:
    allowed = wa_sessions_allowed(user)
    if allowed is not None and session not in allowed:
        raise HTTPException(status_code=403, detail="Puoi collegare solo il numero del tuo negozio")

@api_router.get("/whatsapp/status")
async def whatsapp_status(user: dict = Depends(get_current_user)):
    try:
        async with httpx.AsyncClient(timeout=10) as http_client:
            resp = await http_client.get(f"{WA_SERVICE}/status", headers=wa_headers())
        data = resp.json()
        allowed = wa_sessions_allowed(user)
        if allowed is not None:
            data["sessions"] = [s for s in data.get("sessions", []) if s.get("session") in allowed]
        return data
    except Exception:
        return {"connected": False, "user": None, "has_qr": False, "service_down": True}

@api_router.get("/whatsapp/qr")
async def whatsapp_qr(session: str = "default", user: dict = Depends(get_current_user)):
    check_wa_session(user, session)
    try:
        async with httpx.AsyncClient(timeout=10) as http_client:
            resp = await http_client.get(f"{WA_SERVICE}/qr", params={"session": session}, headers=wa_headers())
        return resp.json()
    except Exception:
        return {"qr": None, "connected": False}

class PairInput(BaseModel):
    phone: str
    session: str = "default"

@api_router.post("/whatsapp/pair")
async def whatsapp_pair(input: PairInput, user: dict = Depends(get_current_user)):
    check_wa_session(user, input.session)
    try:
        async with httpx.AsyncClient(timeout=25) as http_client:
            resp = await http_client.post(f"{WA_SERVICE}/pair", json={"phone": input.phone, "session": input.session}, headers=wa_headers())
        if resp.status_code != 200:
            raise HTTPException(status_code=resp.status_code, detail=resp.json().get("error", "Pairing fallito"))
        return resp.json()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=503, detail="Servizio WhatsApp non raggiungibile")

@api_router.post("/clients/{client_id}/whatsapp/privacy")
async def whatsapp_privacy(client_id: str, user: dict = Depends(get_current_user)):
    c = await get_scoped_client(client_id, user)
    return await invia_privacy_cliente(c)

async def invia_privacy_cliente(c: dict, queue_review: bool = True) -> dict:
    client_id = c["id"]
    if not c.get("telefono"):
        raise HTTPException(status_code=400, detail="Il cliente non ha un numero di telefono")
    registered, reg_error = False, None
    if c.get("email"):
        try:
            store = await db.stores.find_one({"id": c.get("venditore_id")}, {"_id": 0})
            await register_privacy_site(c, (store or {}).get("nome", ""))
            registered = True
        except Exception as e:
            reg_error = getattr(e, "detail", str(e))
            logger.error(f"Registrazione privacy sito fallita per {client_id}: {reg_error}")
    wa_error = None
    store_c = await db.stores.find_one({"id": c.get("venditore_id")}, {"_id": 0})
    sess = await wa_session_cliente(c)
    try:
        await wa_send(c["telefono"], store_msg(store_c, "msg_privacy", nome=c.get("nome", "")), session=sess, tipo="privacy", client_id=client_id)
    except Exception as e:
        wa_error = getattr(e, "detail", str(e))
    if wa_error and not registered:
        raise HTTPException(status_code=502, detail=f"WhatsApp: {wa_error}")
    now = datetime.now(timezone.utc)
    updates = {}
    if not wa_error:
        updates["privacy_msg_sent_at"] = now.isoformat()
        if queue_review and not c.get("no_recensioni"):
            await db.whatsapp_queue.insert_one({
            "id": str(uuid.uuid4()), "client_id": client_id, "phone": c["telefono"],
            "type": "review", "message": await review_msg_cliente(c), "send_after": (now + timedelta(minutes=2)).isoformat(),
            "session": sess,
            "sent": False, "created_at": now.isoformat()})
    if registered:
        updates["privacy_firmata"] = True
        updates["privacy_registered_at"] = now.isoformat()
    if updates:
        await db.clients.update_one({"id": client_id}, {"$set": updates})
    return {"status": "ok", "privacy_registered": registered, "registration_error": reg_error,
            "wa_error": wa_error, "review_scheduled_at": (now + timedelta(minutes=2)).isoformat() if not wa_error else None}

@api_router.post("/clients/{client_id}/whatsapp/review")
async def whatsapp_review(client_id: str, user: dict = Depends(get_current_user)):
    c = await get_scoped_client(client_id, user)
    if not c.get("telefono"):
        raise HTTPException(status_code=400, detail="Il cliente non ha un numero di telefono")
    await wa_send(c["telefono"], await review_msg_cliente(c), session=await wa_session_cliente(c), tipo="recensione", client_id=client_id)
    await db.clients.update_one({"id": client_id},
                                {"$set": {"review_msg_sent_at": datetime.now(timezone.utc).isoformat()}})
    return {"status": "ok"}

async def process_whatsapp_queue():
    now = datetime.now(timezone.utc).isoformat()
    due = await db.whatsapp_queue.find({"sent": False, "send_after": {"$lte": now}}, {"_id": 0}).to_list(100)
    sent_count = 0
    for item in due:
        try:
            await wa_send(item["phone"], item.get("message") or REVIEW_MSG, session=item.get("session", "default"),
                          tipo="recensione", client_id=item.get("client_id", ""), servizio_id=item.get("servizio_id", ""))
            await db.whatsapp_queue.update_one({"id": item["id"]}, {"$set": {"sent": True, "sent_at": now}})
            await db.clients.update_one({"id": item["client_id"]}, {"$set": {"review_msg_sent_at": now}})
            if item.get("servizio_id"):
                await db.servizi.update_one({"id": item["servizio_id"]}, {"$set": {"review_msg_sent_at": now}})
            sent_count += 1
        except Exception as e:
            logger.error(f"WhatsApp queue send fallito per {item['id']}: {getattr(e, 'detail', str(e))}")
    if sent_count:
        logger.info(f"WhatsApp queue: {sent_count} recensioni inviate")

# ---------------- Export Excel ----------------

EXPORT_COLUMNS = [
    ("cognome", "Cognome"), ("nome", "Nome"), ("codice_fiscale", "Codice Fiscale"), ("p_iva", "P.IVA"),
    ("indirizzo", "Indirizzo"), ("telefono", "Telefono"), ("email", "Email"), ("iban", "IBAN"),
    ("tipo_bolletta", "Tipo"), ("pod", "POD"), ("pdr", "PDR"), ("kw_potenza", "kW"),
    ("fornitore_provenienza", "Fornitore attuale"), ("nuovo_fornitore", "Nuovo fornitore"),
    ("tipo_contratto", "Tipo contratto"), ("costo_kwh_nuovo", "€/kWh nuovo"),
    ("costo_smc_nuovo", "€/Smc nuovo"), ("spese_fisse_nuovo", "Spese fisse nuove"),
    ("data_contratto", "Data contratto"), ("data_attivazione", "Attivazione"),
    ("data_rinnovo", "Rinnovo (10 mesi)"), ("data_scadenza", "Scadenza"),
    ("data_verifica", "Data verifica"), ("data_cambio", "Data cambio"),
    ("lavorazione_label", "Lavorazione"), ("pagato_label", "Pagato"),
    ("privacy_label", "Privacy firmata"), ("negozio_nome", "Negozio"), ("note", "Note"),
]

@api_router.get("/export/clients.xlsx")
async def export_clients(user: dict = Depends(get_current_user)):
    scope = client_scope_filter(user)
    proj = {"_id": 0, "id": 1, "cognome": 1, "nome": 1, "codice_fiscale": 1, "p_iva": 1,
            "indirizzo": 1, "telefono": 1, "email": 1, "iban": 1, "tipo_bolletta": 1, "pod": 1,
            "pdr": 1, "kw_potenza": 1, "fornitore_provenienza": 1, "nuovo_fornitore": 1,
            "tipo_contratto": 1, "costo_kwh_nuovo": 1, "costo_smc_nuovo": 1, "spese_fisse_nuovo": 1,
            "data_contratto": 1, "data_verifica": 1, "data_cambio": 1, "lavorazione": 1,
            "pagato": 1, "last_payment_date": 1, "privacy_firmata": 1, "venditore_id": 1, "note": 1}
    clients = await db.clients.find(scope, proj).sort("cognome", 1).to_list(10000)
    stores = {s["id"]: s["nome"] for s in await db.stores.find({}, {"_id": 0}).to_list(500)}
    lav_labels = {"cambiare": "Da Cambiare", "non_cambiare": "Non Cambiare", "cambio_effettuato": "Cambio Effettuato",
                  "in_quotazione": "In Quotazione", "richieste_bollette": "Richieste Bollette",
                  "in_attesa_ok": "In Attesa OK Cliente", "problema_tecnico": "Problema Tecnico",
                  "da_quotare": "Da Quotare", "contattare_cliente": "Contattare Cliente",
                  "non_vuole_cambiare": "Non Vuole Cambiare", "passa_in_negozio": "Passa in Negozio",
                  "attesa_documenti": "Attesa Documenti", "rinnovato": "Rinnovato"}
    rows = []
    for c in clients:
        compute_dates(c)
        c["lavorazione_label"] = lav_labels.get(c.get("lavorazione"), c.get("lavorazione", ""))
        c["pagato_label"] = "Pagato" if c["pagato_effettivo"] else "Non pagato"
        c["privacy_label"] = "Sì" if c.get("privacy_firmata") else "No"
        c["negozio_nome"] = stores.get(c.get("venditore_id"), "")
        rows.append({label: c.get(key, "") for key, label in EXPORT_COLUMNS})
    df = pd.DataFrame(rows, columns=[label for _, label in EXPORT_COLUMNS])
    buf = io.BytesIO()
    df.to_excel(buf, index=False, sheet_name="Clienti")
    buf.seek(0)
    fname = f"report_clienti_{date.today().isoformat()}.xlsx"
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})

# ---------------- Servizi (riparazioni / telefonia) ----------------

RIP_STATI = ["ingresso", "attesa_ricambio_cliente", "attesa_ricambio_carico", "in_attesa_cliente",
             "preventivo", "in_lavorazione", "pronto", "consegnato", "non_riparabile"]

SERVIZIO_TIPI = ["riparazione", "accessori", "vendita", "sim", "internet", "fisso"]

TIPO_TO_SECTION = {"riparazione": "riparazioni", "accessori": "telefonia", "vendita": "telefonia",
                   "sim": "telefonia", "internet": "telefonia", "fisso": "telefonia"}

SIM_OPERATORS = ["WINDTRE", "VERY", "TIM", "KENA", "FASTWEB", "HO", "ILIAD", "LYCA", "DIGI", "ENEL"]
FISSO_OPERATORS = ["EOLO", "WINDTRE", "FASTWEB", "ILIAD", "ENEL"]
INTERNET_OPERATORS = FISSO_OPERATORS

ALL_SECTIONS = ["energia", "riparazioni", "telefonia"]

def user_sections(user: dict) -> list:
    secs = user.get("sections")
    return secs if secs else ALL_SECTIONS

def servizio_scope(user: dict, tipo: str = "") -> dict:
    sections = user_sections(user)
    if user["role"] == "tecnico":
        allowed = ["riparazione"] if "riparazioni" in sections else []
        base = {}
    else:
        allowed = [t for t, sec in TIPO_TO_SECTION.items() if sec in sections]
        if user["role"] == "admin" or user.get("can_view_all"):
            base = {}
        else:
            base = {"venditore_id": {"$in": user.get("store_ids", [])}}
    if tipo:
        if tipo not in allowed:
            raise HTTPException(status_code=403, detail="Sezione non abilitata per questo utente")
        base["tipo"] = tipo
    else:
        base["tipo"] = {"$in": allowed} if allowed else "__nessuno__"
    return base

# ---------------- Regole prezzi riparazioni (per tipologia x marca) + listino fornitori ----------------
REGOLE_PREZZI_DEFAULT = [
    ("display", "apple", "Display iPhone OLED / originale", 40, 95, 100, 350),
    ("display_compatibile", "apple", "Display iPhone compatibile (matrice CPY)", 40, 115, 80, 200),
    ("display", "samsung", "Display Samsung service pack (con frame)", 45, 0, 100, 400),
    ("display", "altri", "Display altri Android (Xiaomi, Realme, ...)", 35, 10, 80, 200),
    ("batteria", "*", "Batteria", 42, 0, 50, 100),
    ("connettore", "*", "Connettore di ricarica / flat", 30, 10, 50, 90),
    ("fotocamera", "*", "Fotocamera / altoparlante / microfono / sensori", 35, 10, 50, 150),
    ("vetro_camera", "*", "Vetrino fotocamera", 18, 10, 30, 60),
    ("vetro_posteriore", "*", "Vetro posteriore / cover batteria", 30, 80, 60, 150),
    ("altro", "*", "Altro componente", 30, 70, 40, 250),
    ("software", "*", "Software / recupero dati / diagnosi (senza ricambio)", 35, 0, 30, 80),
    ("vetro_temperato", "*", "Vetro temperato (prezzo fisso)", 0, 0, 10, 10),
    ("pellicola", "*", "Pellicola (prezzo fisso)", 0, 0, 20, 20),
    ("cover", "*", "Cover (prezzo fisso)", 0, 0, 15, 15),
]
REGOLE_TIPOLOGIE = {"display": "Display OLED / originale", "display_compatibile": "Display compatibile", "batteria": "Batteria",
                    "connettore": "Connettore di ricarica", "fotocamera": "Fotocamera / audio / sensori", "vetro_camera": "Vetrino fotocamera",
                    "vetro_posteriore": "Vetro posteriore", "altro": "Altro componente", "software": "Software / diagnosi",
                    "vetro_temperato": "Vetro temperato", "pellicola": "Pellicola", "cover": "Cover"}
SPESE_SPEDIZIONE_PEZZO = 2.5  # quota spedizione fornitore caricata su ogni ricambio
REGOLE_VERSIONE = 2
_REGOLE_CACHE: list = []

class RegolaPrezzoInput(BaseModel):
    tipologia: str
    marca: str = "*"
    label: str
    manodopera: float
    ricarico_pct: float = 0
    prezzo_min: float = 0
    prezzo_max: float = 0

async def load_regole_prezzi():
    global _REGOLE_CACHE
    state = await db.import_state.find_one({"_id": "regole_prezzi_versione"})
    if await db.regole_prezzi.count_documents({}) == 0 or (state or {}).get("v", 1) < REGOLE_VERSIONE:
        now = datetime.now(timezone.utc).isoformat()
        await db.regole_prezzi.delete_many({})
        await db.regole_prezzi.insert_many([{"id": str(uuid.uuid4()), "tipologia": t, "marca": m, "label": l, "manodopera": man, "ricarico_pct": ric,
                                             "prezzo_min": mn, "prezzo_max": mx, "ordine": i, "updated_at": now}
                                            for i, (t, m, l, man, ric, mn, mx) in enumerate(REGOLE_PREZZI_DEFAULT)])
        await db.import_state.update_one({"_id": "regole_prezzi_versione"}, {"$set": {"v": REGOLE_VERSIONE, "at": now}}, upsert=True)
    _REGOLE_CACHE = await db.regole_prezzi.find({}, {"_id": 0}).sort("ordine", 1).to_list(100)

def marca_dispositivo(dispositivo: str) -> str:
    d = (dispositivo or "").lower()
    if any(k in d for k in ("iphone", "apple", "ipad", "macbook", "airpods", "watch")):
        return "apple"
    if "samsung" in d or "galaxy" in d:
        return "samsung"
    return "altri"

def regola_per(s: dict) -> Optional[dict]:
    tip = (s.get("tipo_ricambio") or "altro") if s.get("con_ricambio") else "software"
    marca = marca_dispositivo(s.get("dispositivo", ""))
    cands = [r for r in _REGOLE_CACHE if r["tipologia"] == tip]
    return next((r for r in cands if r["marca"] == marca), None) or next((r for r in cands if r["marca"] == "*"), None)

def _round5(x: float) -> float:
    return float(5 * round(x / 5))

@api_router.get("/regole-prezzi")
async def get_regole_prezzi(user: dict = Depends(get_current_user)):
    return {"regole": _REGOLE_CACHE, "tipologie": REGOLE_TIPOLOGIE}

@api_router.put("/regole-prezzi")
async def put_regole_prezzi(regole: List[RegolaPrezzoInput], admin: dict = Depends(require_admin)):
    now = datetime.now(timezone.utc).isoformat()
    docs = [{"id": str(uuid.uuid4()), **r.model_dump(), "ordine": i, "updated_at": now, "updated_by": admin["name"]} for i, r in enumerate(regole)]
    await db.regole_prezzi.delete_many({})
    if docs:
        await db.regole_prezzi.insert_many(docs)
    await db.import_state.update_one({"_id": "regole_prezzi_versione"}, {"$set": {"v": REGOLE_VERSIONE, "at": now}}, upsert=True)
    await load_regole_prezzi()
    return {"regole": _REGOLE_CACHE, "tipologie": REGOLE_TIPOLOGIE}

class ListinoRiga(BaseModel):
    codice: str = ""
    descrizione: str
    prezzo_netto: float
    fornitore: str = "Sifar"
    data_fattura: str = ""
    marca: str = ""
    tipologia: str = ""
    modello: str = ""
    brand: str = ""
    qualita: str = ""
    note: str = ""

_PRICE_RX = re.compile(r"(?<![\d.])(\d{1,4}(?:\.\d{3})*,\d{2}|\d{1,4}\.\d{2})(?!\d)")

def _parse_price(t: str) -> float:
    return float(t.replace(".", "").replace(",", ".")) if "," in t else float(t)

def _guess_tipologia(desc: str) -> str:
    d = desc.lower()
    if "vetro temperato" in d:
        return "vetro_temperato"
    if "pellicola" in d:
        return "pellicola"
    if "cover" in d and "batteria" not in d and "vetrino" not in d:
        return "cover"
    if "vetrino" in d or "vetro camera" in d or "vetrino camera" in d:
        return "vetro_camera" if "camera" in d else "vetro_posteriore"
    if ("display" in d or "lcd" in d or "touch" in d) and ("compatibile" in d or "cpy" in d):
        return "display_compatibile"
    for key, words in (("display", ("display", "lcd", "oled", "schermo", "touch", "guscio frontale")), ("batteria", ("batteria", "battery")),
                       ("connettore", ("connettore", "dock", "flat carica", "charging", "ricarica")), ("vetro_posteriore", ("vetro post", "back cover", "back glass", "scocca")),
                       ("fotocamera", ("fotocamera", "camera", "speaker", "altoparlante", "microfono", "buzzer"))):
        if any(w in d for w in words):
            return key
    return "altro"

def parse_fattura_text(text: str) -> list:
    rows = []
    for raw in text.splitlines():
        line = " ".join(raw.split())
        prices = _PRICE_RX.findall(line)
        if len(prices) < 1 or len(line) < 12:
            continue
        first = _PRICE_RX.search(line)
        desc = line[:first.start()].strip(" -|")
        if not desc or not re.search(r"[a-zA-Z]{3}", desc) or re.match(r"^(totale|imponibile|iva|subtotale|spese|trasporto|bollo|pagamento|scadenza)", desc.lower()):
            continue
        m = re.match(r"^([A-Z0-9][A-Z0-9\-_./]{3,})\s+(.*)$", desc)
        codice, descr = (m.group(1), m.group(2)) if m and any(ch.isdigit() for ch in m.group(1)) else ("", desc)
        vals = [_parse_price(p) for p in prices]
        unit = vals[-2] if len(vals) >= 2 else vals[-1]
        qty = re.search(r"\s(\d{1,3})\s*(?:pz|pcs|nr|n\.)?\s*$", desc)
        if qty:
            desc = desc[:qty.start()].strip()
            if m and codice:
                descr = descr[:descr.rfind(qty.group(1))].strip() if descr.rstrip().endswith(qty.group(1)) else descr
            else:
                descr = desc
        if len(vals) >= 2 and vals[-1] > vals[-2] and abs(vals[-1] / vals[-2] - round(vals[-1] / vals[-2])) < 0.01:
            unit = vals[-2]
        if unit <= 0 or unit > 5000:
            continue
        rows.append({"codice": codice, "descrizione": descr.strip(), "prezzo_netto": round(unit, 2), "tipologia": _guess_tipologia(descr),
                     "marca": marca_dispositivo(descr), "qty_hint": qty.group(1) if qty else ""})
    return rows

_SIFAR_MODEL_RX = re.compile(r"^([a-z]+)/([a-z0-9][a-z0-9\-+.]*)$")
_SIFAR_PRICE_RX = re.compile(r"^€\s*([\d.]+,\d{2})$")

def parse_sifar_text(text: str) -> list:
    """Richieste d'ordine / fatture Sifar: codice, descrizione (multi-riga), marca/modello, q.ta, prezzo cad."""
    lines = [" ".join(l.split()) for l in text.splitlines()]
    rows, buf, i = [], [], 0
    while i < len(lines):
        l = lines[i]
        m = _SIFAR_MODEL_RX.match(l)
        if not m:
            if l and not l.startswith(("Note articolo", "Cod. Articolo", "Descrizione", "Q.tà", "Prezzo cad", "Totale", "ordinata", "Disponibile")) \
               and not re.match(r"^\d{1,3}$", l) and not l.startswith("€") and "about:blank" not in l and not re.match(r"^\d{2}/\d{2}/\d{2}", l):
                buf.append(l)
            i += 1
            continue
        desc = " ".join(buf).strip(); buf = []
        desc = re.sub(r"^.*?\bTotale\b\s*", "", desc) if "Q.tà" in desc or "Richiesta d'ordine" in desc else desc
        desc = re.sub(r"^\d/\d\s*", "", desc)
        cm = re.match(r"^([A-Z0-9][A-Z0-9\-_.]{3,})\s+(.*)$", desc)
        if not (cm and (any(ch.isdigit() for ch in cm.group(1)) or len(cm.group(1)) >= 6)):
            inner = re.search(r"\b([A-Z]{2,}[A-Z0-9\-_.]*\d[A-Z0-9\-_.]*)\s+(.*)$", desc)
            cm = inner if inner else None
        codice, descr = (cm.group(1), cm.group(2)) if cm else ("", desc)
        j, qty, prezzo = i + 1, 1, None
        nums = []
        while j < len(lines) and j <= i + 6:
            if re.match(r"^\d{1,3}$", lines[j]):
                nums.append(int(lines[j]))
            elif _SIFAR_PRICE_RX.match(lines[j]):
                prezzo = _parse_price(_SIFAR_PRICE_RX.match(lines[j]).group(1)); j += 1
                break
            j += 1
        if nums:
            qty = nums[0]
        note = ""
        k = j
        while k < len(lines) and k <= j + 3:
            if lines[k].startswith("Note articolo"):
                note = lines[k + 1] if k + 1 < len(lines) and not _SIFAR_MODEL_RX.match(lines[k + 1]) else ""
                k += 2
                break
            k += 1
        i = max(k, j)
        if prezzo is None or not descr:
            continue
        brand, model = m.group(1), m.group(2).replace("-", " ")
        rows.append({"codice": codice, "descrizione": descr, "prezzo_netto": round(prezzo, 2), "qty_hint": str(qty), "note": note,
                     "marca": "apple" if brand == "apple" else "samsung" if brand == "samsung" else "altri", "brand": brand, "modello": model,
                     "tipologia": _guess_tipologia(descr), "qualita": _guess_qualita(descr)})
    return rows

def _guess_qualita(desc: str) -> str:
    d = desc.lower()
    if "matrice compatibile" in d or "cpy" in d or "compatibile" in d:
        return "compatibile"
    if "oled" in d or "eccelsa" in d or "premium" in d or "stessa tecnologia" in d:
        return "premium"
    if "originale" in d or "service pack" in d or "bulk" in d or "frame" in d:
        return "originale"
    return ""

@api_router.post("/listino/parse-fattura")
async def parse_fattura(file: UploadFile = File(...), admin: dict = Depends(require_admin)):
    if (file.content_type or "") != "application/pdf":
        raise HTTPException(status_code=400, detail="Carica un PDF")
    data = await file.read()
    doc = fitz.open("pdf", data)
    text = "\n".join(p.get_text() for p in doc)
    is_sifar = "sifar" in text.lower() or "Codice Utente SN" in text or "Richiesta d'ordine" in text
    fornitore = "Sifar" if is_sifar else ""
    m = re.search(r"del (\d{2}-\d{2}-\d{4})", text) or re.search(r"(\d{2}/\d{2}/\d{4})", text)
    rows = parse_sifar_text(text) if is_sifar else []
    if not rows:
        rows = parse_fattura_text(text)
    numero = re.search(r"(?:ordine|Fattura)\s*(?:nr\.?|n\.?)\s*(\d+)", text, re.I)
    return {"fornitore": fornitore, "data_fattura": (m.group(1).replace("-", "/") if m else ""), "numero": numero.group(1) if numero else "",
            "righe": rows, "righe_testo": len(text.splitlines())}

@api_router.post("/listino")
async def save_listino(righe: List[ListinoRiga], admin: dict = Depends(require_admin)):
    now = datetime.now(timezone.utc).isoformat()
    n = 0
    for r in righe:
        key = {"codice": r.codice, "fornitore": r.fornitore} if r.codice else {"descrizione": r.descrizione, "fornitore": r.fornitore}
        await db.listino.update_one(key, {"$set": {**r.model_dump(), "updated_at": now, "updated_by": admin["name"]},
                                          "$setOnInsert": {"id": str(uuid.uuid4()), "created_at": now}}, upsert=True)
        n += 1
    return {"status": "ok", "salvate": n}

@api_router.get("/listino")
async def search_listino(q: str = "", user: dict = Depends(get_current_user), limit: int = 30):
    flt = {}
    if q:
        toks = [re.escape(t) for t in q.split() if t]
        flt = {"$and": [{"$or": [{"descrizione": {"$regex": t, "$options": "i"}}, {"codice": {"$regex": t, "$options": "i"}},
                                 {"modello": {"$regex": t, "$options": "i"}}, {"brand": {"$regex": t, "$options": "i"}}]} for t in toks]}
    return await db.listino.find(flt, {"_id": 0}).sort("updated_at", -1).to_list(limit)

@api_router.delete("/listino/{riga_id}")
async def delete_listino(riga_id: str, admin: dict = Depends(require_admin)):
    await db.listino.delete_one({"id": riga_id})
    return {"status": "ok"}

def calcola_prezzo_riparazione(s: dict) -> Optional[float]:
    if s.get("tipo") != "riparazione":
        return None
    regola = regola_per(s)
    minuti_raw = int(s.get("minuti_lavoro") or 0)
    if regola:
        costo = (float(s.get("costo_componente") or 0) + SPESE_SPEDIZIONE_PEZZO) if s.get("con_ricambio") else 0.0
        extra_min = max(minuti_raw - 30, 0) * 0.22775
        base = costo * (1 + regola["ricarico_pct"] / 100) + regola["manodopera"] + extra_min
        prezzo = _round5(base * 1.22)
        if regola["prezzo_min"] and (not s.get("con_ricambio") or costo * 1.22 < regola["prezzo_min"]):
            prezzo = max(prezzo, regola["prezzo_min"])
        if regola["prezzo_max"] and costo * 1.22 + regola["manodopera"] * 1.22 <= regola["prezzo_max"]:
            prezzo = min(prezzo, regola["prezzo_max"])
        return round(prezzo, 2)
    if s.get("con_ricambio") and s.get("tipo_ricambio") == "batteria":
        base = float(s.get("costo_componente") or 0) + 2.0 + minuti_raw * 0.22775 + 20.0
        return round(base * 1.22, 2)
    minuti = max(minuti_raw, 30)
    lavoro = minuti * 0.22775
    if s.get("con_ricambio"):
        base = float(s.get("costo_componente") or 0) + 2.0 + lavoro + 60.0
    else:
        base = 30.0 + lavoro
    return round(base * 1.22, 2)

def scadenza_offerta(s: dict) -> Optional[tuple]:
    """(scadenza, annuale). Vincolo>0: attivazione+mesi. Vincolo 0/None: prossimo anniversario annuale."""
    if s.get("tipo") not in ("sim", "internet", "fisso") or not s.get("data_attivazione"):
        return None
    try:
        att = date.fromisoformat(str(s["data_attivazione"])[:10])
    except (ValueError, TypeError):
        return None
    mesi = int(s.get("vincolo_mesi") or 0)
    if mesi > 0:
        return att + relativedelta(months=mesi), False
    scad = att + relativedelta(years=1)
    while scad < date.today() - timedelta(days=30):
        scad += relativedelta(years=1)
    return scad, True

def serialize_servizio(s: dict, client_name: str = "") -> dict:
    out = {k: v for k, v in s.items() if k != "_id" and k not in SECRET_FIELDS and not k.endswith("_enc")}
    out["client_name"] = client_name
    for f in SECRET_FIELDS:
        out["has_" + f] = bool(s.get(f + "_enc") or s.get(f))
    costi = costi_riparazione(s)
    if costi:
        out["costi"] = costi
        pf = s.get("prezzo_finale")
        if pf is not None:
            out["margine_reale"] = round(float(pf) / 1.22 - costi["totale"], 2)
    so = scadenza_offerta(s)
    if so:
        out["scadenza_vincolo"] = so[0].isoformat()
        out["giorni_alla_scadenza"] = (so[0] - date.today()).days
        out["offerta_annuale"] = so[1]
    out["prezzo_consigliato"] = calcola_prezzo_riparazione(s)
    if s.get("tipo") == "riparazione" and not s.get("data_ingresso") and s.get("created_at"):
        out["data_ingresso"] = str(s["created_at"])[:10]
    return out

async def get_scoped_servizio(servizio_id: str, user: dict) -> dict:
    scope = servizio_scope(user)
    scope["id"] = servizio_id
    s = await db.servizi.find_one(scope, {"_id": 0})
    if not s:
        raise HTTPException(status_code=404, detail="Servizio non trovato")
    return s

async def clients_name_map(client_ids: list) -> dict:
    if not client_ids:
        return {}
    docs = await db.clients.find({"id": {"$in": client_ids}}, {"_id": 0, "id": 1, "nome": 1, "cognome": 1}).to_list(10000)
    return {d["id"]: f"{d.get('cognome', '')} {d.get('nome', '')}".strip() for d in docs}

class ServizioInput(BaseModel):
    client_id: str
    tipo: str
    venditore_id: str = ""
    operatore_id: str = ""
    stato: str = "ingresso"
    dispositivo: str = ""
    problema: str = ""
    con_ricambio: bool = False
    tipo_ricambio: str = "altro"
    costo_componente: Optional[float] = None
    minuti_lavoro: Optional[int] = None
    prodotto: str = ""
    operatore_tel: str = ""
    numero: str = ""
    iccid: str = ""
    data_attivazione: Optional[str] = None
    vincolo_mesi: Optional[int] = None
    importo: Optional[float] = None
    pagato: bool = False
    note: str = ""
    codice_sblocco_tipo: str = ""
    codice_sblocco: str = ""
    account_email: str = ""
    account_password: str = ""
    operazioni: str = ""
    cliente_contattato: bool = False
    data_ingresso: Optional[str] = None
    data_lavorazione: Optional[str] = None
    data_uscita: Optional[str] = None
    prezzo_finale: Optional[float] = None

SECRET_FIELDS = ("codice_sblocco", "account_password")

def cifra_segreti(data: dict, old: Optional[dict] = None) -> None:
    """Sposta i campi sensibili in *_enc cifrati. Stringa vuota in update = mantieni valore esistente."""
    for f in SECRET_FIELDS:
        if f not in data:
            continue
        val = data.pop(f)
        if val:
            data[f + "_enc"] = totp_encrypt(val)
        elif old is None:
            data[f + "_enc"] = None

def leggi_segreti(s: dict) -> dict:
    out = {}
    for f in SECRET_FIELDS:
        enc = s.get(f + "_enc")
        out[f] = totp_decrypt(enc) if enc else (s.get(f) or "")
    return out

def cancella_segreti_update() -> dict:
    unset = {f: "" for f in SECRET_FIELDS}
    unset.update({f + "_enc": "" for f in SECRET_FIELDS})
    return {"$unset": unset, "$set": {"segreti_cancellati_at": datetime.now(timezone.utc).isoformat()}}

def costi_riparazione(s: dict) -> Optional[dict]:
    if s.get("tipo") != "riparazione":
        return None
    minuti_raw = int(s.get("minuti_lavoro") or 0)
    batteria = bool(s.get("con_ricambio")) and s.get("tipo_ricambio") == "batteria"
    minuti = minuti_raw if batteria else max(minuti_raw, 30)
    componente = (float(s.get("costo_componente") or 0) + SPESE_SPEDIZIONE_PEZZO) if s.get("con_ricambio") else 0.0
    lavoro = minuti * 0.22775
    return {"componente": round(componente, 2), "lavoro": round(lavoro, 2), "totale": round(componente + lavoro, 2)}


def apply_rip_dates(data: dict, old: Optional[dict] = None) -> None:
    if data.get("tipo", (old or {}).get("tipo")) != "riparazione":
        return
    today = date.today().isoformat()
    old = old or {}
    if not data.get("data_ingresso") and not old.get("data_ingresso"):
        data["data_ingresso"] = today
    stato = data.get("stato")
    if stato and stato != old.get("stato"):
        if stato in ("in_lavorazione", "pronto", "consegnato") and not data.get("data_lavorazione") and not old.get("data_lavorazione"):
            data["data_lavorazione"] = today
        if stato == "consegnato" and not data.get("data_uscita") and not old.get("data_uscita"):
            data["data_uscita"] = today

@api_router.get("/servizi")
async def list_servizi(user: dict = Depends(get_current_user), tipo: str = "", stato: str = "",
                       venditore_id: str = "", q: str = ""):
    scope = servizio_scope(user, tipo)
    if stato:
        scope["stato"] = stato
    if venditore_id and (user["role"] == "admin" or user.get("can_view_all")):
        scope["venditore_id"] = venditore_id
    servizi = await db.servizi.find(scope, {"_id": 0}).sort("created_at", -1).to_list(5000)
    names = await clients_name_map(list({s["client_id"] for s in servizi}))
    if q:
        ql = q.lower()
        servizi = [s for s in servizi if ql in names.get(s["client_id"], "").lower()
                   or ql in s.get("dispositivo", "").lower() or ql in s.get("numero", "").lower()
                   or ql in s.get("operatore_tel", "").lower() or ql in s.get("prodotto", "").lower()]
    return [serialize_servizio(s, names.get(s["client_id"], "")) for s in servizi]

@api_router.post("/servizi")
async def create_servizio(input: ServizioInput, user: dict = Depends(get_current_user)):
    if input.tipo not in SERVIZIO_TIPI:
        raise HTTPException(status_code=400, detail="Tipo servizio non valido")
    servizio_scope(user, input.tipo)  # gate sezione
    client_doc = await db.clients.find_one({"id": input.client_id}, {"_id": 0})
    if not client_doc:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    data = input.model_dump()
    if user["role"] in ("negozio",) and user.get("store_ids"):
        data["venditore_id"] = user["store_ids"][0]
    if data["tipo"] == "riparazione" and data["stato"] not in RIP_STATI:
        data["stato"] = "ingresso"
    if data["tipo"] == "riparazione":
        data["numero_riparazione"] = await next_numero("riparazione", data.get("venditore_id") or "")
    apply_rip_dates(data)
    cifra_segreti(data)
    now = datetime.now(timezone.utc).isoformat()
    data.update({"id": str(uuid.uuid4()), "created_by": user["id"], "created_at": now, "updated_at": now,
                 "last_payment_date": date.today().isoformat() if data.get("pagato") else None,
                 "privacy_firmata": False, "privacy_msg_sent_at": None, "review_msg_sent_at": None})
    await db.servizi.insert_one(data)
    data.pop("_id", None)
    return serialize_servizio(data, f"{client_doc.get('cognome', '')} {client_doc.get('nome', '')}".strip())

@api_router.get("/servizi/{servizio_id}/segreti")
async def get_servizio_segreti(servizio_id: str, user: dict = Depends(get_current_user)):
    s = await get_scoped_servizio(servizio_id, user)
    return {**leggi_segreti(s), "codice_sblocco_tipo": s.get("codice_sblocco_tipo", ""), "account_email": s.get("account_email", ""),
            "segreti_cancellati_at": s.get("segreti_cancellati_at")}

@api_router.get("/servizi/{servizio_id}")
async def get_servizio(servizio_id: str, user: dict = Depends(get_current_user)):
    s = await get_scoped_servizio(servizio_id, user)
    names = await clients_name_map([s["client_id"]])
    client_doc = await db.clients.find_one({"id": s["client_id"]}, {"_id": 0, "telefono": 1, "email": 1, "nome": 1, "cognome": 1, "codice_fiscale": 1, "no_recensioni": 1})
    out = serialize_servizio(s, names.get(s["client_id"], ""))
    out["client_contacts"] = client_doc or {}
    return out

@api_router.patch("/servizi/{servizio_id}")
async def update_servizio(servizio_id: str, input: ServizioInput, background_tasks: BackgroundTasks, user: dict = Depends(get_current_user)):
    old = await get_scoped_servizio(servizio_id, user)
    data = input.model_dump(exclude_unset=True)
    if user["role"] == "negozio":
        data.pop("venditore_id", None)
    apply_rip_dates(data, old)
    cifra_segreti(data, old)
    data["updated_at"] = datetime.now(timezone.utc).isoformat()
    if data.get("pagato") and not old.get("pagato"):
        data["last_payment_date"] = date.today().isoformat()
    await db.servizi.update_one({"id": servizio_id}, {"$set": data})
    if data.get("stato") in ("consegnato", "non_riparabile") and old.get("stato") not in ("consegnato", "non_riparabile"):
        await db.servizi.update_one({"id": servizio_id}, cancella_segreti_update())
    updated = await db.servizi.find_one({"id": servizio_id}, {"_id": 0})
    if updated.get("tipo") == "riparazione" and updated.get("stato") == "pronto" and old.get("stato") != "pronto" \
            and not updated.get("pronto_msg_sent_at"):
        background_tasks.add_task(invia_avviso_pronto, updated)
    if updated.get("tipo") == "riparazione" and updated.get("stato") == "consegnato" and old.get("stato") != "consegnato":
        if await accoda_recensione_riparazione(updated):
            updated["review_queued_at"] = datetime.now(timezone.utc).isoformat()
    names = await clients_name_map([updated["client_id"]])
    return serialize_servizio(updated, names.get(updated["client_id"], ""))

async def accoda_recensione_riparazione(svc: dict) -> bool:
    if svc.get("review_msg_sent_at") or svc.get("review_queued_at"):
        return False
    client_doc = await db.clients.find_one({"id": svc["client_id"]}, {"_id": 0, "telefono": 1, "no_recensioni": 1})
    if not client_doc or not client_doc.get("telefono") or client_doc.get("no_recensioni"):
        return False
    store = await db.stores.find_one({"id": svc.get("venditore_id")}, {"_id": 0})
    if not store or not store.get("review_link"):
        return False
    now = datetime.now(timezone.utc)
    await db.whatsapp_queue.insert_one({
        "id": str(uuid.uuid4()), "client_id": svc["client_id"], "servizio_id": svc["id"],
        "phone": client_doc["telefono"], "type": "review", "message": store_msg(store, "msg_recensione"),
        "session": svc.get("venditore_id") or "default",
        "send_after": (now + timedelta(minutes=2)).isoformat(), "sent": False, "created_at": now.isoformat()})
    await db.servizi.update_one({"id": svc["id"]}, {"$set": {"review_queued_at": now.isoformat()}})
    return True

class BlacklistInput(BaseModel):
    no_recensioni: bool

@api_router.get("/clients/{client_id}/whatsapp-log")
async def client_whatsapp_log(client_id: str, user: dict = Depends(get_current_user)):
    c = await get_scoped_client(client_id, user)
    cond = [{"client_id": client_id}]
    if c.get("telefono"):
        cond.append({"phone": c["telefono"]})
    logs = await db.wa_log.find({"$or": cond}, {"_id": 0}).sort("at", -1).to_list(200)
    return logs

@api_router.get("/servizi/{servizio_id}/whatsapp-log")
async def servizio_whatsapp_log(servizio_id: str, user: dict = Depends(get_current_user)):
    await get_scoped_servizio(servizio_id, user)
    return await db.wa_log.find({"servizio_id": servizio_id}, {"_id": 0}).sort("at", -1).to_list(200)

@api_router.post("/whatsapp-log/{log_id}/resend")
async def resend_whatsapp_log(log_id: str, user: dict = Depends(get_current_user)):
    log = await db.wa_log.find_one({"id": log_id}, {"_id": 0})
    if not log:
        raise HTTPException(status_code=404, detail="Messaggio non trovato")
    if log.get("client_id"):
        await get_scoped_client(log["client_id"], user)
    try:
        await wa_send(log["phone"], log["message"], session=log.get("session") or "default", tipo=log.get("tipo", ""),
                      client_id=log.get("client_id", ""), servizio_id=log.get("servizio_id", ""))
    except HTTPException as e:
        raise HTTPException(status_code=400, detail=f"WhatsApp: {e.detail}")
    now = datetime.now(timezone.utc).isoformat()
    await db.wa_log.update_one({"id": log_id}, {"$set": {"resent_at": now}})
    marks = {"privacy": "privacy_msg_sent_at", "pronto": "pronto_msg_sent_at", "promemoria": "promemoria_msg_sent_at", "recensione": "review_msg_sent_at"}
    field = marks.get(log.get("tipo"))
    if field and log.get("servizio_id"):
        await db.servizi.update_one({"id": log["servizio_id"]}, {"$set": {field: now}, "$unset": {"pronto_msg_error": ""}})
    if field and log.get("client_id") and field in ("privacy_msg_sent_at", "review_msg_sent_at"):
        await db.clients.update_one({"id": log["client_id"]}, {"$set": {field: now}})
    return {"status": "ok"}

@api_router.get("/clients/{client_id}/messaggi-previsti")
async def messaggi_previsti(client_id: str, user: dict = Depends(get_current_user)):
    c = compute_dates(await get_scoped_client(client_id, user))
    today = date.today()
    items = []
    if c.get("data_attivazione"):
        att = date.fromisoformat(c["data_attivazione"])
        d = att + timedelta(days=TRUFFE_DOPO_GIORNI)
        items.append({"tipo": "truffe", "label": "Attenzione truffe", "data": d.isoformat(),
                      "inviato_il": c.get("truffe_msg_sent_at"),
                      "stato": "inviato" if c.get("truffe_msg_sent_at") else ("saltato" if (today - att).days > TRUFFE_FINESTRA_GIORNI else "previsto")})
    if c.get("data_scadenza"):
        scad = date.fromisoformat(c["data_scadenza"])
        d = scad - timedelta(days=RINNOVO_PREAVVISO_GIORNI)
        sent = c.get("rinnovo_msg_sent_for") == c["data_scadenza"]
        items.append({"tipo": "rinnovo_energia", "label": "Rinnovo luce/gas", "data": d.isoformat(), "scadenza": c["data_scadenza"],
                      "inviato_il": c.get("rinnovo_msg_sent_at") if sent else None,
                      "stato": "inviato" if sent else ("saltato" if today > scad else "previsto")})
    servizi = await db.servizi.find({"client_id": client_id, "tipo": {"$in": ["sim", "internet", "fisso"]}}, {"_id": 0}).to_list(100)
    for s in servizi:
        so = scadenza_offerta(s)
        if not so:
            continue
        scad, annuale = so
        d = scad - timedelta(days=VINCOLO_PREAVVISO_GIORNI)
        sent = s.get("vincolo_msg_sent_for") == scad.isoformat()
        items.append({"tipo": "offerta_annuale" if annuale else "vincolo",
                      "label": f"{'Offerta annuale' if annuale else 'Scadenza vincolo'} {s.get('operatore_tel', '')} {s.get('numero', '')}".strip(),
                      "data": d.isoformat(), "scadenza": scad.isoformat(),
                      "inviato_il": s.get("vincolo_msg_sent_at") if sent else None,
                      "stato": "inviato" if sent else ("saltato" if today > scad else "previsto")})
    return sorted(items, key=lambda x: x["data"])

@api_router.get("/clients/{client_id}/gdpr-export")
async def client_gdpr_export(client_id: str, admin: dict = Depends(require_admin)):
    c = await db.clients.find_one({"id": client_id}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    compute_dates(c)
    stores = {s["id"]: s["nome"] for s in await db.stores.find({}, {"_id": 0, "id": 1, "nome": 1}).to_list(200)}
    servizi = await db.servizi.find({"client_id": client_id}, {"_id": 0}).sort("created_at", -1).to_list(500)
    ritiri = await db.ritiri.find({"client_id": client_id}, {"_id": 0}).sort("created_at", -1).to_list(200)
    allegati = await db.attachments.find({"client_id": client_id}, {"_id": 0, "original_filename": 1, "created_at": 1}).to_list(200)
    wa = await db.wa_log.find({"$or": [{"client_id": client_id}, {"phone": c.get("telefono", "__")}]}, {"_id": 0}).sort("at", -1).to_list(500)
    audit = await db.audit_log.find({"entity": "cliente", "entity_id": client_id}, {"_id": 0}).sort("at", -1).to_list(500)

    def _riga(pdf_, txt):
        pdf_.set_x(pdf_.l_margin)
        pdf_.multi_cell(0, 5, txt)

    pdf = _new_pdf("Estratto dati personali (art. 15 GDPR)",
                   f"{c.get('cognome', '')} {c.get('nome', '')}  -  generato il {date.today().strftime('%d/%m/%Y')}")
    pdf.set_font("helvetica", "B", 11); pdf.set_x(pdf.l_margin); pdf.cell(0, 7, "1. Dati anagrafici e di contatto", new_x="LMARGIN", new_y="NEXT"); pdf.set_font("helvetica", "", 10)
    for label, key in [("Nome", "nome"), ("Cognome", "cognome"), ("Codice fiscale", "codice_fiscale"), ("Partita IVA", "piva"),
                       ("Telefono", "telefono"), ("Email", "email"), ("Indirizzo", "indirizzo"), ("IBAN", "iban"),
                       ("Tipo cliente", "tipo_cliente"), ("Negozio di riferimento", None), ("Registrato il", "created_at")]:
        if key is None:
            val = stores.get(c.get("venditore_id"), "-")
        elif key == "created_at":
            val = _fmt_it(c.get(key))
        else:
            val = c.get(key) or "-"
        _pdf_field(pdf, label, str(val))
    pdf.ln(2)
    pdf.set_font("helvetica", "B", 11); pdf.set_x(pdf.l_margin); pdf.cell(0, 7, "2. Consensi e comunicazioni", new_x="LMARGIN", new_y="NEXT"); pdf.set_font("helvetica", "", 10)
    _pdf_field(pdf, "Privacy firmata", "Si" if c.get("privacy_firmata") else "No")
    _pdf_field(pdf, "Privacy inviata via WhatsApp il", _fmt_it(c.get("privacy_msg_sent_at")) or "-")
    _pdf_field(pdf, "Richieste recensione", "Bloccate (blacklist)" if c.get("no_recensioni") else "Consentite")
    pdf.ln(2)
    pdf.set_font("helvetica", "B", 11); pdf.set_x(pdf.l_margin); pdf.cell(0, 7, "3. Contratto energia", new_x="LMARGIN", new_y="NEXT"); pdf.set_font("helvetica", "", 10)
    for label, key in [("Tipo bolletta", "tipo_bolletta"), ("Fornitore attuale", "fornitore_attuale"), ("Nuovo fornitore", "nuovo_fornitore"),
                       ("Lavorazione", "lavorazione"), ("Data contratto", "data_contratto"), ("Attivazione", "data_attivazione"),
                       ("Scadenza", "data_scadenza"), ("POD/PDR", "pod_pdr")]:
        v = c.get(key)
        if v:
            _pdf_field(pdf, label, _fmt_it(v) if key.startswith("data_") else str(v))
    pdf.ln(2)
    pdf.set_font("helvetica", "B", 11); pdf.set_x(pdf.l_margin); pdf.cell(0, 7, f"4. Servizi (riparazioni / telefonia): {len(servizi)}", new_x="LMARGIN", new_y="NEXT"); pdf.set_font("helvetica", "", 10)
    for s in servizi:
        riga = f"{_fmt_it(s.get('created_at'))} - {s.get('tipo', '')} {s.get('numero_riparazione') or s.get('numero') or ''} - {s.get('dispositivo') or s.get('operatore_tel') or ''} - stato {s.get('stato', '')}"
        if s.get("problema"):
            riga += f" - {s['problema'][:80]}"
        _riga(pdf, riga); pdf.ln(0.5)
    pdf.ln(2)
    pdf.set_font("helvetica", "B", 11); pdf.set_x(pdf.l_margin); pdf.cell(0, 7, f"5. Bolle di ritiro usato: {len(ritiri)}", new_x="LMARGIN", new_y="NEXT"); pdf.set_font("helvetica", "", 10)
    for r in ritiri:
        _riga(pdf, f"{_fmt_it(r.get('created_at'))} - {r.get('numero', '')} - {r.get('articolo', '')} IMEI {r.get('imei') or '-'} - EUR {r.get('prezzo_ritiro') or 0}")
    pdf.ln(2)
    pdf.set_font("helvetica", "B", 11); pdf.set_x(pdf.l_margin); pdf.cell(0, 7, f"6. Documenti allegati: {len(allegati)}", new_x="LMARGIN", new_y="NEXT"); pdf.set_font("helvetica", "", 10)
    for a in allegati:
        _riga(pdf, f"{_fmt_it(a.get('created_at'))} - {a.get('original_filename', '')}")
    pdf.ln(2)
    pdf.set_font("helvetica", "B", 11); pdf.set_x(pdf.l_margin); pdf.cell(0, 7, f"7. Messaggi WhatsApp inviati: {len(wa)}", new_x="LMARGIN", new_y="NEXT"); pdf.set_font("helvetica", "", 10)
    for m in wa:
        _riga(pdf, f"{_fmt_it(m.get('at'))} - {m.get('tipo') or 'messaggio'} - {'inviato' if m.get('ok') else 'non inviato'}")
    pdf.ln(2)
    pdf.set_font("helvetica", "B", 11); pdf.set_x(pdf.l_margin); pdf.cell(0, 7, f"8. Accessi ai dati da parte del personale: {len(audit)}", new_x="LMARGIN", new_y="NEXT"); pdf.set_font("helvetica", "", 10)
    for a in audit[:200]:
        _riga(pdf, f"{_fmt_it(a.get('at'))} - {a.get('user_name', '')} - {a.get('action', '')}")
    pdf.ln(4)
    pdf.set_font("helvetica", "I", 8)
    _riga(pdf, "Titolare del trattamento: RS Riparazioni / CambiaOra. Documento generato dal gestionale in risposta a richiesta di accesso ai dati "
                         "(art. 15 Reg. UE 2016/679). Le password e i codici di sblocco dei dispositivi non sono inclusi e vengono cancellati alla consegna.")
    buf = io.BytesIO(bytes(pdf.output()))
    fname = f"dati_personali_{c.get('cognome', '')}_{c.get('nome', '')}.pdf".replace(" ", "_")
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{fname}"'})

@api_router.get("/dashboard/margini-negozi")
async def margini_negozi(admin: dict = Depends(require_admin), mese: str = ""):
    mese = mese or date.today().strftime("%Y-%m")
    prec = (date.fromisoformat(mese + "-01") - relativedelta(months=1)).strftime("%Y-%m")
    rips = await db.servizi.find({"tipo": "riparazione", "stato": {"$in": ["consegnato", "pronto", "in_lavorazione"]}}, {"_id": 0}).to_list(20000)
    stores = {s["id"]: s["nome"] for s in await db.stores.find({}, {"_id": 0, "id": 1, "nome": 1}).to_list(200)}
    acc: dict = {}
    prec_acc: dict = {}
    for s in rips:
        rif = (s.get("data_uscita") or s.get("created_at") or "")[:7]
        if rif not in (mese, prec):
            continue
        costi = costi_riparazione(s) or {"totale": 0}
        prezzo = float(s.get("prezzo_finale") if s.get("prezzo_finale") is not None else (calcola_prezzo_riparazione(s) or 0))
        margine = prezzo / 1.22 - costi["totale"]
        sid = s.get("venditore_id", "")
        if rif == prec:
            prec_acc[sid] = prec_acc.get(sid, 0.0) + margine
            continue
        a = acc.setdefault(sid, {"store_id": sid, "store_name": stores.get(sid, "-"),
                                 "riparazioni": 0, "incasso": 0.0, "costi": 0.0, "margine": 0.0, "consegnate": 0})
        a["riparazioni"] += 1; a["incasso"] += prezzo; a["costi"] += costi["totale"]; a["margine"] += margine
        if s.get("stato") == "consegnato":
            a["consegnate"] += 1
    for sid in prec_acc:
        acc.setdefault(sid, {"store_id": sid, "store_name": stores.get(sid, "-"), "riparazioni": 0, "incasso": 0.0, "costi": 0.0, "margine": 0.0, "consegnate": 0})
    for a in acc.values():
        a["rigenerati"] = 0; a["margine_rigenerati"] = 0.0
    async for v in db.vendite_rigenerati.find({"venduto_at": {"$regex": f"^({re.escape(mese)}|{re.escape(prec)})"}}, {"_id": 0}):
        sid = v.get("store_id", "")
        if v["venduto_at"][:7] == prec:
            prec_acc[sid] = prec_acc.get(sid, 0.0) + v["margine"]
            continue
        a = acc.setdefault(sid, {"store_id": sid, "store_name": stores.get(sid, "-"), "riparazioni": 0, "incasso": 0.0,
                                 "costi": 0.0, "margine": 0.0, "consegnate": 0, "rigenerati": 0, "margine_rigenerati": 0.0})
        a["rigenerati"] += 1; a["margine_rigenerati"] += v["margine"]; a["margine"] += v["margine"]
    out = []
    for a in acc.values():
        mp = round(prec_acc.get(a["store_id"], 0.0), 2)
        out.append({**a, "incasso": round(a["incasso"], 2), "costi": round(a["costi"], 2), "margine": round(a["margine"], 2),
                    "margine_rigenerati": round(a.get("margine_rigenerati", 0.0), 2),
                    "margine_precedente": mp, "delta": round(a["margine"] - mp, 2)})
    return {"mese": mese, "mese_precedente": prec, "negozi": sorted(out, key=lambda x: -x["margine"]),
            "totale": round(sum(a["margine"] for a in out), 2), "totale_precedente": round(sum(prec_acc.values()), 2)}

# ---------------- Password manager (credenziali negozi) ----------------
class PasswordInput(BaseModel):
    servizio: str
    titolo: str = ""
    username: str = ""
    password: Optional[str] = None  # None in update = mantieni
    url: str = ""
    contenuto: Optional[str] = None  # testo libero cifrato (None in update = mantieni)
    store_ids: List[str] = []  # vuoto = visibile solo all'amministratore
    user_ids: List[str] = []  # utenti singoli autorizzati (es. Deborah)

class PasswordImportInput(BaseModel):
    sheet_url: str
    force: bool = False

PASSWORD_PUBLIC = {"_id": 0, "password_enc": 0, "contenuto_enc": 0}

def password_scope(user: dict) -> dict:
    if user["role"] == "admin":
        return {}
    return {"$or": [{"store_ids": {"$in": user.get("store_ids", [])}}, {"user_ids": user["id"]}]}

PW_STORE_KEYS = {"morbegno": "Morbegno", "morbe": "Morbegno", "sondrio": "Sondrio", "gravedona": "Gravedona",
                 "tirano": "Tirano", "sondalo": "Sondalo", "grosio": "Grosio", "ipro": "Ipro",
                 "colico": "Colico", "somaggia": "Somaggia"}
PW_STORE_RX = re.compile("|".join(sorted(PW_STORE_KEYS, key=len, reverse=True)), re.I)
PW_EXTRA_STORES = {"Colico", "Somaggia"}

def _split_contenuto_per_negozio(text: str) -> dict:
    """Ogni riga con un nome negozio apre/continua il blocco di quel negozio; le altre seguono il blocco corrente."""
    blocks: dict = {}
    current = ("generale",)
    for line in text.split("\n"):
        found = tuple(sorted({PW_STORE_KEYS[m.lower()] for m in PW_STORE_RX.findall(line)}))
        if found:
            current = found
        blocks.setdefault(current, []).append(line)
    return {k: "\n".join(v) for k, v in blocks.items()}

@api_router.post("/passwords/dividi-per-negozio")
async def split_passwords_per_negozio(admin: dict = Depends(require_admin)):
    stores = {s["nome"]: s["id"] for s in await db.stores.find({}, {"_id": 0, "id": 1, "nome": 1}).to_list(100)}
    deborah = await db.users.find_one({"email": "deborah@cambiaora.local"}, {"_id": 0, "id": 1})
    extra_users = [deborah["id"]] if deborah else []
    now = datetime.now(timezone.utc).isoformat()
    created, processed = 0, 0
    async for p in db.passwords.find({"import_gid": {"$exists": True}, "split_done": {"$ne": True}}):
        processed += 1
        text = totp_decrypt(p["contenuto_enc"]) if p.get("contenuto_enc") else ""
        blocks = _split_contenuto_per_negozio(text)
        generale = blocks.pop(("generale",), "")
        for names, chunk in blocks.items():
            sids = [stores[n] for n in names if n in stores]
            uids = extra_users if (set(names) & PW_EXTRA_STORES) else []
            await db.passwords.insert_one({
                "id": str(uuid.uuid4()), "servizio": p["servizio"], "titolo": " / ".join(names), "username": "",
                "password_enc": None, "url": "", "contenuto_enc": totp_encrypt(chunk), "store_ids": sids,
                "user_ids": uids, "import_gid": p["import_gid"], "import_source": p.get("import_source"),
                "split_done": True, "created_by": admin["id"], "created_at": now, "updated_at": now})
            created += 1
        if generale.strip() or not blocks:
            await db.passwords.update_one({"id": p["id"]}, {"$set": {
                "titolo": "generale" if blocks else p.get("titolo", ""), "split_done": True, "updated_at": now,
                "contenuto_enc": totp_encrypt(generale) if generale.strip() else p.get("contenuto_enc")}})
        else:
            await db.passwords.delete_one({"id": p["id"]})
    return {"status": "ok", "voci_analizzate": processed, "schede_create": created}

def serialize_password(p: dict) -> dict:
    out = {k: v for k, v in p.items() if k not in ("_id", "password_enc", "contenuto_enc")}
    out["has_password"] = bool(p.get("password_enc"))
    out["has_contenuto"] = bool(p.get("contenuto_enc"))
    return out

@api_router.get("/passwords")
async def list_passwords(user: dict = Depends(get_current_user), q: str = ""):
    scope = password_scope(user)
    if q:
        rx = {"$regex": re.escape(q), "$options": "i"}
        scope = {"$and": [scope, {"$or": [{"servizio": rx}, {"titolo": rx}, {"username": rx}, {"url": rx}]}]}
    rows = await db.passwords.find(scope).sort([("servizio", 1), ("titolo", 1)]).to_list(2000)
    return [serialize_password(p) for p in rows]

@api_router.get("/passwords/{password_id}/reveal")
async def reveal_password(password_id: str, user: dict = Depends(get_current_user)):
    p = await db.passwords.find_one({**password_scope(user), "id": password_id})
    if not p:
        raise HTTPException(status_code=404, detail="Voce non trovata")
    return {"id": p["id"], "password": totp_decrypt(p["password_enc"]) if p.get("password_enc") else "",
            "contenuto": totp_decrypt(p["contenuto_enc"]) if p.get("contenuto_enc") else ""}

def _password_doc(data: dict) -> dict:
    out = {k: v for k, v in data.items() if k not in ("password", "contenuto")}
    out["servizio"] = out["servizio"].strip()
    if data.get("password") is not None:
        out["password_enc"] = totp_encrypt(data["password"]) if data["password"] else None
    if data.get("contenuto") is not None:
        out["contenuto_enc"] = totp_encrypt(data["contenuto"]) if data["contenuto"] else None
    return out

@api_router.post("/passwords")
async def create_password(input: PasswordInput, admin: dict = Depends(require_admin)):
    now = datetime.now(timezone.utc).isoformat()
    doc = _password_doc(input.model_dump())
    doc.update({"id": str(uuid.uuid4()), "created_by": admin["id"], "created_at": now, "updated_at": now})
    await db.passwords.insert_one(doc)
    return serialize_password(doc)

@api_router.patch("/passwords/{password_id}")
async def update_password(password_id: str, input: PasswordInput, admin: dict = Depends(require_admin)):
    doc = _password_doc(input.model_dump())
    doc.update({"updated_at": datetime.now(timezone.utc).isoformat(), "updated_by": admin["id"]})
    res = await db.passwords.update_one({"id": password_id}, {"$set": doc})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Voce non trovata")
    return serialize_password(await db.passwords.find_one({"id": password_id}))

@api_router.delete("/passwords/{password_id}")
async def delete_password(password_id: str, admin: dict = Depends(require_admin)):
    res = await db.passwords.delete_one({"id": password_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Voce non trovata")
    return {"status": "ok"}

def _sheet_rows_to_text(csv_text: str) -> str:
    lines = []
    for row in csv.reader(io.StringIO(csv_text)):
        cells = [c.strip() for c in row if c and c.strip()]
        if cells:
            lines.append("  |  ".join(cells))
    return "\n".join(lines)

@api_router.post("/passwords/import-sheet")
async def import_passwords_sheet(input: PasswordImportInput, admin: dict = Depends(require_admin)):
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9\-_]+)", input.sheet_url)
    if not m:
        raise HTTPException(status_code=400, detail="Link Google Sheet non valido")
    sheet_id = m.group(1)
    marker_id = f"passwords_sheet:{sheet_id}"
    if not input.force and await db.import_state.find_one({"_id": marker_id}):
        return {"status": "gia_importato", "imported": 0}
    async with httpx.AsyncClient(timeout=60, follow_redirects=True) as http:
        page = await http.get(f"https://docs.google.com/spreadsheets/d/{sheet_id}/htmlview")
        if page.status_code != 200:
            raise HTTPException(status_code=400, detail="Foglio non raggiungibile: verifica la condivisione 'Chiunque con il link'")
        tabs = re.findall(r'items\.push\(\{name: "([^"]*)".*?gid: "(-?\d+)"', page.text)
        if not tabs:
            raise HTTPException(status_code=400, detail="Nessuna pagina trovata nel foglio")
        now = datetime.now(timezone.utc).isoformat()
        imported, skipped = 0, 0
        for name, gid in tabs:
            r = await _fetch_tab_csv(http, sheet_id, gid=gid)
            text = _sheet_rows_to_text(r.text) if r.status_code == 200 else ""
            name = name.replace("\\/", "/").strip()
            if not text or await db.passwords.find_one({"servizio": name, "import_gid": gid}):
                skipped += 1
                continue
            await db.passwords.insert_one({
                "id": str(uuid.uuid4()), "servizio": name, "titolo": "", "username": "", "password_enc": None,
                "url": "", "contenuto_enc": totp_encrypt(text), "store_ids": [], "import_gid": gid,
                "import_source": sheet_id, "created_by": admin["id"], "created_at": now, "updated_at": now})
            imported += 1
    await db.import_state.update_one({"_id": marker_id}, {"$set": {"at": now, "imported": imported}}, upsert=True)
    return {"status": "ok", "imported": imported, "skipped": skipped, "pagine": len(tabs)}

# ---------------- Portali operatori ----------------
class PortaleInput(BaseModel):
    sezione: str  # energia | mobile | fisso | riparazioni
    operatore: str  # "*" = valido per tutta la sezione
    nome: str = ""
    url: str = ""
    note: str = ""
    flag_label: str = ""  # es. "CARICATO SU JOY"
    flag_url: str = ""

PORTALI_SEED = [
    ("mobile", "WINDTRE", "Kolme", "https://spazio.kolme.it/", "", "", ""),
    ("fisso", "WINDTRE", "Kolme", "https://spazio.kolme.it/", "", "", ""),
    ("mobile", "VERY", "Kolme", "https://spazio.kolme.it/", "", "", ""),
    ("mobile", "KENA", "App Kena", "", "Apri l'app Kena per l'inserimento", "", ""),
    ("mobile", "HO", "Oscar ho.", "https://oscar.ho-mobile.it/login.html", "", "", ""),
    ("mobile", "DIGI", "Partner Digi", "https://parteneri.digimobil.it/cgi-bin/index.cgi", "", "", ""),
    ("mobile", "ILIAD", "Planet Iliad", "https://planet.iliad.it/login", "", "", ""),
    ("fisso", "ILIAD", "Planet Iliad", "https://planet.iliad.it/login", "", "", ""),
    ("mobile", "FASTWEB", "Portale Fastweb", "https://logon.fastweb.it/oam/server/obrareq.cgi?encquery%3DjXkCxgqyzKOrBAEASswVNrDrA6ALYQoV2QdTqOUAtYLisdDvOzgttcfRwtZmOUuullJrr5rxzgkYTaZEI86Fq2ubyUU6RfQ3ZEe1PHeuWgE5ApWesUZPJ5YrbjHD49Op6cl0CUrisMRCxReH%2BF3mLkoSLMt1shXPkuVkEywb3lS3Cq3qj9E4fKt9vWMTNmoYPrNPOMaz4431sISeDH23LrW%2BdCrHH8kcsw9u%2BoY9o15Tl%2Bh1Bd%2BWLutKDWI9M6274MMSyLle7Kb3jgoAhc0NFQ%3D%3D%20agentid%3DFront-End-OAM%20ver%3D1%20crmethod%3D2",
     "Dopo l'inserimento carica su JOY (serve per il pagamento)", "CARICATO SU JOY", "https://logon.fastweb.it/oam/server/obrareq.cgi?encquery%3DjXkCxgqyzKOrBAEASswVNrDrA6ALYQoV2QdTqOUAtYLisdDvOzgttcfRwtZmOUuullJrr5rxzgkYTaZEI86Fq2ubyUU6RfQ3ZEe1PHeuWgE5ApWesUZPJ5YrbjHD49Op6cl0CUrisMRCxReH%2BF3mLkoSLMt1shXPkuVkEywb3lS3Cq3qj9E4fKt9vWMTNmoYPrNPOMaz4431sISeDH23LrW%2BdCrHH8kcsw9u%2BoY9o15Tl%2Bh1Bd%2BWLutKDWI9M6274MMSyLle7Kb3jgoAhc0NFQ%3D%3D%20agentid%3DFront-End-OAM%20ver%3D1%20crmethod%3D2"),
    ("fisso", "FASTWEB", "Portale Fastweb", "https://logon.fastweb.it/oam/server/obrareq.cgi?encquery%3DjXkCxgqyzKOrBAEASswVNrDrA6ALYQoV2QdTqOUAtYLisdDvOzgttcfRwtZmOUuullJrr5rxzgkYTaZEI86Fq2ubyUU6RfQ3ZEe1PHeuWgE5ApWesUZPJ5YrbjHD49Op6cl0CUrisMRCxReH%2BF3mLkoSLMt1shXPkuVkEywb3lS3Cq3qj9E4fKt9vWMTNmoYPrNPOMaz4431sISeDH23LrW%2BdCrHH8kcsw9u%2BoY9o15Tl%2Bh1Bd%2BWLutKDWI9M6274MMSyLle7Kb3jgoAhc0NFQ%3D%3D%20agentid%3DFront-End-OAM%20ver%3D1%20crmethod%3D2",
     "Dopo l'inserimento carica su JOY (serve per il pagamento)", "CARICATO SU JOY", "https://logon.fastweb.it/oam/server/obrareq.cgi?encquery%3DjXkCxgqyzKOrBAEASswVNrDrA6ALYQoV2QdTqOUAtYLisdDvOzgttcfRwtZmOUuullJrr5rxzgkYTaZEI86Fq2ubyUU6RfQ3ZEe1PHeuWgE5ApWesUZPJ5YrbjHD49Op6cl0CUrisMRCxReH%2BF3mLkoSLMt1shXPkuVkEywb3lS3Cq3qj9E4fKt9vWMTNmoYPrNPOMaz4431sISeDH23LrW%2BdCrHH8kcsw9u%2BoY9o15Tl%2Bh1Bd%2BWLutKDWI9M6274MMSyLle7Kb3jgoAhc0NFQ%3D%3D%20agentid%3DFront-End-OAM%20ver%3D1%20crmethod%3D2"),
    ("fisso", "EOLO", "Eolo", "https://www.eolo.it/", "", "", ""),
    ("mobile", "LYCA", "Lyca Retailer", "https://retailer.lycamobile.it/", "", "", ""),
    ("energia", "*", "CambiaOra - primo passo attivazione", "https://swi.tc/pom/4jb", "Primo passo per attivare l'utenza", "", ""),
    ("energia", "ENEL", "Portale Enel", "https://login.microsoftonline.com/d539d4bf-5610-471a-afc2-1c76685cfefa/saml2?SAMLRequest=hZNdc6owEIb%2FCpN7MHwrU%2B2gaK2KX6BtvXECJkKBBEj86q8%2FjG1nes5Fz87sRWY375vNPPvweC1y6YxrnjLaBaoCgYRpzA4pPXbBJhzJbfDYe%2BCoyLXScU8ioWtcnTAXUnORcuez0gWnmjoM8ZQ7FBWYOyJ2AtefOZoCnbJmgsUsB5LLOa5FYzVglJ8KXAe4Pqcx3qxnXZAIUXKn1cIU53FdCKW4KRzlmBNWx1iJWfHIWRdCTzPgPcpwcAKS1zwmpUjcB%2FjWyNkxpUqRxjXjjAhG85TeJVoHU%2B8cjIjIpqVC2bBVJCMSa7Ia25bVNmOCCWrdpwLSs9cFe23gNoGG0YhsVd9zL%2F9mU%2B33z%2FhsbytyvanlnCRLHeFrnq5V8rJezL1ZcayW%2B%2FWUjiZZdNx5%2BSWYuLBzXOux7R5W5eS6omrWfqmFp2uZbkQ3G1XRuAwPoyc5CzSqz1B7mp2tKNkWtyla7TcTw73tbTLIWX9HPhbB5OKaBg%2FzQWLPN35N39Txcfi%2BfyKBGN2GQWbvVf9FJ5mbvVn%2BIEu83Uobmpr5VC6z8aa%2FXOhWUXWi4Sucm2pmw1dxYf5HOfUP20R%2FppNDVnVuhvke%2B6bH9tXrxw72g%2FlmZZ63b%2BM8Kt3mtzg%2F4WfKBaKiCzSoWTLsyCoMVejolqNZit6BOyAtv3Dop%2FQTs9%2FYiT6buDMOw6W8XAQhkLbfuDYN4AtO5%2B5e%2F6Tyd2H0jSLo%2FR%2B8h9ZPk97X8e%2BF6P0B&RelayState=%2F&SigAlg=http%3A%2F%2Fwww.w3.org%2F2001%2F04%2Fxmldsig-more%23rsa-sha256&Signature=OpjLgPTQBa9YaWoW33forzaqln1Ub11q6GvS8Vvii4UHJy6cZJ%2BbkkVcfot9KD5DCxO3GIddyn0MhZRixbvyFSbXza9fjkJuKDGvcWmvsYuAOyinfm4QDhIC1VBLjUHcJZ8u4sOx4xOKxp7aOfZsjwt29Ft%2Bd79CrzB7F71PuYIFAZhhpdFThd4yyH%2B%2FD0tuTAxy%2Fzjb1Us1owI62Qvi2wh9Y9EmVhT5Ogtmpr7VyBG3rb9yFsT6rDaRSpaNI0lAS9TMJnkCL96Pa1GqNRcv8drGSO2%2Bboq7%2BYZKaarhyZvqIWBYGh6Fk7j2%2FrPBoaWpUV2fmLwSrNpjEK2b94MwHQ%3D%3D&sso_reload=true", "", "", ""),
    ("riparazioni", "*", "SIFAR (ricambi)", "https://www.sifar.it/it", "", "", ""),
    ("riparazioni", "*", "MobileSentrix (ricambi)", "https://www.mobilesentrix.eu/", "", "", ""),
    ("riparazioni", "*", "New Best (ricambi)", "https://www.newbest-ricambi.com/index.php", "", "", ""),
    ("riparazioni", "*", "El Hope (ricambi)", "https://www.el-hope.com/shop/authentication?back=my-account", "", "", ""),
    ("riparazioni", "*", "New Net (ricambi)", "https://newnetsrl.com/", "", "", ""),
    ("riparazioni", "*", "5G (ricambi)", "https://www.5g-m.com/it/login?back=my-account", "", "", ""),
    ("riparazioni", "*", "Phone Click (telefoni nuovi)", "https://www.phoneclick.it/index.asp", "Acquisto telefoni nuovi", "", ""),
    ("riparazioni", "*", "MIWO (telefoni nuovi e rigenerati)", "https://www.miwo.it/", "Acquisto telefoni nuovi e rigenerati", "", ""),
    ("riparazioni", "*", "CDR International (telefoni nuovi)", "http://www.cdrinternational.com/", "Acquisto telefoni nuovi", "", ""),
]

async def seed_portali() -> None:
    if await db.portali.count_documents({}):
        return
    now = datetime.now(timezone.utc).isoformat()
    await db.portali.insert_many([{"id": str(uuid.uuid4()), "sezione": s, "operatore": o, "nome": n, "url": u, "note": nt,
                                   "flag_label": fl, "flag_url": fu, "created_at": now} for s, o, n, u, nt, fl, fu in PORTALI_SEED])
    logger.info(f"Portali operatori seedati: {len(PORTALI_SEED)}")

PORTALE_SEZIONE_BY_TIPO = {"sim": "mobile", "internet": "fisso", "fisso": "fisso", "riparazione": "riparazioni"}

@api_router.get("/portali")
async def list_portali(user: dict = Depends(get_current_user)):
    return await db.portali.find({}, {"_id": 0}).sort([("sezione", 1), ("operatore", 1)]).to_list(500)

@api_router.post("/portali")
async def create_portale(input: PortaleInput, admin: dict = Depends(require_admin)):
    data = input.model_dump()
    data["operatore"] = data["operatore"].strip().upper()
    data.update({"id": str(uuid.uuid4()), "created_at": datetime.now(timezone.utc).isoformat()})
    await db.portali.insert_one(data)
    data.pop("_id", None)
    return data

@api_router.patch("/portali/{portale_id}")
async def update_portale(portale_id: str, input: PortaleInput, admin: dict = Depends(require_admin)):
    data = input.model_dump()
    data["operatore"] = data["operatore"].strip().upper()
    res = await db.portali.update_one({"id": portale_id}, {"$set": data})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Portale non trovato")
    return await db.portali.find_one({"id": portale_id}, {"_id": 0})

@api_router.delete("/portali/{portale_id}")
async def delete_portale(portale_id: str, admin: dict = Depends(require_admin)):
    await db.portali.delete_one({"id": portale_id})
    return {"status": "ok"}

async def portale_per(sezione: str, operatore: str) -> Optional[dict]:
    if not operatore:
        return None
    return await db.portali.find_one({"sezione": sezione, "operatore": operatore.strip().upper()}, {"_id": 0})

class PortaleFlagInput(BaseModel):
    inserito: bool
    campo: str = "inserito"  # inserito | extra (es. CARICATO SU JOY)

def _portale_flag_update(input: PortaleFlagInput, user: dict) -> dict:
    now = datetime.now(timezone.utc).isoformat()
    at, by = ("portale_extra_at", "portale_extra_da") if input.campo == "extra" else ("portale_inserito_at", "portale_inserito_da")
    return {"$set": {at: now, by: user["name"]}} if input.inserito else {"$unset": {at: "", by: ""}}

@api_router.post("/servizi/{servizio_id}/portale-inserito")
async def servizio_portale_inserito(servizio_id: str, input: PortaleFlagInput, user: dict = Depends(get_current_user)):
    await get_scoped_servizio(servizio_id, user)
    await db.servizi.update_one({"id": servizio_id}, _portale_flag_update(input, user))
    return {"status": "ok"}

@api_router.post("/clients/{client_id}/portale-inserito")
async def client_portale_inserito(client_id: str, input: PortaleFlagInput, user: dict = Depends(get_current_user)):
    await get_scoped_client(client_id, user)
    await db.clients.update_one({"id": client_id}, _portale_flag_update(input, user))
    return {"status": "ok"}

@api_router.get("/portali/flag-scaduti")
async def portali_flag_scaduti(user: dict = Depends(get_current_user), giorni: int = 3):
    """Servizi/clienti con flag aggiuntivo del portale (es. CARICATO SU JOY) non spuntato da oltre N giorni."""
    return await flag_scaduti_items(servizio_scope(user), giorni)

async def flag_scaduti_items(scope: dict, giorni: int = 3) -> list:
    portali = [p for p in await db.portali.find({"flag_label": {"$nin": [None, ""]}}, {"_id": 0}).to_list(200)]
    if not portali:
        return []
    limite = (datetime.now(timezone.utc) - timedelta(days=giorni)).isoformat()
    by_key = {(p["sezione"], p["operatore"]): p for p in portali}
    out = []
    scope = dict(scope)
    scope.update({"tipo": {"$in": ["sim", "internet", "fisso"]}, "portale_extra_at": {"$in": [None]}, "created_at": {"$lte": limite},
                  "operatore_tel": {"$in": sorted({p["operatore"] for p in portali})}})
    svcs = await db.servizi.find(scope, {"_id": 0}).sort("created_at", 1).to_list(2000)
    names = await clients_name_map(list({s["client_id"] for s in svcs}))
    for s in svcs:
        p = by_key.get((PORTALE_SEZIONE_BY_TIPO.get(s["tipo"]), (s.get("operatore_tel") or "").upper()))
        if p:
            out.append({"kind": "servizio", "id": s["id"], "client_id": s["client_id"], "client_name": names.get(s["client_id"], ""),
                        "store_id": s.get("venditore_id", ""),
                        "operatore": s.get("operatore_tel", ""), "tipo": s["tipo"], "flag_label": p["flag_label"], "flag_url": p.get("flag_url", ""),
                        "created_at": s.get("created_at"), "giorni": (datetime.now(timezone.utc) - datetime.fromisoformat(s["created_at"])).days})
    return out

@api_router.get("/portali/da-inserire")
async def portali_da_inserire(user: dict = Depends(get_current_user)):
    portali = await db.portali.find({}, {"_id": 0}).to_list(500)
    if not portali:
        return []
    keys = {(p["sezione"], p["operatore"]) for p in portali if p["operatore"] != "*"}
    out = []
    scope = servizio_scope(user)
    scope["tipo"] = {"$in": ["sim", "internet", "fisso"]}
    scope["portale_inserito_at"] = {"$in": [None]}
    svcs = await db.servizi.find(scope, {"_id": 0}).sort("created_at", -1).to_list(2000)
    names = await clients_name_map(list({s["client_id"] for s in svcs}))
    for s in svcs:
        sez = PORTALE_SEZIONE_BY_TIPO.get(s["tipo"])
        if (sez, (s.get("operatore_tel") or "").upper()) in keys:
            out.append({"kind": "servizio", "id": s["id"], "client_id": s["client_id"], "client_name": names.get(s["client_id"], ""),
                        "sezione": sez, "operatore": s.get("operatore_tel", ""), "created_at": s.get("created_at")})
    cscope = client_scope_filter(user)
    cscope.update({"nuovo_fornitore": {"$nin": [None, ""]}, "portale_inserito_at": {"$in": [None]},
                   "lavorazione": {"$in": ["da_inserire", "inserito", "da_quotare", "quotato"]}})
    for c in await db.clients.find(cscope, {"_id": 0, "id": 1, "nome": 1, "cognome": 1, "nuovo_fornitore": 1, "created_at": 1}).sort("created_at", -1).to_list(2000):
        if ("energia", (c.get("nuovo_fornitore") or "").upper()) in keys:
            out.append({"kind": "cliente", "id": c["id"], "client_id": c["id"], "client_name": f"{c.get('cognome', '')} {c.get('nome', '')}".strip(),
                        "sezione": "energia", "operatore": c.get("nuovo_fornitore", ""), "created_at": c.get("created_at")})
    return out

@api_router.post("/clients/{client_id}/blacklist-recensioni")
async def blacklist_recensioni(client_id: str, input: BlacklistInput, user: dict = Depends(get_current_user)):
    c = await db.clients.find_one({"id": client_id}, {"_id": 0, "id": 1})
    if not c:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    await db.clients.update_one({"id": client_id}, {"$set": {"no_recensioni": input.no_recensioni}})
    if input.no_recensioni:
        await db.whatsapp_queue.delete_many({"client_id": client_id, "type": "review", "sent": False})
        await db.servizi.update_many({"client_id": client_id, "review_msg_sent_at": None}, {"$unset": {"review_queued_at": ""}})
    return {"status": "ok", "no_recensioni": input.no_recensioni}

PROMEMORIA_GIORNI = 7

async def invia_promemoria_ritiro() -> dict:
    limite = (datetime.now(timezone.utc) - timedelta(days=PROMEMORIA_GIORNI)).isoformat()
    svcs = await db.servizi.find({"tipo": "riparazione", "stato": "pronto", "pronto_msg_sent_at": {"$ne": None, "$lte": limite},
                                  "promemoria_msg_sent_at": {"$in": [None]}}, {"_id": 0}).to_list(1000)
    report = []
    for svc in svcs:
        if svc.get("promemoria_msg_sent_at"):
            continue
        client_doc = await db.clients.find_one({"id": svc["client_id"]}, {"_id": 0, "nome": 1, "telefono": 1})
        if not client_doc or not client_doc.get("telefono"):
            continue
        store = await db.stores.find_one({"id": svc.get("venditore_id")}, {"_id": 0})
        msg = store_msg(store, "msg_promemoria", nome=client_doc.get("nome", ""), dispositivo=svc.get("dispositivo", ""),
                        numero=svc.get("numero_riparazione", "-"))
        try:
            await wa_send(client_doc["telefono"], msg, session=svc.get("venditore_id") or "default", tipo="promemoria", client_id=svc["client_id"], servizio_id=svc["id"])
            await db.servizi.update_one({"id": svc["id"]}, {"$set": {"promemoria_msg_sent_at": datetime.now(timezone.utc).isoformat()}})
            report.append({"numero": svc.get("numero_riparazione"), "inviato": True})
        except HTTPException as e:
            report.append({"numero": svc.get("numero_riparazione"), "inviato": False, "errore": str(e.detail)})
    if report:
        await db.cron_log.insert_one({"job": "promemoria-ritiro", "at": datetime.now(timezone.utc).isoformat(), "report": report})
    return {"report": report}

VINCOLO_PREAVVISO_GIORNI = 30

async def invia_avvisi_vincolo() -> dict:
    svcs = await db.servizi.find({"tipo": {"$in": ["sim", "internet", "fisso"]}, "data_attivazione": {"$nin": [None, ""]}},
                                 {"_id": 0}).to_list(10000)
    report = []
    for svc in svcs:
        so = scadenza_offerta(svc)
        if not so:
            continue
        scad, annuale = so
        giorni = (scad - date.today()).days
        if not (0 <= giorni <= VINCOLO_PREAVVISO_GIORNI) or svc.get("vincolo_msg_sent_for") == scad.isoformat():
            continue
        client_doc = await db.clients.find_one({"id": svc["client_id"]}, {"_id": 0, "nome": 1, "cognome": 1, "telefono": 1})
        if not client_doc or not client_doc.get("telefono"):
            continue
        store = await db.stores.find_one({"id": svc.get("venditore_id")}, {"_id": 0})
        msg = store_msg(store, "msg_offerta_annuale" if annuale else "msg_vincolo",
                        nome=client_doc.get("nome", ""), cognome=client_doc.get("cognome", ""))
        try:
            await wa_send(client_doc["telefono"], msg, session=svc.get("venditore_id") or "default",
                          tipo="offerta_annuale" if annuale else "vincolo", client_id=svc["client_id"], servizio_id=svc["id"])
            await db.servizi.update_one({"id": svc["id"]}, {"$set": {"vincolo_msg_sent_for": scad.isoformat(),
                                                                     "vincolo_msg_sent_at": datetime.now(timezone.utc).isoformat()}})
            report.append({"servizio": svc["id"], "inviato": True, "annuale": annuale, "scadenza": scad.isoformat()})
        except HTTPException as e:
            report.append({"servizio": svc["id"], "inviato": False, "errore": str(e.detail)})
    if report:
        await db.cron_log.insert_one({"job": "avvisi-vincolo", "at": datetime.now(timezone.utc).isoformat(), "report": report})
    return {"report": report}

@api_router.post("/telefonia/avvisi-vincolo/invia-ora")
async def avvisi_vincolo_ora(admin: dict = Depends(require_admin)):
    return await invia_avvisi_vincolo()

RINNOVO_PREAVVISO_GIORNI = 60

async def invia_avvisi_rinnovo_energia() -> dict:
    clients = await db.clients.find({"data_contratto": {"$nin": [None, ""]}, "telefono": {"$nin": [None, ""]}}, {"_id": 0}).to_list(20000)
    report = []
    for c in clients:
        compute_dates(c)
        if not c.get("data_scadenza"):
            continue
        giorni = (date.fromisoformat(c["data_scadenza"]) - date.today()).days
        if not (0 <= giorni <= RINNOVO_PREAVVISO_GIORNI) or c.get("rinnovo_msg_sent_for") == c["data_scadenza"]:
            continue
        store = await db.stores.find_one({"id": c.get("venditore_id")}, {"_id": 0})
        msg = store_msg(store, "msg_rinnovo_energia", nome=c.get("nome", ""), cognome=c.get("cognome", ""))
        try:
            await wa_send(c["telefono"], msg, session=await wa_session_cliente(c), tipo="rinnovo_energia", client_id=c["id"])
            await db.clients.update_one({"id": c["id"]}, {"$set": {"rinnovo_msg_sent_for": c["data_scadenza"],
                                                                   "rinnovo_msg_sent_at": datetime.now(timezone.utc).isoformat()}})
            report.append({"client_id": c["id"], "inviato": True, "scadenza": c["data_scadenza"]})
        except HTTPException as e:
            report.append({"client_id": c["id"], "inviato": False, "errore": str(e.detail)})
    if report:
        await db.cron_log.insert_one({"job": "rinnovo-energia", "at": datetime.now(timezone.utc).isoformat(), "report": report})
    return {"report": report}

@api_router.post("/energia/avvisi-rinnovo/invia-ora")
async def avvisi_rinnovo_ora(admin: dict = Depends(require_admin)):
    return await invia_avvisi_rinnovo_energia()

TRUFFE_DOPO_GIORNI = 10
TRUFFE_FINESTRA_GIORNI = 40

async def invia_avvisi_truffe() -> dict:
    clients = await db.clients.find({"data_contratto": {"$nin": [None, ""]}, "telefono": {"$nin": [None, ""]},
                                     "truffe_msg_sent_at": {"$in": [None]}}, {"_id": 0}).to_list(20000)
    report = []
    for c in clients:
        compute_dates(c)
        if not c.get("data_attivazione"):
            continue
        giorni = (date.today() - date.fromisoformat(c["data_attivazione"])).days
        if not (TRUFFE_DOPO_GIORNI <= giorni <= TRUFFE_FINESTRA_GIORNI):
            continue
        store = await db.stores.find_one({"id": c.get("venditore_id")}, {"_id": 0})
        msg = ENEL_TRUFFE_MSG if is_enel(c) else store_msg(store, "msg_truffe", nome=c.get("nome", ""), cognome=c.get("cognome", ""))
        try:
            await wa_send(c["telefono"], msg, session=await wa_session_cliente(c), tipo="truffe", client_id=c["id"])
            await db.clients.update_one({"id": c["id"]}, {"$set": {"truffe_msg_sent_at": datetime.now(timezone.utc).isoformat()}})
            report.append({"client_id": c["id"], "inviato": True})
        except HTTPException as e:
            report.append({"client_id": c["id"], "inviato": False, "errore": str(e.detail)})
    if report:
        await db.cron_log.insert_one({"job": "avviso-truffe", "at": datetime.now(timezone.utc).isoformat(), "report": report})
    return {"report": report}

@api_router.post("/energia/avvisi-truffe/invia-ora")
async def avvisi_truffe_ora(admin: dict = Depends(require_admin)):
    return await invia_avvisi_truffe()

PRONTO_MSG = MSG_DEFAULTS["msg_pronto"]

async def invia_avviso_pronto(svc: dict) -> None:
    client_doc = await db.clients.find_one({"id": svc["client_id"]}, {"_id": 0, "nome": 1, "telefono": 1})
    if not client_doc or not client_doc.get("telefono"):
        await db.servizi.update_one({"id": svc["id"]}, {"$set": {"pronto_msg_error": "Cliente senza numero di telefono"}})
        return
    store = await db.stores.find_one({"id": svc.get("venditore_id")}, {"_id": 0})
    msg = store_msg(store, "msg_pronto", nome=client_doc.get("nome", ""), dispositivo=svc.get("dispositivo", "dispositivo"),
                    numero=svc.get("numero_riparazione", "-"))
    try:
        await wa_send(client_doc["telefono"], msg, session=svc.get("venditore_id") or "default", tipo="pronto", client_id=svc["client_id"], servizio_id=svc["id"])
        await db.servizi.update_one({"id": svc["id"]}, {"$set": {"pronto_msg_sent_at": datetime.now(timezone.utc).isoformat(),
                                                                 "pronto_msg_error": None}})
    except HTTPException as e:
        await db.servizi.update_one({"id": svc["id"]}, {"$set": {"pronto_msg_error": str(e.detail)}})

@api_router.post("/servizi/{servizio_id}/whatsapp/pronto")
async def whatsapp_pronto_servizio(servizio_id: str, user: dict = Depends(get_current_user)):
    svc = await get_scoped_servizio(servizio_id, user)
    if svc.get("tipo") != "riparazione":
        raise HTTPException(status_code=400, detail="Solo per riparazioni")
    await invia_avviso_pronto(svc)
    updated = await db.servizi.find_one({"id": servizio_id}, {"_id": 0, "pronto_msg_sent_at": 1, "pronto_msg_error": 1})
    if updated.get("pronto_msg_error"):
        raise HTTPException(status_code=400, detail=f"WhatsApp: {updated['pronto_msg_error']}")
    return {"status": "ok", "pronto_msg_sent_at": updated.get("pronto_msg_sent_at")}

@api_router.delete("/servizi/{servizio_id}")
async def delete_servizio(servizio_id: str, admin: dict = Depends(require_admin)):
    res = await db.servizi.delete_one({"id": servizio_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Servizio non trovato")
    await db.attachments.update_many({"servizio_id": servizio_id}, {"$set": {"is_deleted": True}})
    return {"status": "ok"}

@api_router.post("/servizi/{servizio_id}/mark-paid")
async def mark_servizio_paid(servizio_id: str, user: dict = Depends(get_current_user)):
    await get_scoped_servizio(servizio_id, user)
    await db.servizi.update_one({"id": servizio_id},
                                {"$set": {"pagato": True, "last_payment_date": date.today().isoformat()}})
    return {"status": "ok"}

SVC_CATEGORIA = {"riparazione": ("TELEFONIA", "Riparazione"), "accessori": ("TELEFONIA", "Accessori"),
                 "vendita": ("TELEFONIA", "Vendita"), "sim": ("SIM", None),
                 "internet": ("INTERNET", None), "fisso": ("INTERNET", None)}

async def register_privacy_servizio(client: dict, svc: dict, store_name: str) -> dict:
    cat, det = SVC_CATEGORIA.get(svc["tipo"], ("CambiaOra", "CambiaOra"))
    if det is None:
        op = (svc.get("operatore_tel") or "").strip().upper()
        allowed = SIM_OPERATORS if svc["tipo"] == "sim" else [o.upper() for o in INTERNET_OPERATORS]
        det = op if op in allowed else allowed[0]
    return await register_privacy_site(client, store_name, categoria=cat, dettaglio=det)

def review_message_for_store(store: Optional[dict]) -> Optional[str]:
    if not store or not store.get("review_link"):
        return None
    return store_msg(store, "msg_recensione")

@api_router.post("/servizi/{servizio_id}/whatsapp/privacy")
async def whatsapp_privacy_servizio(servizio_id: str, user: dict = Depends(get_current_user)):
    svc = await get_scoped_servizio(servizio_id, user)
    client_doc = await db.clients.find_one({"id": svc["client_id"]}, {"_id": 0})
    if not client_doc or not client_doc.get("telefono"):
        raise HTTPException(status_code=400, detail="Il cliente non ha un numero di telefono")
    store = await db.stores.find_one({"id": svc.get("venditore_id")}, {"_id": 0})
    registered, reg_error = False, None
    if client_doc.get("email"):
        try:
            await register_privacy_servizio(client_doc, svc, (store or {}).get("nome", ""))
            registered = True
        except Exception as e:
            reg_error = getattr(e, "detail", str(e))
            logger.error(f"Registrazione privacy sito fallita per servizio {servizio_id}: {reg_error}")
    wa_error = None
    try:
        await wa_send(client_doc["telefono"], store_msg(store, "msg_privacy", nome=client_doc.get("nome", "")), session=client_doc.get("venditore_id") or "default", tipo="privacy", client_id=svc["client_id"], servizio_id=servizio_id)
    except Exception as e:
        wa_error = getattr(e, "detail", str(e))
    if wa_error and not registered:
        raise HTTPException(status_code=502, detail=f"WhatsApp: {wa_error}")
    now = datetime.now(timezone.utc)
    updates = {}
    review_queued = False
    if not wa_error:
        updates["privacy_msg_sent_at"] = now.isoformat()
        review_msg = review_message_for_store(store)
        # Per le riparazioni la recensione parte alla consegna del dispositivo, non dopo la privacy
        if review_msg and svc.get("tipo") != "riparazione" and not client_doc.get("no_recensioni"):
            await db.whatsapp_queue.insert_one({
                "id": str(uuid.uuid4()), "client_id": svc["client_id"], "servizio_id": servizio_id,
                "phone": client_doc["telefono"], "type": "review", "message": review_msg,
                "session": svc.get("venditore_id") or "default",
                "send_after": (now + timedelta(minutes=2)).isoformat(), "sent": False,
                "created_at": now.isoformat()})
            review_queued = True
    if registered:
        updates["privacy_firmata"] = True
        updates["privacy_registered_at"] = now.isoformat()
    if updates:
        await db.servizi.update_one({"id": servizio_id}, {"$set": updates})
        await db.clients.update_one({"id": svc["client_id"]},
                                    {"$set": {"privacy_firmata": True}} if registered else {"$set": {}})
    return {"status": "ok", "privacy_registered": registered, "registration_error": reg_error,
            "wa_error": wa_error, "review_queued": review_queued}

@api_router.post("/servizi/{servizio_id}/attachments")
async def upload_servizio_attachment(servizio_id: str, file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    await get_scoped_servizio(servizio_id, user)
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else "bin"
    if ext not in ("pdf", "jpg", "jpeg", "png", "webp"):
        raise HTTPException(status_code=400, detail="Formato non supportato (PDF, JPG, PNG, WEBP)")
    data = await file.read()
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File troppo grande (max 15 MB)")
    path = f"{APP_NAME}/servizi/{servizio_id}/{uuid.uuid4()}.{ext}"
    result = put_object(path, data, file.content_type or "application/octet-stream")
    doc = {"id": str(uuid.uuid4()), "servizio_id": servizio_id, "storage_path": result["path"],
           "original_filename": file.filename, "content_type": file.content_type,
           "size": result["size"], "is_deleted": False, "uploaded_by": user["id"],
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.attachments.insert_one(doc)
    doc.pop("_id", None)
    return doc

@api_router.get("/servizi/{servizio_id}/attachments")
async def list_servizio_attachments(servizio_id: str, user: dict = Depends(get_current_user)):
    await get_scoped_servizio(servizio_id, user)
    return await db.attachments.find({"servizio_id": servizio_id, "is_deleted": False}, {"_id": 0}).to_list(200)

@api_router.get("/clients/{client_id}/servizi")
async def client_servizi(client_id: str, user: dict = Depends(get_current_user)):
    await get_scoped_client(client_id, user)
    servizi = await db.servizi.find({"client_id": client_id}, {"_id": 0}).sort("created_at", -1).to_list(500)
    return [serialize_servizio(s) for s in servizi]

@api_router.get("/vincoli")
async def report_vincoli(user: dict = Depends(get_current_user), giorni: int = 90):
    if "telefonia" not in user_sections(user):
        raise HTTPException(status_code=403, detail="Sezione non abilitata")
    scope = servizio_scope(user)
    scope["tipo"] = {"$in": [t for t in ["sim", "internet", "fisso"] if t in (scope["tipo"].get("$in", []) if isinstance(scope.get("tipo"), dict) else [scope.get("tipo")])]}
    scope["data_attivazione"] = {"$ne": None}
    servizi = await db.servizi.find(scope, {"_id": 0}).to_list(5000)
    names = await clients_name_map(list({s["client_id"] for s in servizi}))
    out = []
    for s in servizi:
        ser = serialize_servizio(s, names.get(s["client_id"], ""))
        if ser.get("giorni_alla_scadenza") is not None and ser["giorni_alla_scadenza"] <= giorni:
            out.append(ser)
    out.sort(key=lambda x: x["giorni_alla_scadenza"])
    return out

@api_router.get("/telefonia/proposte")
async def telefonia_proposte(user: dict = Depends(get_current_user)):
    """Clienti con mobile ma senza fisso e senza contratto energia: da proporre."""
    if "telefonia" not in user_sections(user):
        raise HTTPException(status_code=403, detail="Sezione non abilitata")
    scope = servizio_scope(user)
    scope["tipo"] = {"$in": ["sim", "internet", "fisso"]}
    servizi = await db.servizi.find(scope, {"_id": 0, "client_id": 1, "tipo": 1, "operatore_tel": 1, "venditore_id": 1}).to_list(10000)
    per_client: dict = {}
    for s in servizi:
        per_client.setdefault(s["client_id"], {"tipi": set(), "venditore_id": s.get("venditore_id", ""), "operatore": s.get("operatore_tel", "")})
        per_client[s["client_id"]]["tipi"].add(s["tipo"])
    ids = [cid for cid, v in per_client.items() if "sim" in v["tipi"] and not ({"fisso", "internet"} & v["tipi"])]
    if not ids:
        return []
    clients = await db.clients.find({"id": {"$in": ids}}, {"_id": 0, "id": 1, "nome": 1, "cognome": 1, "telefono": 1,
                                                          "lavorazione": 1, "nuovo_fornitore": 1, "no_recensioni": 1}).to_list(10000)
    out = []
    for c in clients:
        ha_energia = c.get("lavorazione") == "cambio_effettuato" or bool(c.get("nuovo_fornitore"))
        out.append({"client_id": c["id"], "client_name": f"{c.get('cognome', '')} {c.get('nome', '')}".strip(),
                    "telefono": c.get("telefono", ""), "operatore_mobile": per_client[c["id"]]["operatore"],
                    "venditore_id": per_client[c["id"]]["venditore_id"],
                    "manca_fisso": True, "manca_energia": not ha_energia})
    return sorted(out, key=lambda x: (not x["manca_energia"], x["client_name"]))

# ---------------- Cron ----------------

@api_router.post("/cron/daily-digest")
async def cron_daily_digest(request: Request, background_tasks: BackgroundTasks):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else ""
    secret = os.environ.get("WEBHOOK_CRON_SECRET", "")
    if not token or not secret or not hmac.compare_digest(token, secret):
        raise HTTPException(status_code=401, detail="Unauthorized")
    background_tasks.add_task(send_digest_email)
    return {"status": "accepted"}

@api_router.post("/cron/whatsapp-due")
async def cron_whatsapp_due(request: Request, background_tasks: BackgroundTasks):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else ""
    secret = os.environ.get("WEBHOOK_CRON_SECRET", "")
    if not token or not secret or not hmac.compare_digest(token, secret):
        raise HTTPException(status_code=401, detail="Unauthorized")
    background_tasks.add_task(process_whatsapp_queue)
    return {"status": "accepted"}

RIP_FERME_GIORNI = 7

async def riparazioni_ferme_per_negozio() -> dict:
    today = date.today()
    rips = await db.servizi.find({"tipo": "riparazione", "stato": {"$nin": ["consegnato", "non_riparabile"]}},
                                 {"_id": 0, "venditore_id": 1, "numero_riparazione": 1, "dispositivo": 1,
                                  "stato": 1, "client_id": 1, "data_ingresso": 1, "created_at": 1}).to_list(10000)
    names = await clients_name_map(list({r["client_id"] for r in rips}))
    out: dict = {}
    for r in rips:
        try:
            ingresso = date.fromisoformat(str(r.get("data_ingresso") or r.get("created_at"))[:10])
        except (ValueError, TypeError):
            continue
        giorni = (today - ingresso).days
        if giorni > RIP_FERME_GIORNI:
            out.setdefault(r.get("venditore_id", ""), []).append({
                "numero": r.get("numero_riparazione", "-"), "dispositivo": r.get("dispositivo", ""),
                "cliente": names.get(r["client_id"], ""), "stato": r.get("stato", ""), "giorni": giorni})
    for lst in out.values():
        lst.sort(key=lambda x: -x["giorni"])
    return out

RIP_STATO_LABEL = {"ingresso": "Ingresso", "attesa_ricambio_cliente": "Attesa ricambio (cliente)",
                   "attesa_ricambio_carico": "Attesa ricambio (in carico)", "in_attesa_cliente": "In attesa cliente",
                   "preventivo": "Preventivo", "in_lavorazione": "In lavorazione", "pronto": "Pronto"}

def messaggio_riparazioni_ferme(store_name: str, items: list, joy: Optional[list] = None) -> str:
    parti = [f"Buongiorno {store_name}!"]
    if items:
        righe = [f"- {i['numero']} {i['dispositivo']} ({i['cliente']}) - {RIP_STATO_LABEL.get(i['stato'], i['stato'])} - {i['giorni']} gg"
                 for i in items[:25]]
        extra = f"\n...e altre {len(items) - 25}" if len(items) > 25 else ""
        parti.append(f"Riparazioni ferme da oltre {RIP_FERME_GIORNI} giorni: {len(items)}\n" + "\n".join(righe) + extra)
    if joy:
        righe = [f"- {j['client_name']} ({j['operatore']} {j['tipo']}) - inserito {j['giorni']} gg fa" for j in joy[:25]]
        extra = f"\n...e altri {len(joy) - 25}" if len(joy) > 25 else ""
        parti.append(f"Contratti {joy[0]['operatore']} ancora da caricare su {joy[0]['flag_label'].replace('CARICATO SU ', '')}: {len(joy)}\n"
                     + "\n".join(righe) + extra)
    parti.append("Controllali nel gestionale, grazie.")
    return "\n".join(parti)

async def invia_avvisi_riparazioni_ferme() -> dict:
    ferme = await riparazioni_ferme_per_negozio()
    joy_per_store: dict = {}
    for j in await flag_scaduti_items({}, 3):
        joy_per_store.setdefault(j["store_id"], []).append(j)
    stores = await db.stores.find({"telefono_avvisi": {"$nin": [None, ""]}}, {"_id": 0}).to_list(200)
    report = []
    for s in stores:
        items = ferme.get(s["id"], [])
        joy = joy_per_store.get(s["id"], [])
        if not items and not joy:
            report.append({"store": s["nome"], "inviato": False, "ferme": 0, "joy": 0})
            continue
        try:
            await wa_send(s["telefono_avvisi"], messaggio_riparazioni_ferme(s["nome"], items, joy), session=s["id"], tipo="avviso_negozio")
            report.append({"store": s["nome"], "inviato": True, "ferme": len(items), "joy": len(joy)})
        except HTTPException as e:
            report.append({"store": s["nome"], "inviato": False, "ferme": len(items), "joy": len(joy), "errore": str(e.detail)})
    await db.cron_log.insert_one({"job": "riparazioni-ferme", "at": datetime.now(timezone.utc).isoformat(), "report": report})
    return {"report": report}

@api_router.post("/cron/riparazioni-ferme")
async def cron_riparazioni_ferme(request: Request, background_tasks: BackgroundTasks):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else ""
    secret = os.environ.get("WEBHOOK_CRON_SECRET", "")
    if not token or not secret or not hmac.compare_digest(token, secret):
        raise HTTPException(status_code=401, detail="Unauthorized")
    background_tasks.add_task(invia_avvisi_riparazioni_ferme)
    background_tasks.add_task(invia_promemoria_ritiro)
    background_tasks.add_task(invia_avvisi_vincolo)
    background_tasks.add_task(invia_avvisi_rinnovo_energia)
    background_tasks.add_task(invia_avvisi_truffe)
    background_tasks.add_task(pulizia_retention_gdpr)
    return {"status": "accepted"}

@api_router.get("/riparazioni-ferme")
async def get_riparazioni_ferme(user: dict = Depends(get_current_user)):
    scope = servizio_scope(user, "riparazione")
    ferme = await riparazioni_ferme_per_negozio()
    allowed = scope.get("venditore_id", {}).get("$in") if isinstance(scope.get("venditore_id"), dict) else None
    stores = {s["id"]: s for s in await db.stores.find({}, {"_id": 0, "id": 1, "nome": 1, "telefono_avvisi": 1}).to_list(200)}
    return [{"store_id": sid, "store_name": stores.get(sid, {}).get("nome", "-"),
             "telefono_avvisi": stores.get(sid, {}).get("telefono_avvisi", ""), "items": items}
            for sid, items in ferme.items() if allowed is None or sid in allowed]

@api_router.post("/riparazioni-ferme/invia-ora")
async def invia_riparazioni_ferme_ora(admin: dict = Depends(require_admin)):
    return await invia_avvisi_riparazioni_ferme()

@api_router.get("/")
async def root():
    return {"message": "Gestionale Utenze API"}

# ---------------- Seed ----------------

SEED_STORES = [
    {"nome": "Tirano", "referente": "Michael", "tipo": "negozio"},
    {"nome": "Sondalo", "referente": "Lorenzo", "tipo": "negozio"},
    {"nome": "Sondrio", "referente": "Michael", "tipo": "negozio"},
    {"nome": "Grosio", "referente": "Michael", "tipo": "negozio"},
    {"nome": "Deriu", "referente": "Deriu", "tipo": "negozio"},
    {"nome": "Gravedona", "referente": "Kevin", "tipo": "negozio"},
    {"nome": "Devis (Freelance)", "referente": "Devis", "tipo": "freelance"},
]

SEED_USERS = [
    {"name": "Deborah", "email": "deborah@cambiaora.local", "password": "Deborah2026!",
     "role": "operatore", "can_view_all": True, "stores": []},
    {"name": "Michael", "email": "michael@cambiaora.local", "password": "Michael2026!",
     "role": "negozio", "can_view_all": False, "stores": ["Tirano"]},
    {"name": "Seba", "email": "seba@cambiaora.local", "password": "Seba2026!",
     "role": "negozio", "can_view_all": False, "stores": ["Grosio"]},
    {"name": "Lorenzo", "email": "lorenzo@cambiaora.local", "password": "Lorenzo2026!",
     "role": "negozio", "can_view_all": False, "stores": ["Sondalo"]},
    {"name": "Kevin", "email": "kevin@cambiaora.local", "password": "Kevin2026!",
     "role": "negozio", "can_view_all": False, "stores": ["Gravedona"]},
]

SEED_CLIENT_SAMPLES = [
    ("Mario", "Rossi", "privato", "luce", "Enel Energia", "cambio_effettuato", "Tirano", 10, True, 7),
    ("Giuseppe", "Bianchi", "privato", "gas", "Eni Plenitude", "rinnovato", "Sondalo", 11, True, 2),
    ("Valtellina Eco", "Srl", "business", "luce", "A2A Energia", "in_quotazione", "Tirano", 1, False, None),
    ("Anna", "Verdi", "privato", "luce", "Sorgenia", "richieste_bollette", "Gravedona", 3, False, None),
    ("Luca", "Fontana", "privato", "gas", "Hera Comm", "attesa_documenti", "Sondrio", 9, True, 5),
    ("Bar al Lago", "Sas", "business", "luce", "Illumia", "contattare_cliente", "Gravedona", 0, False, None),
    ("Paola", "Neri", "privato", "gas", "NeN", "da_quotare", "Sondalo", 2, False, None),
    ("Franco", "Colombo", "privato", "luce", "Octopus Energy", "in_attesa_ok", "Deriu", 10, True, 4),
    ("Agriturismo Pizzo", "Srl", "business", "gas", "Dolomiti Energia", "passa_in_negozio", "Grosio", 6, False, None),
    ("Sara", "Galli", "privato", "luce", "Tate", "problema_tecnico", "Tirano", 4, False, None),
]

async def _seed_indexes():
    await db.users.create_index("email", unique=True)
    await db.login_attempts.create_index("identifier")
    await db.clients.create_index("venditore_id")
    await db.clients.create_index("lavorazione")
    await db.lavorazioni_log.create_index("operatore_id")

async def _seed_admin():
    admin_email = os.environ["ADMIN_EMAIL"].strip().lower()
    admin_password = os.environ["ADMIN_PASSWORD"]
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        await db.users.insert_one({"id": str(uuid.uuid4()), "name": "Enrico", "email": admin_email,
                                   "password_hash": hash_password(admin_password), "role": "admin",
                                   "store_ids": [], "can_view_all": True, "active": True,
                                   "created_at": datetime.now(timezone.utc).isoformat()})
    elif not verify_password(admin_password, existing["password_hash"]):
        await db.users.update_one({"email": admin_email},
                                  {"$set": {"password_hash": hash_password(admin_password)}})

async def _seed_stores():
    if await db.stores.count_documents({}) == 0:
        for s in SEED_STORES:
            await db.stores.insert_one({"id": str(uuid.uuid4()), **s, "note": "", "pagato": False,
                                        "last_payment_date": None,
                                        "created_at": datetime.now(timezone.utc).isoformat()})

async def _seed_users():
    stores = {s["nome"]: s["id"] for s in await db.stores.find({}, {"_id": 0}).to_list(100)}
    for su in SEED_USERS:
        if not await db.users.find_one({"email": su["email"]}):
            await db.users.insert_one({"id": str(uuid.uuid4()), "name": su["name"], "email": su["email"],
                                       "password_hash": hash_password(su["password"]), "role": su["role"],
                                       "store_ids": [stores[n] for n in su["stores"] if n in stores],
                                       "can_view_all": su["can_view_all"], "active": True,
                                       "created_at": datetime.now(timezone.utc).isoformat()})

def _sample_client_doc(nome, cognome, tipo_c, bolletta, forn, lav, store_id, dc, pagato, last_pay, op_id, now):
    return {"id": str(uuid.uuid4()), "nome": nome, "cognome": cognome, "tipo_cliente": tipo_c,
            "codice_fiscale": "RSSMRA80A01F205X" if tipo_c == "privato" else "",
            "p_iva": "01234567890" if tipo_c == "business" else "",
            "indirizzo": "Via Roma 1, 23100 Sondrio", "pod": "IT001E12345678" if bolletta == "luce" else "",
            "pdr": "12345678901234" if bolletta == "gas" else "", "iban": "IT60X0542811101000000123456",
            "email": f"{nome.lower().replace(' ', '')}@esempio.it", "telefono": "3331234567",
            "kw_potenza": 3.0 if bolletta == "luce" else None, "tipo_bolletta": bolletta,
            "fornitore_provenienza": forn,
            "costo_kwh_attuale": 0.32 if bolletta == "luce" else None,
            "spese_fisse_attuale": 10.0,
            "costo_smc_attuale": 1.15 if bolletta == "gas" else None,
            "data_contratto": dc.isoformat(), "data_verifica": None, "data_cambio": None,
            "tipo_contratto": "fisso", "nuovo_fornitore": "Sorgenia",
            "costo_kwh_nuovo": 0.24 if bolletta == "luce" else None, "spese_fisse_nuovo": 8.0,
            "costo_smc_nuovo": 0.89 if bolletta == "gas" else None,
            "privacy_firmata": True, "note": "", "lavorazione": lav,
            "venditore_id": store_id, "operatore_id": op_id,
            "pagato": pagato, "last_payment_date": last_pay,
            "created_by": op_id, "created_at": now, "updated_at": now}

async def _seed_clients():
    if await db.clients.count_documents({}) > 0:
        return
    deborah = await db.users.find_one({"email": "deborah@cambiaora.local"}, {"_id": 0})
    op_id = deborah["id"] if deborah else ""
    stores = {s["nome"]: s["id"] for s in await db.stores.find({}, {"_id": 0}).to_list(100)}
    today = date.today()
    now = datetime.now(timezone.utc).isoformat()
    for nome, cognome, tipo_c, bolletta, forn, lav, store_name, mesi_fa, pagato, pag_mesi_fa in SEED_CLIENT_SAMPLES:
        dc = today - relativedelta(months=mesi_fa)
        last_pay = (today - relativedelta(months=pag_mesi_fa)).isoformat() if pag_mesi_fa is not None else None
        doc = _sample_client_doc(nome, cognome, tipo_c, bolletta, forn, lav,
                                 stores.get(store_name, ""), dc, pagato, last_pay, op_id, now)
        await db.clients.insert_one(doc)
        await db.lavorazioni_log.insert_one({
            "id": str(uuid.uuid4()), "client_id": doc["id"],
            "client_name": f"{cognome} {nome}", "operatore_id": op_id,
            "operatore_name": "Deborah", "status": lav, "note": "Cliente inserito",
            "created_at": now})

async def _seed_stores_v2():
    if not await db.stores.find_one({"nome": "Morbegno"}):
        await db.stores.insert_one({"id": str(uuid.uuid4()), "nome": "Morbegno", "referente": "",
                                    "tipo": "negozio", "note": "", "review_link": "", "pagato": False,
                                    "last_payment_date": None,
                                    "created_at": datetime.now(timezone.utc).isoformat()})
    review_links = {"Tirano": "https://g.page/r/CQedddebLpxpEAE/review",
                    "Sondrio": "https://g.page/r/CeQBdxW0AtlqEAE/review",
                    "Gravedona": "https://g.page/r/CWOpR0o29GE-EAE/review",
                    "Morbegno": "https://g.page/r/CT15UUSXSQoCEAE/review"}
    for nome, link in review_links.items():
        await db.stores.update_one({"nome": nome, "$or": [{"review_link": {"$exists": False}}, {"review_link": ""}]},
                                   {"$set": {"review_link": link}})

async def _seed_user_sections():
    await db.users.update_many({"sections": {"$exists": False}},
                               {"$set": {"sections": ["energia", "riparazioni", "telefonia"]}})

async def seed_data():
    await _seed_indexes()
    await _seed_admin()
    await _seed_stores()
    await _seed_users()
    if os.environ.get("ENVIRONMENT", "production") == "development":
        await _seed_clients()
    await _seed_stores_v2()
    await _seed_user_sections()

@app.on_event("startup")
async def startup():
    await seed_data()
    await seed_formazione()
    await load_regole_prezzi()
    await seed_portali()
    await migra_segreti_in_chiaro()
    try:
        init_storage()
        logger.info("Object storage inizializzato")
    except Exception as e:
        logger.error(f"Storage init fallito: {e}")

app.include_router(api_router)

# ---------------- Registro accessi (audit GDPR) ----------------
async def migra_segreti_in_chiaro() -> None:
    """Cifra i codici sblocco/password salvati in chiaro e cancella quelli di riparazioni chiuse."""
    cur = db.servizi.find({"$or": [{"codice_sblocco": {"$nin": [None, ""]}}, {"account_password": {"$nin": [None, ""]}}]}, {"_id": 0})
    n = 0
    async for s in cur:
        if s.get("stato") in ("consegnato", "non_riparabile"):
            await db.servizi.update_one({"id": s["id"]}, cancella_segreti_update())
        else:
            upd = {f + "_enc": totp_encrypt(s[f]) for f in SECRET_FIELDS if s.get(f)}
            await db.servizi.update_one({"id": s["id"]}, {"$set": upd, "$unset": {f: "" for f in SECRET_FIELDS}})
        n += 1
    if n:
        logger.info(f"Segreti dispositivi migrati/cancellati: {n}")

AUDIT_ENTITIES = [
    (re.compile(r"^/api/clients/([^/]+)/anonimizza$"), "cliente", "delete"),
    (re.compile(r"^/api/clients/([^/]+)/gdpr-export$"), "cliente", "export"),
    (re.compile(r"^/api/clients/([^/]+)/whatsapp-log$"), "cliente", "view"),
    (re.compile(r"^/api/clients/([^/]+)/messaggi-previsti$"), None, None),
    (re.compile(r"^/api/clients/([^/]+)/blacklist-recensioni$"), "cliente", "update"),
    (re.compile(r"^/api/clients/([^/]+)/(whatsapp|allegati|foto)"), "cliente", "update"),
    (re.compile(r"^/api/clients/([^/]+)$"), "cliente", None),
    (re.compile(r"^/api/clients$"), "cliente", None),
    (re.compile(r"^/api/servizi/([^/]+)/segreti$"), "servizio", "view_segreti"),
    (re.compile(r"^/api/passwords/([^/]+)/reveal$"), "password", "view_segreti"),
    (re.compile(r"^/api/passwords/import-sheet$"), "password", "import"),
    (re.compile(r"^/api/passwords/([^/]+)$"), "password", None),
    (re.compile(r"^/api/passwords$"), "password", None),
    (re.compile(r"^/api/servizi/([^/]+)/whatsapp-log$"), "servizio", "view"),
    (re.compile(r"^/api/servizi/([^/]+)/(whatsapp|allegati|foto|ricambi|scheda)"), "servizio", "update"),
    (re.compile(r"^/api/servizi/([^/]+)$"), "servizio", None),
    (re.compile(r"^/api/servizi$"), "servizio", None),
    (re.compile(r"^/api/ritiri/([^/]+)/pdf$"), "ritiro", "export"),
    (re.compile(r"^/api/ritiri/([^/]+)$"), "ritiro", None),
    (re.compile(r"^/api/ritiri$"), "ritiro", None),
    (re.compile(r"^/api/export/"), "clienti", "export"),
    (re.compile(r"^/api/(magazzino|ritiri)/export$"), "export", "export"),
    (re.compile(r"^/api/import/"), "clienti", "import"),
    (re.compile(r"^/api/users/([^/]+)"), "utente", None),
    (re.compile(r"^/api/users$"), "utente", None),
    (re.compile(r"^/api/auth/login/mfa$"), "accesso", "login"),
    (re.compile(r"^/api/auth/logout$"), "accesso", "logout"),
    (re.compile(r"^/api/auth/2fa/(enroll/confirm|disable)$"), "sicurezza", "update"),
]
METHOD_ACTION = {"POST": "create", "PATCH": "update", "PUT": "update", "DELETE": "delete", "GET": "view"}

def _audit_classify(method: str, path: str):
    for rx, entity, action in AUDIT_ENTITIES:
        m = rx.match(path)
        if not m:
            continue
        if entity is None:
            return None
        act = action or METHOD_ACTION.get(method)
        if act == "view" and not m.groups():
            return None  # elenco: non tracciato (troppo rumore)
        return entity, act, (m.group(1) if m.groups() else "")
    return None

@app.middleware("http")
async def audit_middleware(request: Request, call_next):
    response = await call_next(request)
    if request.method == "OPTIONS" or response.status_code >= 400:
        return response
    cls = _audit_classify(request.method, request.url.path)
    if not cls:
        return response
    user = getattr(request.state, "user", None)
    if not user and request.url.path == "/api/auth/login/mfa":
        user = getattr(request.state, "mfa_user", None)
    if not user:
        return response
    entity, action, entity_id = cls
    label = ""
    try:
        if entity == "cliente" and entity_id:
            c = await db.clients.find_one({"id": entity_id}, {"_id": 0, "nome": 1, "cognome": 1})
            label = f"{c.get('cognome', '')} {c.get('nome', '')}".strip() if c else ""
        elif entity == "servizio" and entity_id:
            s = await db.servizi.find_one({"id": entity_id}, {"_id": 0, "numero_riparazione": 1, "tipo": 1, "dispositivo": 1})
            label = f"{s.get('numero_riparazione') or s.get('tipo', '')} {s.get('dispositivo', '')}".strip() if s else ""
        elif entity == "ritiro" and entity_id:
            r = await db.ritiri.find_one({"id": entity_id}, {"_id": 0, "numero": 1, "cognome": 1})
            label = f"{r.get('numero', '')} {r.get('cognome', '')}".strip() if r else ""
        elif entity == "utente" and entity_id:
            u = await db.users.find_one({"id": entity_id}, {"_id": 0, "name": 1})
            label = u.get("name", "") if u else ""
        elif entity == "password" and entity_id:
            p = await db.passwords.find_one({"id": entity_id}, {"_id": 0, "servizio": 1, "titolo": 1})
            label = f"{p.get('servizio', '')} {p.get('titolo', '')}".strip() if p else ""
    except Exception:
        pass
    fwd = request.headers.get("x-forwarded-for", "")
    await db.audit_log.insert_one({
        "id": str(uuid.uuid4()), "at": datetime.now(timezone.utc).isoformat(),
        "user_id": user["id"], "user_name": user.get("name", ""), "user_role": user.get("role", ""),
        "action": action, "entity": entity, "entity_id": entity_id, "label": label,
        "method": request.method, "path": request.url.path[:200],
        "ip": (fwd.split(",")[0].strip() if fwd else (request.client.host if request.client else "")),
    })
    return response

@api_router.get("/audit-log")
async def audit_log_list(admin: dict = Depends(require_admin), q: str = "", user_id: str = "", action: str = "",
                         entity: str = "", dal: str = "", al: str = "", limit: int = 200):
    f: dict = {}
    if user_id:
        f["user_id"] = user_id
    if action:
        f["action"] = action
    if entity:
        f["entity"] = entity
    if dal or al:
        f["at"] = {}
        if dal:
            f["at"]["$gte"] = dal
        if al:
            f["at"]["$lte"] = al + "T23:59:59"
    if q:
        f["$or"] = [{"label": {"$regex": re.escape(q), "$options": "i"}}, {"user_name": {"$regex": re.escape(q), "$options": "i"}},
                    {"path": {"$regex": re.escape(q), "$options": "i"}}]
    rows = await db.audit_log.find(f, {"_id": 0}).sort("at", -1).to_list(min(max(limit, 1), 1000))
    return rows

app.include_router(api_router)

cors_origins = os.environ.get("CORS_ORIGINS", "*")
allow_origins = ["*"] if cors_origins == "*" else [o.strip() for o in cors_origins.split(",")]
extra_origins = [os.environ.get("FRONTEND_URL", ""), "http://localhost:3000"]
for o in extra_origins:
    if o and o not in allow_origins and allow_origins != ["*"]:
        allow_origins.append(o)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=allow_origins != ["*"],
    allow_origins=allow_origins,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
